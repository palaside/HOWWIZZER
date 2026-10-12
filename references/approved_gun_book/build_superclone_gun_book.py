# -*- coding: utf-8 -*-
"""
Superclone Production-Grade Engine for สมุดประวัติการใช้ปืนใหญ่ (6.สมุดประวัติการใช้ปืนใหญ่.pdf)
Form: สมุดประวัติการใช้ปืนใหญ่ (๕ หน้า - หัวข้อละ ๑ หน้าพอดีเป๊ะ)
Specifications:
- 5-Page A4 Landscape (297mm x 210mm)
- Page 1: หน้าปก (พร้อมตราสัญลักษณ์กองพันทหารปืนใหญ่ที่ ๔)
- Page 2: ประวัติอาวุธ (ป.๖)
- Page 3: การคำนวณอายุลำกล้อง (ข้อมูลลำกล้อง + ตาราง ๔ คอลัมน์ ๑๒ แถว)
- Page 4: ประวัติการใช้ปืนใหญ่ (ตารางปฏิบัติการและซ่อมบำรุง หัวซ้อน ๖ คอลัมน์ ๑๓ แถว)
- Page 5: บันทึกหลักฐานการยิง (ตารางบันทึกการยิง หัวซ้อน ๙ คอลัมน์ ๑๓ แถว)
"""

import os, sys, json
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

WORKDIR = r"d:\Project\WEPAPP"

with open(os.path.join(WORKDIR, "table_layout_gun_book.json"), "r", encoding="utf-8") as f:
    blueprint = json.load(f)

# ----------------------------------------------------------------------
# 1. HTML Generation (A4 Landscape, 5 Pages Exactly)
# ----------------------------------------------------------------------
def generate_html():
    html_content = f"""<!DOCTYPE html>
<html lang="th">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>สมุดประวัติการใช้ปืนใหญ่</title>
    <link href="https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap" rel="stylesheet">
    <style>
        @page {{
            size: A4 landscape;
            margin: 10mm 15mm;
        }}
        * {{
            box-sizing: border-box;
            font-family: 'TH Sarabun PSK', Sarabun, sans-serif;
        }}
        body {{
            margin: 0;
            padding: 20px;
            background: #e2e8f0;
            color: #000;
            font-size: 16pt;
            line-height: 1.15;
            -webkit-print-color-adjust: exact;
            print-color-adjust: exact;
            font-feature-settings: "locl" 1, "numr" 1;
            font-variant-numeric: tabular-nums;
        }}

        .sheet-container {{
            width: 297mm;
            min-height: 210mm;
            margin: 0 auto 30px auto;
            background: #fff;
            padding: 12mm 15mm;
            box-shadow: 0 4px 12px rgba(0,0,0,0.15);
            position: relative;
            page-break-after: always;
            display: flex;
            flex-direction: column;
        }}

        /* Typography */
        .text-center {{ text-align: center; }}
        .text-left {{ text-align: left; }}
        .text-right {{ text-align: right; }}
        .font-bold {{ font-weight: bold; }}

        /* PAGE 1: Cover Page */
        .cover-page {{
            display: flex;
            flex-direction: column;
            align-items: center;
            justify-content: center;
            height: 100%;
            min-height: 185mm;
            text-align: center;
        }}
        .cover-emblem {{
            width: 140px;
            height: auto;
            margin-bottom: 24px;
        }}
        .cover-title {{
            font-size: 32pt;
            font-weight: bold;
            margin: 0 0 12px 0;
        }}
        .cover-unit {{
            font-size: 22pt;
            margin: 0 0 16px 0;
        }}
        .cover-divider {{
            letter-spacing: 2px;
            margin: 10px 0 24px 0;
            font-size: 18pt;
        }}
        .cover-details {{
            font-size: 20pt;
            line-height: 1.8;
            margin-top: 10px;
        }}

        /* PAGE 2: Weapon History (ป.6) */
        .page-corner-code {{
            position: absolute;
            top: 15mm;
            right: 20mm;
            font-size: 20pt;
            font-weight: bold;
        }}
        .page-title {{
            text-align: center;
            font-size: 24pt;
            font-weight: bold;
            margin-top: 25px;
            margin-bottom: 6px;
        }}
        .p6-container {{
            max-width: 680px;
            margin: 40px auto 0 auto;
            font-size: 18pt;
            line-height: 2.3;
        }}
        .p6-row {{
            display: flex;
            align-items: baseline;
            margin-bottom: 12px;
        }}
        .p6-label {{
            width: 200px;
            flex-shrink: 0;
        }}
        .p6-value {{
            flex-grow: 1;
            border-bottom: 1px dotted #555;
            padding-left: 15px;
        }}

        /* PAGE 3: Barrel Life Calculation */
        .page3-layout {{
            display: flex;
            gap: 30px;
            margin-top: 25px;
            align-items: flex-start;
        }}
        .page3-left {{
            width: 38%;
            padding: 15px 15px;
            font-size: 15pt;
            line-height: 1.4;
        }}
        .page3-left-row {{
            margin-bottom: 22px;
        }}
        .page3-label {{
            font-size: 15pt;
            color: #000;
            margin-bottom: 6px;
            display: block;
        }}
        .page3-field {{
            border-bottom: 1px dotted #444;
            display: block;
            width: 100%;
            height: 28px;
            font-size: 15pt;
            padding-left: 10px;
        }}
        .page3-right {{
            width: 62%;
        }}

        /* TABLES */
        .standard-table {{
            width: 100%;
            table-layout: fixed;
            border-collapse: collapse;
            border: 1px solid #000;
            font-size: 15pt;
        }}
        .standard-table th, .standard-table td {{
            border: 1px solid #000;
            padding: 4px 6px;
            vertical-align: middle;
            line-height: 1.15;
            white-space: nowrap;
        }}
        .standard-table th {{
            text-align: center;
            font-weight: normal;
            background: #fafafa;
        }}
        .standard-table td {{
            height: 27px;
        }}

        /* PAGE 4 & 5 Headings */
        .operation-header {{
            margin-bottom: 16px;
        }}
        .operation-subhead {{
            display: flex;
            justify-content: space-between;
            font-size: 16pt;
            margin-top: 10px;
            margin-bottom: 12px;
            padding: 0 5px;
        }}

        @media print {{
            body {{ background: transparent; padding: 0; }}
            .sheet-container {{
                box-shadow: none;
                margin: 0;
                padding: 10mm 15mm;
                width: 100%;
                min-height: auto;
            }}
        }}
    </style>
</head>
<body>

    <!-- ========================================== -->
    <!-- PAGE 1: หน้าปก (COVER)                     -->
    <!-- ========================================== -->
    <div class="sheet-container">
        <div class="cover-page">
            <img src="artillery_emblem.png" class="cover-emblem" alt="ตรากองพันทหารปืนใหญ่ที่ ๔">
            <div class="cover-title">สมุดประวัติการใช้ปืนใหญ่</div>
            <div class="cover-unit">หน่วย ป.๔ พัน.๔ ร้อย.๒</div>
            <div class="cover-divider">…………………………………</div>
            <div class="cover-details">
                <div>ปบค.๙๕ ขนาด ๑๐๕ มม.</div>
                <div>แบบ M.101 A.1 ปรับปรุง</div>
                <div style="margin-top: 8px;">หมายเลข &nbsp;&nbsp;&nbsp;&nbsp; ๘๕๑๖ &nbsp;&nbsp;&nbsp;&nbsp;</div>
                <div style="margin-top: 8px;">สถานภาพ &nbsp;&nbsp;&nbsp;&nbsp; พร้อมใช้งาน &nbsp;&nbsp;&nbsp;&nbsp;</div>
            </div>
        </div>
    </div>

    <!-- ========================================== -->
    <!-- PAGE 2: ประวัติอาวุธ (ป.๖)                 -->
    <!-- ========================================== -->
    <div class="sheet-container">
        <div class="page-corner-code">ป.๖</div>
        <div class="page-title">ประวัติอาวุธ</div>
        <div class="text-center" style="letter-spacing: 2px; margin-bottom: 30px;">................................................</div>
        <div class="p6-container">
            <div class="p6-row">
                <div class="p6-label">๑. ชนิดอาวุธ</div>
                <div class="p6-value">ปบค.๙๕ ขนาด ๑๐๕ มม.</div>
            </div>
            <div class="p6-row">
                <div class="p6-label">๒. แบบ</div>
                <div class="p6-value">M.101 A.1 ปรับปรุง</div>
            </div>
            <div class="p6-row">
                <div class="p6-label">๓. หมายเลข</div>
                <div class="p6-value">๘๕๑๖</div>
            </div>
            <div class="p6-row">
                <div class="p6-label">๔. ค่า ตร.</div>
                <div class="p6-value">-๙.๑๒</div>
            </div>
            <div class="p6-row">
                <div class="p6-label">๕. บริษัทผู้สร้าง</div>
                <div class="p6-value">....................................................................</div>
            </div>
        </div>
    </div>

    <!-- ========================================== -->
    <!-- PAGE 3: การคำนวณอายุลำกล้อง               -->
    <!-- ========================================== -->
    <div class="sheet-container">
        <div class="page-title">การคำนวณอายุลำกล้อง</div>
        <div class="page3-layout">
            <div class="page3-left">
                <div class="page3-left-row">
                    <span class="page3-label">ชนิดลำกล้อง</span>
                    <span class="page3-field">M2A2</span>
                </div>
                <div class="page3-left-row">
                    <span class="page3-label">หมายเลขลำกล้อง</span>
                    <span class="page3-field">๑๔๐๓๖</span>
                </div>
                <div class="page3-left-row">
                    <span class="page3-label">อายุลำกล้องสูงสุด</span>
                    <span class="page3-field">๗,๕๐๐ นัด</span>
                </div>
            </div>
            <div class="page3-right">
                <table class="standard-table">
                    <colgroup>
                        <col style="width: 30%;">
                        <col style="width: 24%;">
                        <col style="width: 20%;">
                        <col style="width: 26%;">
                    </colgroup>
                    <thead>
                        <tr>
                            <th colspan="3" style="padding: 6px;">อัตราการคำนวณกระสุน</th>
                            <th rowspan="2" style="vertical-align: middle;">หมายเหตุ</th>
                        </tr>
                        <tr>
                            <th>ดินส่งกระสุน</th>
                            <th>ส่วนบรรจุ</th>
                            <th>ค่า EFC</th>
                        </tr>
                    </thead>
                    <tbody>
                        {"".join(["<tr><td></td><td></td><td></td><td></td></tr>" for _ in range(12)])}
                    </tbody>
                </table>
            </div>
        </div>
    </div>

    <!-- ========================================== -->
    <!-- PAGE 4: ประวัติการใช้ปืนใหญ่ (บันทึกซ่อม)  -->
    <!-- ========================================== -->
    <div class="sheet-container">
        <div class="operation-header">
            <div class="page-title">ประวัติการใช้ปืนใหญ่</div>
            <div class="operation-subhead">
                <div>ปบค.๙๕ ขนาด ๑๐๕ มม. แบบ M.101 A.1 ปรับปรุง หมายเลข ๘๕๑๖</div>
                <div>หน่วย ป.๔ พัน.๔ ร้อย.๒</div>
            </div>
        </div>
        <table class="standard-table">
            <colgroup>
                <col style="width: 12%;">
                <col style="width: 34%;">
                <col style="width: 20%;">
                <col style="width: 11%;">
                <col style="width: 11%;">
                <col style="width: 12%;">
            </colgroup>
            <thead>
                <tr>
                    <th rowspan="2" style="vertical-align: middle;">วัน/เดือน/ปี</th>
                    <th rowspan="2" style="vertical-align: middle;">ภารกิจ/การปฏิบัติ</th>
                    <th rowspan="2" style="vertical-align: middle;">ปัญหาข้อขัดข้อง</th>
                    <th colspan="2" style="padding: 4px;">การดำเนินการ</th>
                    <th rowspan="2" style="vertical-align: middle;">หมายเหตุ</th>
                </tr>
                <tr>
                    <th>แก้ไข</th>
                    <th>ส่งซ่อม</th>
                </tr>
            </thead>
            <tbody>
                {"".join(["<tr><td></td><td></td><td></td><td></td><td></td><td></td></tr>" for _ in range(13)])}
            </tbody>
        </table>
    </div>

    <!-- ========================================== -->
    <!-- PAGE 5: บันทึกหลักฐานการยิง                -->
    <!-- ========================================== -->
    <div class="sheet-container">
        <div class="page-title" style="margin-top: 10px; margin-bottom: 16px;">บันทึกหลักฐานการยิง</div>
        <table class="standard-table">
            <colgroup>
                <col style="width: 8%;">
                <col style="width: 12%;">
                <col style="width: 10%;">
                <col style="width: 7%;">
                <col style="width: 6%;">
                <col style="width: 11%;">
                <col style="width: 9%;">
                <col style="width: 17%;">
                <col style="width: 20%;">
            </colgroup>
            <thead>
                <tr>
                    <th rowspan="2" style="vertical-align: middle;">วัน/<br>เดือน/ปี</th>
                    <th rowspan="2" style="vertical-align: middle;">ชนิดกระสุน</th>
                    <th rowspan="2" style="vertical-align: middle;">ส่วนบรรจุ</th>
                    <th rowspan="2" style="vertical-align: middle;">จำนวน<br>นัด</th>
                    <th rowspan="2" style="vertical-align: middle;">ค่า<br>EFC</th>
                    <th>ค่า EFC</th>
                    <th rowspan="2" style="vertical-align: middle;">อายุลำ<br>กล้อง</th>
                    <th rowspan="2" style="vertical-align: middle;">ลายเซ็น</th>
                    <th rowspan="2" style="vertical-align: middle;">หมายเหตุ</th>
                </tr>
                <tr>
                    <th>X จำนวน<br>นัด</th>
                </tr>
            </thead>
            <tbody>
                {"".join(["<tr><td></td><td></td><td></td><td></td><td></td><td></td><td></td><td></td><td></td></tr>" for _ in range(13)])}
            </tbody>
        </table>
    </div>

</body>
</html>
    """

    html_path = os.path.join(WORKDIR, "สมุดประวัติการใช้ปืนใหญ่_Rebuild_Thai.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print("Saved HTML:", html_path)
    return html_path

generate_html()

# ----------------------------------------------------------------------
# 2. Microsoft Excel Generation (.xlsx)
# ----------------------------------------------------------------------
def generate_excel():
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # remove default sheet

    font_title = Font(name="TH Sarabun PSK", size=20, bold=True)
    font_bold = Font(name="TH Sarabun PSK", size=16, bold=True)
    font_regular = Font(name="TH Sarabun PSK", size=16)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    fill_header = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    thin = Side(border_style="thin", color="000000")
    box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Sheet 1: หน้าปก
    ws1 = wb.create_sheet(title="หน้าปก")
    ws1.views.sheetView[0].showGridLines = True
    ws1.column_dimensions['A'].width = 15
    ws1.column_dimensions['B'].width = 45
    ws1.column_dimensions['C'].width = 15

    ws1["B4"] = "สมุดประวัติการใช้ปืนใหญ่"
    ws1["B4"].font = Font(name="TH Sarabun PSK", size=26, bold=True)
    ws1["B4"].alignment = align_center

    ws1["B5"] = "หน่วย ป.๔ พัน.๔ ร้อย.๒"
    ws1["B5"].font = Font(name="TH Sarabun PSK", size=18, bold=True)
    ws1["B5"].alignment = align_center

    ws1["B6"] = "…………………………………"
    ws1["B6"].font = font_regular
    ws1["B6"].alignment = align_center

    ws1["B8"] = "ปบค.๙๕ ขนาด ๑๐๕ มม."
    ws1["B8"].font = Font(name="TH Sarabun PSK", size=18)
    ws1["B8"].alignment = align_center

    ws1["B9"] = "แบบ M.101 A.1 ปรับปรุง"
    ws1["B9"].font = Font(name="TH Sarabun PSK", size=18)
    ws1["B9"].alignment = align_center

    ws1["B11"] = "หมายเลข     ๘๕๑๖"
    ws1["B11"].font = Font(name="TH Sarabun PSK", size=18)
    ws1["B11"].alignment = align_center

    ws1["B12"] = "สถานภาพ     พร้อมใช้งาน"
    ws1["B12"].font = Font(name="TH Sarabun PSK", size=18)
    ws1["B12"].alignment = align_center

    # Sheet 2: ประวัติอาวุธ ป.๖
    ws2 = wb.create_sheet(title="ประวัติอาวุธ (ป.๖)")
    ws2.views.sheetView[0].showGridLines = True
    ws2.column_dimensions['A'].width = 5
    ws2.column_dimensions['B'].width = 25
    ws2.column_dimensions['C'].width = 40
    ws2.column_dimensions['D'].width = 15

    ws2["D2"] = "ป.๖"
    ws2["D2"].font = font_bold
    ws2["D2"].alignment = align_right

    ws2["B3"] = "ประวัติอาวุธ"
    ws2["B3"].font = font_title
    ws2["B3"].alignment = align_center
    ws2.merge_cells("B3:C3")

    p6_items = [
        ("๑. ชนิดอาวุธ", "ปบค.๙๕ ขนาด ๑๐๕ มม."),
        ("๒. แบบ", "M.101 A.1 ปรับปรุง"),
        ("๓. หมายเลข", "๘๕๑๖"),
        ("๔. ค่า ตร.", "-๙.๑๒"),
        ("๕. บริษัทผู้สร้าง", "....................................................")
    ]
    for r_idx, (lab, val) in enumerate(p6_items, start=5):
        ws2[f"B{r_idx}"] = lab
        ws2[f"B{r_idx}"].font = font_regular
        ws2[f"C{r_idx}"] = val
        ws2[f"C{r_idx}"].font = font_regular

    # Sheet 3: การคำนวณอายุลำกล้อง
    ws3 = wb.create_sheet(title="คำนวณอายุลำกล้อง")
    ws3.views.sheetView[0].showGridLines = True
    ws3.column_dimensions['A'].width = 18
    ws3.column_dimensions['B'].width = 16
    ws3.column_dimensions['C'].width = 5
    ws3.column_dimensions['D'].width = 22
    ws3.column_dimensions['E'].width = 16
    ws3.column_dimensions['F'].width = 14
    ws3.column_dimensions['G'].width = 18

    ws3["A2"] = "การคำนวณอายุลำกล้อง"
    ws3["A2"].font = font_title
    ws3["A2"].alignment = align_center
    ws3.merge_cells("A2:G2")

    font_label_p3 = Font(name="TH Sarabun PSK", size=15)
    font_val_p3 = Font(name="TH Sarabun PSK", size=15, bold=True)

    # Separate rows for each label and value
    ws3["A4"] = "ชนิดลำกล้อง"
    ws3["A4"].font = font_label_p3
    ws3["A5"] = "M2A2"
    ws3["A5"].font = font_val_p3

    ws3["A7"] = "หมายเลขลำกล้อง"
    ws3["A7"].font = font_label_p3
    ws3["A8"] = "๑๔๐๓๖"
    ws3["A8"].font = font_val_p3

    ws3["A10"] = "อายุลำกล้องสูงสุด"
    ws3["A10"].font = font_label_p3
    ws3["A11"] = "๗,๕๐๐ นัด"
    ws3["A11"].font = font_val_p3

    # Right Table
    ws3["D4"] = "อัตราการคำนวณกระสุน"
    ws3["D4"].font = font_regular
    ws3["D4"].alignment = align_center
    ws3["D4"].fill = fill_header
    ws3.merge_cells("D4:F4")

    ws3["G4"] = "หมายเหตุ"
    ws3["G4"].font = font_regular
    ws3["G4"].alignment = align_center
    ws3["G4"].fill = fill_header
    ws3.merge_cells("G4:G5")

    ws3["D5"] = "ดินส่งกระสุน"
    ws3["E5"] = "ส่วนบรรจุ"
    ws3["F5"] = "ค่า EFC"
    for col in ["D", "E", "F"]:
        ws3[f"{col}5"].font = font_regular
        ws3[f"{col}5"].alignment = align_center
        ws3[f"{col}5"].fill = fill_header

    for r in range(4, 18):
        for c in ["D", "E", "F", "G"]:
            ws3[f"{c}{r}"].border = box_border

    # Sheet 4: ประวัติการใช้ปืนใหญ่
    ws4 = wb.create_sheet(title="ประวัติการใช้ปืนใหญ่")
    ws4.views.sheetView[0].showGridLines = True
    cols_w4 = {'A': 16, 'B': 36, 'C': 25, 'D': 14, 'E': 14, 'F': 16}
    for c, w in cols_w4.items():
        ws4.column_dimensions[c].width = w

    ws4["A2"] = "ประวัติการใช้ปืนใหญ่"
    ws4["A2"].font = font_title
    ws4["A2"].alignment = align_center
    ws4.merge_cells("A2:F2")

    ws4["A3"] = "ปบค.๙๕ ขนาด ๑๐๕ มม. แบบ M.101 A.1 ปรับปรุง หมายเลข ๘๕๑๖"
    ws4["A3"].font = font_regular
    ws4.merge_cells("A3:C3")
    ws4["D3"] = "หน่วย ป.๔ พัน.๔ ร้อย.๒"
    ws4["D3"].font = font_regular
    ws4["D3"].alignment = align_right
    ws4.merge_cells("D3:F3")

    ws4["A4"] = "วัน/เดือน/ปี"; ws4.merge_cells("A4:A5")
    ws4["B4"] = "ภารกิจ/การปฏิบัติ"; ws4.merge_cells("B4:B5")
    ws4["C4"] = "ปัญหาข้อขัดข้อง"; ws4.merge_cells("C4:C5")
    ws4["D4"] = "การดำเนินการ"; ws4.merge_cells("D4:E4")
    ws4["D5"] = "แก้ไข"
    ws4["E5"] = "ส่งซ่อม"
    ws4["F4"] = "หมายเหตุ"; ws4.merge_cells("F4:F5")

    for r in range(4, 6):
        for c in ['A', 'B', 'C', 'D', 'E', 'F']:
            ws4[f"{c}{r}"].font = font_regular
            ws4[f"{c}{r}"].alignment = align_center
            ws4[f"{c}{r}"].fill = fill_header

    for r in range(4, 19):
        for c in ['A', 'B', 'C', 'D', 'E', 'F']:
            ws4[f"{c}{r}"].border = box_border

    # Sheet 5: บันทึกหลักฐานการยิง
    ws5 = wb.create_sheet(title="บันทึกหลักฐานการยิง")
    ws5.views.sheetView[0].showGridLines = True
    cols_w5 = {'A': 14, 'B': 18, 'C': 14, 'D': 12, 'E': 10, 'F': 16, 'G': 14, 'H': 20, 'I': 22}
    for c, w in cols_w5.items():
        ws5.column_dimensions[c].width = w

    ws5["A2"] = "บันทึกหลักฐานการยิง"
    ws5["A2"].font = font_title
    ws5["A2"].alignment = align_center
    ws5.merge_cells("A2:I2")

    ws5["A4"] = "วัน/\nเดือน/ปี"; ws5.merge_cells("A4:A5")
    ws5["B4"] = "ชนิดกระสุน"; ws5.merge_cells("B4:B5")
    ws5["C4"] = "ส่วนบรรจุ"; ws5.merge_cells("C4:C5")
    ws5["D4"] = "จำนวน\nนัด"; ws5.merge_cells("D4:D5")
    ws5["E4"] = "ค่า\nEFC"; ws5.merge_cells("E4:E5")
    ws5["F4"] = "ค่า EFC\nX จำนวนนัด"; ws5.merge_cells("F4:F5")
    ws5["G4"] = "อายุลำ\nกล้อง"; ws5.merge_cells("G4:G5")
    ws5["H4"] = "ลายเซ็น"; ws5.merge_cells("H4:H5")
    ws5["I4"] = "หมายเหตุ"; ws5.merge_cells("I4:I5")

    for r in range(4, 6):
        for c in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
            ws5[f"{c}{r}"].font = font_regular
            ws5[f"{c}{r}"].alignment = align_center
            ws5[f"{c}{r}"].fill = fill_header

    for r in range(4, 19):
        for c in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
            ws5[f"{c}{r}"].border = box_border

    xlsx_path = os.path.join(WORKDIR, "สมุดประวัติการใช้ปืนใหญ่_Rebuild_Thai.xlsx")
    wb.save(xlsx_path)
    print("Saved XLSX:", xlsx_path)

generate_excel()

# ----------------------------------------------------------------------
# 3. Microsoft Word Generation (.docx)
# ----------------------------------------------------------------------
def format_docx_run(run, font_name="TH Sarabun PSK", size_pt=16, bold=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}" w:eastAsia="{font_name}"/>')
    rPr.append(rFonts)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top="single", bottom="single", left="single", right="single", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="000000"/><w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="000000"/><w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="000000"/><w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="000000"/></w:tcBorders>')
    tcPr.append(borders)

def generate_word():
    doc = docx.Document()

    # Landscape A4 (11.69" x 8.27")
    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)

    # PAGE 1: Cover
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(40)
    p_img.paragraph_format.space_after = Pt(14)
    run_img = p_img.add_run()
    run_img.add_picture(os.path.join(WORKDIR, "artillery_emblem.png"), width=Inches(1.5))

    p_c1 = doc.add_paragraph()
    p_c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c1.paragraph_format.space_after = Pt(8)
    r = p_c1.add_run("สมุดประวัติการใช้ปืนใหญ่")
    format_docx_run(r, size_pt=26, bold=True)

    p_c2 = doc.add_paragraph()
    p_c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c2.paragraph_format.space_after = Pt(12)
    r = p_c2.add_run("หน่วย ป.๔ พัน.๔ ร้อย.๒")
    format_docx_run(r, size_pt=20, bold=True)

    p_c3 = doc.add_paragraph()
    p_c3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c3.paragraph_format.space_after = Pt(16)
    r = p_c3.add_run("…………………………………")
    format_docx_run(r, size_pt=18)

    p_c4 = doc.add_paragraph()
    p_c4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c4.paragraph_format.space_after = Pt(6)
    r = p_c4.add_run("ปบค.๙๕ ขนาด ๑๐๕ มม.\nแบบ M.101 A.1 ปรับปรุง\nหมายเลข     ๘๕๑๖\nสถานภาพ     พร้อมใช้งาน")
    format_docx_run(r, size_pt=18)

    # PAGE 2: ประวัติอาวุธ (ป.๖)
    doc.add_page_break()

    p_code = doc.add_paragraph()
    p_code.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_code.paragraph_format.space_before = Pt(10)
    p_code.paragraph_format.space_after = Pt(20)
    r = p_code.add_run("ป.๖")
    format_docx_run(r, size_pt=18, bold=True)

    p_t2 = doc.add_paragraph()
    p_t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t2.paragraph_format.space_after = Pt(8)
    r = p_t2.add_run("ประวัติอาวุธ")
    format_docx_run(r, size_pt=22, bold=True)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(30)
    r = p_div.add_run("................................................")
    format_docx_run(r, size_pt=18)

    p6_texts = [
        "๑. ชนิดอาวุธ                     ปบค.๙๕ ขนาด ๑๐๕ มม.",
        "๒. แบบ                             M.101 A.1 ปรับปรุง",
        "๓. หมายเลข                      ๘๕๑๖",
        "๔. ค่า ตร.                         -๙.๑๒",
        "๕. บริษัทผู้สร้าง                 ...................................................."
    ]
    for pt in p6_texts:
        p_row = doc.add_paragraph()
        p_row.paragraph_format.left_indent = Inches(2.5)
        p_row.paragraph_format.space_after = Pt(16)
        r = p_row.add_run(pt)
        format_docx_run(r, size_pt=18)

    # PAGE 3: การคำนวณอายุลำกล้อง
    doc.add_page_break()

    p_t3 = doc.add_paragraph()
    p_t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t3.paragraph_format.space_before = Pt(10)
    p_t3.paragraph_format.space_after = Pt(20)
    r = p_t3.add_run("การคำนวณอายุลำกล้อง")
    format_docx_run(r, size_pt=22, bold=True)

    tbl3 = doc.add_table(rows=14, cols=6)
    tbl3.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl3.autofit = False

    # Widths: left info (Col 0-1), right table (Col 2-5)
    w_p3 = [Inches(2.0), Inches(2.0), Inches(2.0), Inches(1.5), Inches(1.3), Inches(1.69)]
    for row in tbl3.rows:
        for idx, w in enumerate(w_p3):
            row.cells[idx].width = w

    # Left info cells: label and value on separate lines
    tbl3.cell(0, 0).text = "ชนิดลำกล้อง"
    tbl3.cell(1, 0).text = "M2A2"
    tbl3.cell(3, 0).text = "หมายเลขลำกล้อง"
    tbl3.cell(4, 0).text = "๑๔๐๓๖"
    tbl3.cell(6, 0).text = "อายุลำกล้องสูงสุด"
    tbl3.cell(7, 0).text = "๗,๕๐๐ นัด"

    # Merge Col 0 and Col 1 for left info rows
    for r in [0, 1, 3, 4, 6, 7]:
        tbl3.cell(r, 0).merge(tbl3.cell(r, 1))
        p_c = tbl3.cell(r, 0).paragraphs[0]
        if p_c.runs:
            is_val = (r in [1, 4, 7])
            format_docx_run(p_c.runs[0], size_pt=15, bold=is_val)

    # Right table header: merge row 0 col 2-4 for อัตราการคำนวณกระสุน
    c_hdr = tbl3.cell(0, 2)
    c_hdr.merge(tbl3.cell(0, 4))
    c_hdr.text = "อัตราการคำนวณกระสุน"
    c_hdr.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_docx_run(c_hdr.paragraphs[0].runs[0], size_pt=16)

    c_rem = tbl3.cell(0, 5)
    c_rem.merge(tbl3.cell(1, 5))
    c_rem.text = "หมายเหตุ"
    c_rem.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_docx_run(c_rem.paragraphs[0].runs[0], size_pt=16)

    sub_hdrs = ["ดินส่งกระสุน", "ส่วนบรรจุ", "ค่า EFC"]
    for i, sh in enumerate(sub_hdrs):
        cell = tbl3.cell(1, 2 + i)
        cell.text = sh
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_docx_run(cell.paragraphs[0].runs[0], size_pt=16)

    # Set borders for right table
    for r in range(0, 14):
        for c in range(2, 6):
            cell = tbl3.cell(r, c)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            set_cell_borders(cell, top="single", bottom="single", left="single", right="single")

    # PAGE 4: ประวัติการใช้ปืนใหญ่
    doc.add_page_break()

    p_t4 = doc.add_paragraph()
    p_t4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t4.paragraph_format.space_before = Pt(10)
    p_t4.paragraph_format.space_after = Pt(8)
    r = p_t4.add_run("ประวัติการใช้ปืนใหญ่")
    format_docx_run(r, size_pt=22, bold=True)

    p_sub4 = doc.add_paragraph()
    p_sub4.paragraph_format.space_after = Pt(14)
    r1 = p_sub4.add_run("ปบค.๙๕ ขนาด ๑๐๕ มม. แบบ M.101 A.1 ปรับปรุง หมายเลข ๘๕๑๖")
    format_docx_run(r1, size_pt=16)
    r2 = p_sub4.add_run("\t\t\t\t\t\t\t\tหน่วย ป.๔ พัน.๔ ร้อย.๒")
    format_docx_run(r2, size_pt=16)

    tbl4 = doc.add_table(rows=15, cols=6)
    tbl4.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl4.autofit = False
    w_p4 = [Inches(1.26), Inches(3.57), Inches(2.1), Inches(1.15), Inches(1.15), Inches(1.26)]
    for row in tbl4.rows:
        for idx, w in enumerate(w_p4):
            row.cells[idx].width = w

    # Merge headers
    tbl4.cell(0, 0).merge(tbl4.cell(1, 0)).text = "วัน/เดือน/ปี"
    tbl4.cell(0, 1).merge(tbl4.cell(1, 1)).text = "ภารกิจ/การปฏิบัติ"
    tbl4.cell(0, 2).merge(tbl4.cell(1, 2)).text = "ปัญหาข้อขัดข้อง"
    tbl4.cell(0, 3).merge(tbl4.cell(0, 4)).text = "การดำเนินการ"
    tbl4.cell(1, 3).text = "แก้ไข"
    tbl4.cell(1, 4).text = "ส่งซ่อม"
    tbl4.cell(0, 5).merge(tbl4.cell(1, 5)).text = "หมายเหตุ"

    for r in range(15):
        for c in range(6):
            cell = tbl4.cell(r, c)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            set_cell_borders(cell, top="single", bottom="single", left="single", right="single")
            if r < 2 and cell.paragraphs[0].runs:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                format_docx_run(cell.paragraphs[0].runs[0], size_pt=16)

    # PAGE 5: บันทึกหลักฐานการยิง
    doc.add_page_break()

    p_t5 = doc.add_paragraph()
    p_t5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t5.paragraph_format.space_before = Pt(10)
    p_t5.paragraph_format.space_after = Pt(14)
    r = p_t5.add_run("บันทึกหลักฐานการยิง")
    format_docx_run(r, size_pt=22, bold=True)

    tbl5 = doc.add_table(rows=15, cols=9)
    tbl5.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl5.autofit = False
    w_p5 = [Inches(0.84), Inches(1.26), Inches(1.05), Inches(0.74), Inches(0.63), Inches(1.15), Inches(0.95), Inches(1.78), Inches(2.1)]
    for row in tbl5.rows:
        for idx, w in enumerate(w_p5):
            row.cells[idx].width = w

    tbl5.cell(0, 0).merge(tbl5.cell(1, 0)).text = "วัน/\nเดือน/ปี"
    tbl5.cell(0, 1).merge(tbl5.cell(1, 1)).text = "ชนิดกระสุน"
    tbl5.cell(0, 2).merge(tbl5.cell(1, 2)).text = "ส่วนบรรจุ"
    tbl5.cell(0, 3).merge(tbl5.cell(1, 3)).text = "จำนวน\nนัด"
    tbl5.cell(0, 4).merge(tbl5.cell(1, 4)).text = "ค่า\nEFC"
    tbl5.cell(0, 5).text = "ค่า EFC"
    tbl5.cell(1, 5).text = "X จำนวนนัด"
    tbl5.cell(0, 6).merge(tbl5.cell(1, 6)).text = "อายุลำ\nกล้อง"
    tbl5.cell(0, 7).merge(tbl5.cell(1, 7)).text = "ลายเซ็น"
    tbl5.cell(0, 8).merge(tbl5.cell(1, 8)).text = "หมายเหตุ"

    for r in range(15):
        for c in range(9):
            cell = tbl5.cell(r, c)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            set_cell_borders(cell, top="single", bottom="single", left="single", right="single")
            if r < 2 and cell.paragraphs[0].runs:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                format_docx_run(cell.paragraphs[0].runs[0], size_pt=15)

    docx_path = os.path.join(WORKDIR, "สมุดประวัติการใช้ปืนใหญ่_Rebuild_Thai.docx")
    doc.save(docx_path)
    print("Saved DOCX:", docx_path)

generate_word()
print("All 3 Formats Generated Successfully!")
