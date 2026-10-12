#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
================================================================================
SUPERCLONE PRODUCTION-GRADE ENGINE
Form: ทบ. ๔๖๘ – ๓๖๐ (สพ. ๔๖๐) กำหนดการปรนนิบัติบำรุงยุทธภัณฑ์
Target: ป.๔ พัน.๔ ร้อย.๒ (ปบค.๙๕ ขนาด ๑๐๕ มม. เอ็ม ๑๐๑ เอ ๑)
Compiler: Multi-Format Generator (HTML, DOCX, XLSX)
================================================================================
"""

import os
import sys
import json

# Ensure stdout handles UTF-8 on Windows
if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load SSOT Blueprint
BLUEPRINT_PATH = os.path.join(os.path.dirname(__file__), "table_layout_468_360.json")
with open(BLUEPRINT_PATH, "r", encoding="utf-8") as f:
    bp = json.load(f)

# -----------------------------------------------------------------------------
# 1. GENERATE PRODUCTION HTML
# -----------------------------------------------------------------------------
def generate_html(output_path):
    p1 = bp["page_1_cover"]
    p2 = bp["page_2_equipment_table"]
    p3 = bp["page_3_calendar_matrix"]
    p4 = bp["page_4_instructions"]

    # Table 2 rows HTML (ตัวบาง / normal weight)
    rows_p2_html = ""
    for r in p2["rows"]:
        rows_p2_html += f"""
        <tr>
            <td class="text-center">{r["seq"]}</td>
            <td>{r["driver"]}</td>
            <td>{r["supervisor"]}</td>
            <td>{r["equipment_name"]}</td>
            <td>{r["last_pm_date_mile"]}</td>
            <td>{r["unit_internal_no"]}</td>
            <td>{r["next_pm_date"]}</td>
            <td class="text-center">{r["registration_no"]}</td>
        </tr>
        """

    # Table 2 columns header HTML (บรรทัดเดียว & ตัวบาง)
    headers_p2_html = ""
    for col in p2["columns"]:
        headers_p2_html += f'<th style="width: {col["width"]};">{col["title"]}</th>\n'

    # Table 3 (Calendar) columns and rows HTML (บรรทัดเดียว & ตัวบาง)
    days_th_html = '<th class="cal-seq-col"></th>'
    for d in range(1, 32):
        thai_d = str(d).translate(str.maketrans("0123456789", "๐๑๒๓๔๕๖๗๘๙"))
        days_th_html += f'<th class="cal-day-col"><span class="d-single">{thai_d}</span></th>'
    days_th_html += '<th class="cal-seq-col"></th>'

    # Rows 1 to 6 (ตัวบาง)
    rows_p3_html = ""
    for r in p3["rows"]:
        thai_seq = r["seq"]
        cells = f'<td class="text-center cal-seq-cell">{thai_seq}</td>'
        for d in range(1, 32):
            cells += '<td class="cal-cell"></td>'
        cells += f'<td class="text-center cal-seq-cell">{thai_seq}</td>'
        rows_p3_html += f'<tr>{cells}</tr>\n'

    # Instructions HTML
    inst_html = ""
    for rule in p4["rules"]:
        inst_html += f"""
        <div class="inst-rule">
            <p class="inst-p">
                <span class="font-bold">{rule["index"]}. {rule["title"]} – </span>
                <span>{rule["content"]}</span>
            </p>
        """
        if "abbreviations" in rule:
            inst_html += '<div class="abbr-list">\n'
            for abb in rule["abbreviations"]:
                inst_html += f"""
                <div class="abbr-item">
                    <span class="abbr-code font-bold">{abb["code"]}</span>
                    <span class="abbr-desc">{abb["description"]}</span>
                </div>
                """
            inst_html += '</div>\n'
        inst_html += '</div>\n'

    html_content = f"""<!DOCTYPE html>
<html lang="th">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{bp["document_title"]}</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Sarabun:ital,wght@0,300;0,400;0,600;0,700;1,400&display=swap" rel="stylesheet">
    <style>
        @page {{
            size: A4 landscape;
            margin: 0;
        }}
        * {{
            box-sizing: border-box;
            -webkit-print-color-adjust: exact;
            print-color-adjust: exact;
        }}
        body {{
            margin: 0;
            padding: 20px 0;
            background-color: #525659;
            font-family: 'TH Sarabun PSK', 'THSarabunNew', 'Sarabun', sans-serif;
            font-feature-settings: "locl" 1, "numr" 1;
            font-variant-numeric: tabular-nums;
            color: #000000;
            line-height: 1.25;
            display: flex;
            flex-direction: column;
            align-items: center;
        }}
        .page {{
            width: 297mm;
            height: 210mm;
            background-color: #ffffff;
            margin-bottom: 20mm;
            box-shadow: 0 4px 15px rgba(0, 0, 0, 0.35);
            position: relative;
            page-break-after: always;
            padding: 12mm 16mm 10mm 16mm;
            overflow: hidden;
        }}
        @media print {{
            body {{
                background-color: #ffffff;
                padding: 0;
            }}
            .page {{
                margin: 0;
                box-shadow: none;
                page-break-after: always;
            }}
        }}

        /* Universal Typography */
        .text-center {{ text-align: center; }}
        .text-left {{ text-align: left; }}
        .text-right {{ text-align: right; }}
        .font-bold {{ font-weight: 700; }}
        
        /* PAGE 1: COVER */
        .p1-header {{
            text-align: right;
            font-size: 20pt;
            font-weight: 700;
            margin-top: 5mm;
            margin-right: 5mm;
        }}
        .p1-title-container {{
            text-align: center;
            margin-top: 16mm;
            margin-bottom: 12mm;
        }}
        .p1-title {{
            font-size: 38pt;
            font-weight: 700;
            letter-spacing: 0.5px;
        }}
        .p1-slogans {{
            text-align: center;
            font-size: 20pt;
            font-weight: 700;
            line-height: 1.6;
            margin-bottom: 10mm;
        }}
        .p1-box-wrapper {{
            display: flex;
            justify-content: center;
            margin-bottom: 15mm;
        }}
        .p1-box-table {{
            border-collapse: collapse;
            border: 1.5px solid #000;
            width: 95mm;
            text-align: center;
        }}
        .p1-box-table td {{
            padding: 4mm 6mm;
            border: 1px solid #000;
        }}
        .p1-box-unit {{
            font-size: 19pt;
            font-weight: 700;
        }}
        .p1-box-month {{
            font-size: 18pt;
            font-weight: 700;
        }}
        .p1-footer {{
            position: absolute;
            bottom: 12mm;
            left: 16mm;
            right: 16mm;
            display: flex;
            justify-content: space-between;
            align-items: flex-end;
        }}
        .p1-sub-id {{
            font-size: 20pt;
            font-weight: 700;
        }}
        .p1-authority {{
            font-size: 15pt;
            font-weight: 700;
        }}

        /* PAGE 2: EQUIPMENT TABLE */
        .page-2-container {{
            padding-top: 6mm;
        }}
        .equip-table {{
            width: 100%;
            border-collapse: collapse;
            border: 1.5px solid #000;
            table-layout: fixed;
        }}
        .equip-table th, .equip-table td {{
            border: 1px solid #000;
            padding: 2px 3px;
            font-size: 14pt;
            vertical-align: middle;
            font-weight: normal; /* ตัวบางตามคำสั่ง */
        }}
        .equip-table th {{
            background-color: #fafafa;
            font-weight: normal; /* หัวตารางตัวบาง */
            font-size: 11.5pt;   /* บรรทัดเดียว */
            text-align: center;
            height: 11mm;
            line-height: 1.0;
            white-space: nowrap; /* บรรทัดเดียว */
            padding: 2px 2px;
        }}
        .equip-table td {{
            height: 22mm;
        }}

        /* PAGE 3: CALENDAR MATRIX */
        .p3-title {{
            text-align: center;
            font-size: 16pt;
            font-weight: 700;
            margin-bottom: 4mm;
            margin-top: 4mm;
        }}
        .dotted-fill {{
            display: inline-block;
            border-bottom: 1.5px dotted #000;
            padding: 0 4px;
        }}
        .cal-table {{
            width: 100%;
            border-collapse: collapse;
            border: 1.5px solid #000;
            table-layout: fixed;
        }}
        .cal-table th, .cal-table td {{
            border: 1px solid #000;
            padding: 0;
            text-align: center;
            vertical-align: middle;
            font-weight: normal; /* ตัวบางตามคำสั่ง */
        }}
        .cal-seq-col {{
            width: 32px;
            height: 10mm;
            background-color: #f7f7f7;
        }}
        .cal-day-col {{
            width: calc((100% - 64px) / 31);
            font-size: 11pt;     /* ขนาดพอดี 1 บรรทัด */
            font-weight: normal; /* ตัวบางตามคำสั่ง */
            height: 10mm;
            background-color: #fafafa;
            white-space: nowrap; /* บรรทัดเดียว */
        }}
        .d-single {{
            display: block;
            line-height: 1.0;
        }}
        .cal-seq-cell {{
            font-size: 14pt;
            font-weight: normal; /* ตัวบาง */
            background-color: #f7f7f7;
        }}
        .cal-table tr td {{
            height: 24mm;
        }}

        /* PAGE 4: INSTRUCTIONS */
        .p4-title {{
            text-align: center;
            font-size: 26pt;
            font-weight: 700;
            margin-top: 3mm;
            margin-bottom: 6mm;
        }}
        .inst-rule {{
            margin-bottom: 4mm;
            font-size: 15pt;
            line-height: 1.35;
        }}
        .inst-p {{
            margin: 0 0 2mm 0;
            text-align: justify;
            text-indent: 12mm;
        }}
        .abbr-list {{
            margin-left: 20mm;
            margin-bottom: 2mm;
        }}
        .abbr-item {{
            display: flex;
            margin-bottom: 1mm;
            line-height: 1.25;
        }}
        .abbr-code {{
            width: 55mm;
            flex-shrink: 0;
        }}
        .abbr-desc {{
            flex-grow: 1;
        }}
    </style>
</head>
<body>

    <!-- PAGE 1: COVER -->
    <div class="page" id="page-1">
        <div class="p1-header">{p1["form_id"]}</div>
        
        <div class="p1-title-container">
            <div class="p1-title">{p1["main_title"]}</div>
        </div>

        <div class="p1-slogans">
            <div>{p1["slogans"][0]}</div>
            <div>{p1["slogans"][1]}</div>
        </div>

        <div class="p1-box-wrapper">
            <table class="p1-box-table">
                <tr>
                    <td class="p1-box-unit">{p1["unit_box"]["unit_name"]}</td>
                </tr>
                <tr>
                    <td class="p1-box-month">{p1["unit_box"]["month_year"]}</td>
                </tr>
            </table>
        </div>

        <div class="p1-footer">
            <div class="p1-sub-id">{p1["sub_form_id"]}</div>
            <div class="p1-authority">{p1["footnote_authority"]}</div>
        </div>
    </div>

    <!-- PAGE 2: EQUIPMENT TABLE (ลบ - ๒ - ออกแล้ว, ตัวบาง, หัวตารางบรรทัดเดียว) -->
    <div class="page" id="page-2">
        <div class="page-2-container">
            <table class="equip-table">
                <thead>
                    <tr>
                        {headers_p2_html}
                    </tr>
                </thead>
                <tbody>
                    {rows_p2_html}
                </tbody>
            </table>
        </div>
    </div>

    <!-- PAGE 3: CALENDAR MATRIX (ตัวบาง, หัวตาราง ๑ ถึง ๓๑ บรรทัดเดียว) -->
    <div class="page" id="page-3">
        <div class="p3-title">ตารางการปรนนิบัติบำรุงประจำ <span class="dotted-fill">ปบค.๙๕ เอ็ม. ๑๐๑ เอ. ๑</span> เดือน <span class="dotted-fill">&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;พ.ค.&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;</span> พ.ศ. <span class="dotted-fill">&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;๒๕๖๙&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;</span> .</div>

        <table class="cal-table">
            <thead>
                <tr>
                    {days_th_html}
                </tr>
            </thead>
            <tbody>
                {rows_p3_html}
            </tbody>
        </table>
    </div>

    <!-- PAGE 4: INSTRUCTIONS -->
    <div class="page" id="page-4">
        <div class="p4-title">{p4["title"]}</div>
        {inst_html}
    </div>

</body>
</html>
"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"Generated HTML: {output_path}")


# -----------------------------------------------------------------------------
# 2. GENERATE PRODUCTION DOCX
# -----------------------------------------------------------------------------
def set_cell_margins(cell, top=80, bottom=80, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b = OxmlElement(f'w:{edge}')
            b.set(qn('w:val'), edge_data.get('val', 'single'))
            b.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), edge_data.get('color', '000000'))
            tcBorders.append(b)
        else:
            b = OxmlElement(f'w:{edge}')
            b.set(qn('w:val'), 'none')
            tcBorders.append(b)
    tcPr.append(tcBorders)

def generate_docx(output_path):
    p1 = bp["page_1_cover"]
    p2 = bp["page_2_equipment_table"]
    p3 = bp["page_3_calendar_matrix"]
    p4 = bp["page_4_instructions"]

    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width = Inches(11.69)   # A4 Landscape
    sec.page_height = Inches(8.27)
    sec.top_margin = Inches(0.45)
    sec.bottom_margin = Inches(0.40)
    sec.left_margin = Inches(0.60)
    sec.right_margin = Inches(0.60)

    border_all = {'val': 'single', 'sz': 8, 'color': '000000'}

    # PAGE 1: COVER
    p_id = doc.add_paragraph()
    p_id.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_id = p_id.add_run(p1["form_id"])
    r_id.font.name = "TH Sarabun PSK"
    r_id.font.size = Pt(20)
    r_id.font.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(24)
    r_t = p_title.add_run(p1["main_title"])
    r_t.font.name = "TH Sarabun PSK"
    r_t.font.size = Pt(36)
    r_t.font.bold = True

    for slogan in p1["slogans"]:
        p_s = doc.add_paragraph()
        p_s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_s.paragraph_format.space_before = Pt(0)
        p_s.paragraph_format.space_after = Pt(4)
        r_s = p_s.add_run(slogan)
        r_s.font.name = "TH Sarabun PSK"
        r_s.font.size = Pt(20)
        r_s.font.bold = True

    tbl_box = doc.add_table(rows=2, cols=1)
    tbl_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_box.autofit = False
    cell_u = tbl_box.cell(0, 0)
    cell_u.width = Inches(4.0)
    cell_u.text = p1["unit_box"]["unit_name"]
    p_u = cell_u.paragraphs[0]
    p_u.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_u = p_u.runs[0]
    r_u.font.name = "TH Sarabun PSK"
    r_u.font.size = Pt(18)
    r_u.font.bold = True

    cell_m = tbl_box.cell(1, 0)
    cell_m.width = Inches(4.0)
    cell_m.text = p1["unit_box"]["month_year"]
    p_m = cell_m.paragraphs[0]
    p_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_m = p_m.runs[0]
    r_m.font.name = "TH Sarabun PSK"
    r_m.font.size = Pt(18)
    r_m.font.bold = True

    set_cell_border(cell_u, top=border_all, bottom=border_all, left=border_all, right=border_all)
    set_cell_border(cell_m, top=border_all, bottom=border_all, left=border_all, right=border_all)
    set_cell_margins(cell_u, top=100, bottom=100, left=150, right=150)
    set_cell_margins(cell_m, top=100, bottom=100, left=150, right=150)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(80)

    tbl_foot = doc.add_table(rows=1, cols=2)
    tbl_foot.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_foot.autofit = False
    cf_l = tbl_foot.cell(0, 0)
    cf_l.width = Inches(5.0)
    cf_l.text = p1["sub_form_id"]
    p_fl = cf_l.paragraphs[0]
    r_fl = p_fl.runs[0]
    r_fl.font.name = "TH Sarabun PSK"
    r_fl.font.size = Pt(20)
    r_fl.font.bold = True

    cf_r = tbl_foot.cell(0, 1)
    cf_r.width = Inches(5.0)
    cf_r.text = p1["footnote_authority"]
    p_fr = cf_r.paragraphs[0]
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_fr = p_fr.runs[0]
    r_fr.font.name = "TH Sarabun PSK"
    r_fr.font.size = Pt(15)
    r_fr.font.bold = True

    set_cell_border(cf_l)
    set_cell_border(cf_r)

    # PAGE 2: EQUIPMENT TABLE (ลบ - ๒ - ออกแล้ว, ตัวบาง, หัวตารางบรรทัดเดียว)
    doc.add_page_break()
    p_space_p2 = doc.add_paragraph()
    p_space_p2.paragraph_format.space_before = Pt(10)

    tbl_eq = doc.add_table(rows=7, cols=8)
    tbl_eq.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths_eq = [Inches(0.42), Inches(0.85), Inches(1.25), Inches(1.15), Inches(2.10), Inches(1.45), Inches(1.35), Inches(1.90)]
    hdr_titles = [col["title"] for col in p2["columns"]]

    for c_idx, title in enumerate(hdr_titles):
        cell = tbl_eq.cell(0, c_idx)
        cell.width = col_widths_eq[c_idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.text = title
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = "TH Sarabun PSK"
        p.runs[0].font.size = Pt(11.5)   # บรรทัดเดียว
        p.runs[0].font.bold = False      # ตัวบางตามคำสั่ง
        set_cell_border(cell, top=border_all, bottom=border_all, left=border_all, right=border_all)
        set_cell_margins(cell, top=50, bottom=50, left=40, right=40)

    for r_idx, r_data in enumerate(p2["rows"]):
        row = tbl_eq.rows[r_idx + 1]
        vals = [r_data["seq"], r_data["driver"], r_data["supervisor"], r_data["equipment_name"], r_data["last_pm_date_mile"], r_data["unit_internal_no"], r_data["next_pm_date"], r_data["registration_no"]]
        for c_idx, val in enumerate(vals):
            cell = row.cells[c_idx]
            cell.width = col_widths_eq[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = val
            p = cell.paragraphs[0]
            if c_idx in (0, 7):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if p.runs:
                p.runs[0].font.name = "TH Sarabun PSK"
                p.runs[0].font.size = Pt(14)
                p.runs[0].font.bold = False # ตัวบางตามคำสั่ง
            set_cell_border(cell, top=border_all, bottom=border_all, left=border_all, right=border_all)
            set_cell_margins(cell, top=140, bottom=140, left=50, right=50)

    # PAGE 3: CALENDAR MATRIX (หัวตาราง ๑ ถึง ๓๑ บรรทัดเดียว & ตัวบาง)
    doc.add_page_break()
    p3_title = doc.add_paragraph()
    p3_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_p3t = p3_title.add_run(p3["table_title"])
    r_p3t.font.name = "TH Sarabun PSK"
    r_p3t.font.size = Pt(16)
    r_p3t.font.bold = True

    tbl_cal = doc.add_table(rows=7, cols=33)
    tbl_cal.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cal.autofit = False

    w_seq = Inches(0.4)
    w_day = Inches(0.3)
    
    # Header cells
    c0 = tbl_cal.cell(0, 0)
    c0.width = w_seq
    c0.text = ""
    set_cell_border(c0, top=border_all, bottom=border_all, left=border_all, right=border_all)

    for d in range(1, 32):
        thai_d = str(d).translate(str.maketrans("0123456789", "๐๑๒๓๔๕๖๗๘๙"))
        c = tbl_cal.cell(0, d)
        c.width = w_day
        c.text = thai_d # บรรทัดเดียว
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = "TH Sarabun PSK"
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.bold = False # ตัวบาง
        set_cell_border(c, top=border_all, bottom=border_all, left=border_all, right=border_all)

    c_end = tbl_cal.cell(0, 32)
    c_end.width = w_seq
    c_end.text = ""
    set_cell_border(c_end, top=border_all, bottom=border_all, left=border_all, right=border_all)

    for r_idx, r_data in enumerate(p3["rows"]):
        row = tbl_cal.rows[r_idx + 1]
        cl = row.cells[0]
        cl.width = w_seq
        cl.text = r_data["seq"]
        cl.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cl.paragraphs[0].runs[0].font.name = "TH Sarabun PSK"
        cl.paragraphs[0].runs[0].font.size = Pt(14)
        cl.paragraphs[0].runs[0].font.bold = False # ตัวบาง
        set_cell_border(cl, top=border_all, bottom=border_all, left=border_all, right=border_all)
        set_cell_margins(cl, top=150, bottom=150, left=15, right=15)

        for d in range(1, 32):
            cd = row.cells[d]
            cd.width = w_day
            cd.text = ""
            set_cell_border(cd, top=border_all, bottom=border_all, left=border_all, right=border_all)
            set_cell_margins(cd, top=150, bottom=150, left=10, right=10)

        cr = row.cells[32]
        cr.width = w_seq
        cr.text = r_data["seq"]
        cr.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr.paragraphs[0].runs[0].font.name = "TH Sarabun PSK"
        cr.paragraphs[0].runs[0].font.size = Pt(14)
        cr.paragraphs[0].runs[0].font.bold = False # ตัวบาง
        set_cell_border(cr, top=border_all, bottom=border_all, left=border_all, right=border_all)
        set_cell_margins(cr, top=150, bottom=150, left=15, right=15)

    # PAGE 4: INSTRUCTIONS
    doc.add_page_break()
    p4_title = doc.add_paragraph()
    p4_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_p4t = p4_title.add_run(p4["title"])
    r_p4t.font.name = "TH Sarabun PSK"
    r_p4t.font.size = Pt(24)
    r_p4t.font.bold = True
    p4_title.paragraph_format.space_after = Pt(12)

    for rule in p4["rules"]:
        pr = doc.add_paragraph()
        pr.paragraph_format.first_line_indent = Inches(0.5)
        pr.paragraph_format.space_after = Pt(4)
        r_head = pr.add_run(f"{rule['index']}. {rule['title']} – ")
        r_head.font.name = "TH Sarabun PSK"
        r_head.font.size = Pt(15)
        r_head.font.bold = True

        r_body = pr.add_run(rule["content"])
        r_body.font.name = "TH Sarabun PSK"
        r_body.font.size = Pt(15)

        if "abbreviations" in rule:
            for abb in rule["abbreviations"]:
                p_abb = doc.add_paragraph()
                p_abb.paragraph_format.left_indent = Inches(0.8)
                p_abb.paragraph_format.space_after = Pt(2)
                r_ac = p_abb.add_run(f"{abb['code']}  ")
                r_ac.font.name = "TH Sarabun PSK"
                r_ac.font.size = Pt(15)
                r_ac.font.bold = True

                r_ad = p_abb.add_run(abb["description"])
                r_ad.font.name = "TH Sarabun PSK"
                r_ad.font.size = Pt(15)

    doc.save(output_path)
    print(f"Generated DOCX: {output_path}")


# -----------------------------------------------------------------------------
# 3. GENERATE PRODUCTION XLSX
# -----------------------------------------------------------------------------
def generate_xlsx(output_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    p1 = bp["page_1_cover"]
    p2 = bp["page_2_equipment_table"]
    p3 = bp["page_3_calendar_matrix"]
    p4 = bp["page_4_instructions"]

    font_title = Font(name="TH Sarabun PSK", size=24, bold=True)
    font_header = Font(name="TH Sarabun PSK", size=16, bold=True)
    font_regular = Font(name="TH Sarabun PSK", size=14, bold=False)
    font_small_regular = Font(name="TH Sarabun PSK", size=11, bold=False)

    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=False)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=False)

    # SHEET 1: หน้าปก_ทบ468_360
    ws1 = wb.create_sheet(title="หน้าปก_ทบ468_360")
    ws1.views.sheetView[0].showGridLines = True
    ws1["H2"] = p1["form_id"]
    ws1["H2"].font = font_header
    ws1["H2"].alignment = Alignment(horizontal='right')

    ws1["C6"] = p1["main_title"]
    ws1["C6"].font = font_title
    ws1["C6"].alignment = Alignment(horizontal='center', vertical='center')

    ws1["C9"] = p1["slogans"][0]
    ws1["C9"].font = font_header
    ws1["C9"].alignment = Alignment(horizontal='center', vertical='center')

    ws1["C10"] = p1["slogans"][1]
    ws1["C10"].font = font_header
    ws1["C10"].alignment = Alignment(horizontal='center', vertical='center')

    ws1["D13"] = p1["unit_box"]["unit_name"]
    ws1["D13"].font = font_header
    ws1["D13"].alignment = Alignment(horizontal='center', vertical='center')
    ws1["D13"].border = thin_border

    ws1["D14"] = p1["unit_box"]["month_year"]
    ws1["D14"].font = font_header
    ws1["D14"].alignment = Alignment(horizontal='center', vertical='center')
    ws1["D14"].border = thin_border

    ws1["B20"] = p1["sub_form_id"]
    ws1["B20"].font = font_header
    ws1["B20"].alignment = Alignment(horizontal='left')

    ws1["G20"] = p1["footnote_authority"]
    ws1["G20"].font = font_regular
    ws1["G20"].alignment = Alignment(horizontal='right')

    # SHEET 2: ตารางคุมยุทธภัณฑ์ (ลบ - ๒ -, ตัวบาง, บรรทัดเดียว)
    ws2 = wb.create_sheet(title="ตารางคุมยุทธภัณฑ์")
    ws2.views.sheetView[0].showGridLines = True

    headers_p2 = [col["title"] for col in p2["columns"]]
    col_widths_ws2 = [8, 16, 22, 20, 36, 26, 24, 32]

    for col_idx, (h_text, w) in enumerate(zip(headers_p2, col_widths_ws2), start=1):
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = w
        cell = ws2.cell(row=2, column=col_idx, value=h_text)
        cell.font = font_regular  # ตัวบางตามคำสั่ง
        cell.alignment = align_center
        cell.border = thin_border
        cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    ws2.row_dimensions[2].height = 25

    for r_idx, r_data in enumerate(p2["rows"], start=3):
        ws2.row_dimensions[r_idx].height = 35
        row_vals = [r_data["seq"], r_data["driver"], r_data["supervisor"], r_data["equipment_name"], r_data["last_pm_date_mile"], r_data["unit_internal_no"], r_data["next_pm_date"], r_data["registration_no"]]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws2.cell(row=r_idx, column=col_idx, value=val)
            cell.font = font_regular # ตัวบาง
            cell.border = thin_border
            if col_idx in (1, 8):
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # SHEET 3: ปฏิทินปรนนิบัติ_31วัน (ตัวบาง, บรรทัดเดียว)
    ws3 = wb.create_sheet(title="ปฏิทินปรนนิบัติ_31วัน")
    ws3.views.sheetView[0].showGridLines = True
    ws3["A1"] = p3["table_title"]
    ws3["A1"].font = font_header
    ws3["A1"].alignment = Alignment(horizontal='center', vertical='center')

    ws3.column_dimensions['A'].width = 6
    cell_a3 = ws3.cell(row=3, column=1, value="")
    cell_a3.border = thin_border
    cell_a3.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    for d in range(1, 32):
        col_l = get_column_letter(d + 1)
        ws3.column_dimensions[col_l].width = 4.5
        thai_d = str(d).translate(str.maketrans("0123456789", "๐๑๒๓๔๕๖๗๘๙"))
        c = ws3.cell(row=3, column=d + 1, value=thai_d)
        c.font = font_small_regular # ตัวบาง
        c.alignment = align_center
        c.border = thin_border
        c.fill = PatternFill(start_color="FAFAFA", end_color="FAFAFA", fill_type="solid")

    col_last = get_column_letter(33)
    ws3.column_dimensions[col_last].width = 6
    c_end = ws3.cell(row=3, column=33, value="")
    c_end.border = thin_border
    c_end.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    ws3.row_dimensions[3].height = 25

    for r_idx, r_data in enumerate(p3["rows"], start=4):
        ws3.row_dimensions[r_idx].height = 35
        c_l = ws3.cell(row=r_idx, column=1, value=r_data["seq"])
        c_l.font = font_regular # ตัวบาง
        c_l.alignment = align_center
        c_l.border = thin_border
        c_l.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

        for d in range(1, 32):
            c_d = ws3.cell(row=r_idx, column=d + 1, value="")
            c_d.border = thin_border

        c_r = ws3.cell(row=r_idx, column=33, value=r_data["seq"])
        c_r.font = font_regular # ตัวบาง
        c_r.alignment = align_center
        c_r.border = thin_border
        c_r.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    # SHEET 4: คำแนะนำการปฏิบัติ
    ws4 = wb.create_sheet(title="คำแนะนำการปฏิบัติ")
    ws4.views.sheetView[0].showGridLines = True
    ws4.column_dimensions['A'].width = 6
    ws4.column_dimensions['B'].width = 100

    ws4["B2"] = p4["title"]
    ws4["B2"].font = font_title
    ws4["B2"].alignment = align_left

    cur_row = 4
    for rule in p4["rules"]:
        ws4.cell(row=cur_row, column=2, value=f"{rule['index']}. {rule['title']} – {rule['content']}")
        ws4.cell(row=cur_row, column=2).font = font_regular
        ws4.cell(row=cur_row, column=2).alignment = align_left
        cur_row += 2
        if "abbreviations" in rule:
            for abb in rule["abbreviations"]:
                ws4.cell(row=cur_row, column=2, value=f"    {abb['code']}    {abb['description']}")
                ws4.cell(row=cur_row, column=2).font = font_regular
                ws4.cell(row=cur_row, column=2).alignment = align_left
                cur_row += 1
            cur_row += 1

    wb.save(output_path)
    print(f"Generated XLSX: {output_path}")


# -----------------------------------------------------------------------------
# MAIN RUNNER
# -----------------------------------------------------------------------------
if __name__ == "__main__":
    base_name = "ทบ_468_360_ปรนนิบัติบำรุง_Rebuild_Thai"
    html_out = f"{base_name}.html"
    docx_out = f"{base_name}.docx"
    xlsx_out = f"{base_name}.xlsx"

    generate_html(html_out)
    generate_docx(docx_out)
    generate_xlsx(xlsx_out)
    print("\nAll formats compiled successfully!")
