import fitz # PyMuPDF
import docx
import os

pdf_path = 'docs/thai-docx.skill/ระบบเอกสาร/11.ทบ468-360 ปรนนิบัติบำรุง.pdf'
doc = fitz.open(pdf_path)

output_dir = 'C:/Users/EVE/.gemini/antigravity-ide/brain/7ca9cd9b-6ae7-44b7-ab50-280c021eef14'

for i, page in enumerate(doc):
    pix = page.get_pixmap(dpi=150)
    out_file = os.path.join(output_dir, f'orig_tb468_360_p{i+1}.png')
    pix.save(out_file)
    print(f'Page {i+1} saved to {out_file}, rect: {page.rect}')

docx_path = 'docs/thai-docx.skill/โรงปืน/Ref_เอกสารโรงปืน/Ref/02_แบบฟอร์มมาตรฐาน_Standard_DOCX_and_PDF/ทบ_468_360_ปรนนิบัติบำรุง.docx'
if os.path.exists(docx_path):
    d = docx.Document(docx_path)
    print("\n--- DOCX Inspection ---")
    print(f"Paragraphs: {len(d.paragraphs)}, Tables: {len(d.tables)}")
    for j, p in enumerate(d.paragraphs[:15]):
        if p.text.strip():
            print(f"P{j}: {p.text.strip()}")
    for t_idx, table in enumerate(d.tables):
        print(f"Table {t_idx}: {len(table.rows)} rows, {len(table.columns)} cols")
        for r_idx, r in enumerate(table.rows[:3]):
            cells_text = [c.text.strip().replace('\n', ' ') for c in r.cells]
            print(f"  Row {r_idx}: {cells_text[:6]}")
