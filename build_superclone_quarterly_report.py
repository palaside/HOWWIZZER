#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SUPERCLONE Production Engine
Document: รายงานผลการตรวจสอบอาการถอยและจำลองการยิง ป. และเครื่องให้ทาง สูง-ต่ำ, ซ้าย-ขวา ป. (ไตรมาส ก.ค. - ก.ย. ๖๙)
Authority: คำสั่ง ป.๔ พัน.๔ ที่ ๘๓ / ๒๕๔๔ ลง ๔ พ.ค. ๔๔
Unit: ป.๔ พัน.๔ ร้อย.๒
Formats: HTML (Pixel-Perfect Multi-page Mixed Orientation), DOCX (2 Sections Portrait/Landscape), XLSX (3 Sheets)
"""

import os
import sys
import json
import base64

if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

# Load SSOT Blueprint
BLUEPRINT_PATH = os.path.join(os.path.dirname(__file__), "table_layout_quarterly_report.json")
with open(BLUEPRINT_PATH, "r", encoding="utf-8") as f:
    bp = json.load(f)

# Garuda Crest (From D:\Project\DOCS HOWWIZ\Garuda_Vector.svg.png)
GARUDA_PATH = os.path.join(os.path.dirname(__file__), "garuda_vector_official.png")
if not os.path.exists(GARUDA_PATH):
    GARUDA_PATH = os.path.join(os.path.dirname(__file__), "garuda_tr.png")
if not os.path.exists(GARUDA_PATH):
    GARUDA_PATH = r"D:\Project\DOCS HOWWIZ\Garuda_Vector.svg.png"

with open(GARUDA_PATH, "rb") as f:
    GARUDA_B64 = base64.b64encode(f.read()).decode("utf-8")


# -----------------------------------------------------------------------------
# 1. GENERATE HTML (Pixel-Perfect Mixed Portrait / Landscape)
# -----------------------------------------------------------------------------
def generate_html(output_path):
    p1 = bp["page_1_memo"]
    p2 = bp["page_2_levels_inspection"]
    p3 = bp["page_3_recoil_inspection"]

    # Page 2 Rows HTML
    p2_rows_html = ""
    for r_idx, r in enumerate(p2["rows"]):
        val_cells = "".join([f'<td class="text-center">{v}</td>' for v in r["values"]])
        if r_idx == 0:
            p2_rows_html += f"""
            <tr>
                <td class="text-center font-bold p2-item-col">{r["item"]}</td>
                {val_cells}
                <td rowspan="4" class="text-center p2-examiner-cell font-bold">{p2["examiner"]}</td>
            </tr>
            """
        else:
            p2_rows_html += f"""
            <tr>
                <td class="text-center font-bold p2-item-col">{r["item"]}</td>
                {val_cells}
            </tr>
            """

    # Page 3 Table HTML
    p3_cols_th = ""
    for c in p3["columns"][2:]:
        if "serial" in c:
            p3_cols_th += f"""
            <th class="p3-gun-th">
                <div>{c["title"]}</div>
                <div class="p3-sub-th">หมายเลข</div>
                <div class="p3-serial">{c["serial"]}</div>
            </th>
            """
        else:
            p3_cols_th += f'<th class="p3-cmd-th">{c["title"]}</th>'

    # Page 3 Row 1 Cells
    r1 = p3["row_1_result"]
    p3_r1_cells = ""
    for res in r1["results"]:
        status_cls = "text-danger font-bold" if res["is_repair"] else ""
        p3_r1_cells += f"""
        <td class="text-center p3-val-cell">
            <div class="p3-cell-date">{res["date"]}</div>
            <div class="{status_cls}">{res["status"]}</div>
        </td>
        """
    # Empty cells for commanders in row 1
    p3_r1_cells += '<td class="p3-val-cell"></td><td class="p3-val-cell"></td>'

    # Page 3 Row 2 Cells (Officers)
    r2 = p3["row_2_officers"]
    p3_r2_cells = ""
    for off in r2["officers"]:
        p3_r2_cells += f"""
        <td class="p3-officer-cell">
            <div class="p3-officer-rank">{off["rank"]}</div>
            <div class="p3-officer-name">({off["name"]})</div>
        </td>
        """
    p3_r2_cells += f"""
    <td class="p3-officer-cell">
        <div class="p3-officer-rank">{r2["coy_commander"]["rank"]}</div>
        <div class="p3-officer-name">({r2["coy_commander"]["name"]})</div>
    </td>
    <td class="p3-officer-cell">
        <div class="p3-officer-rank">{r2["bn_commander"]["rank"]}</div>
        <div class="p3-officer-name">({r2["bn_commander"]["name"]})</div>
    </td>
    """

    html_content = f"""<!DOCTYPE html>
<html lang="th">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{bp["document_title"]}</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Sarabun:ital,wght@0,100;0,200;0,300;0,400;0,600;0,700;1,400&display=swap" rel="stylesheet">
    <style>
        @page portrait-page {{
            size: A4 portrait;
            margin: 0;
        }}
        @page landscape-page {{
            size: A4 landscape;
            margin: 0;
        }}
        * {{
            box-sizing: border-box;
            -webkit-print-color-adjust: exact;
            print-color-adjust: exact;
        }}
        body {{
            margin: 0;
            padding: 20px 0;
            background-color: #525659;
            font-family: 'Sarabun', 'TH Sarabun PSK', 'THSarabunNew', sans-serif;
            font-size: 12pt;
            font-weight: 300; /* ตัวบาง (Light / Thin) */
            font-feature-settings: "locl" 1, "numr" 1;
            font-variant-numeric: tabular-nums;
            color: #000000;
            line-height: 1.5;
            display: flex;
            flex-direction: column;
            align-items: center;
        }}
        .page-portrait {{
            width: 210mm;
            height: 297mm;
            background-color: #ffffff;
            margin-bottom: 20mm;
            box-shadow: 0 4px 15px rgba(0, 0, 0, 0.35);
            position: relative;
            page-break-after: always;
            padding: 20mm 15mm 15mm 25mm; /* ขอบบน 2cm, ขวา 1.5cm, ล่าง 1.5cm, ซ้าย 2.5cm มาตรฐานงานสารบรรณ */
            overflow: hidden;
            page: portrait-page;
        }}
        .page-landscape {{
            width: 297mm;
            height: 210mm;
            background-color: #ffffff;
            margin-bottom: 20mm;
            box-shadow: 0 4px 15px rgba(0, 0, 0, 0.35);
            position: relative;
            page-break-after: always;
            padding: 15mm 18mm 15mm 18mm;
            overflow: hidden;
            page: landscape-page;
        }}
        @media print {{
            body {{
                background-color: #ffffff;
                padding: 0;
            }}
            .page-portrait, .page-landscape {{
                margin: 0;
                box-shadow: none;
                page-break-after: always;
            }}
        }}

        /* UTILITIES */
        .text-center {{ text-align: center; }}
        .text-left {{ text-align: left; }}
        .text-right {{ text-align: right; }}
        .font-bold {{ font-weight: bold; }}
        .font-normal {{ font-weight: normal; }}
        .text-danger {{ color: #d32f2f; }}

        /* ================= PAGE 1: MEMO ================= */
        .memo-copy-label {{
            position: absolute;
            top: 7mm;
            left: 0;
            right: 0;
            text-align: center;
            font-size: 12pt;
            line-height: 1.5;
        }}
        .memo-header-container {{
            position: relative;
            width: 100%;
            height: 30mm;
            margin-top: -2mm;
        }}
        .memo-garuda {{
            position: absolute;
            left: 0;
            top: 0;
            width: 30mm;
            height: 30mm;
            object-fit: contain;
        }}
        .memo-title {{
            position: absolute;
            left: 0;
            right: 0;
            top: 6mm;
            text-align: center;
            font-size: 24pt;
            font-weight: bold;
            line-height: 1;
            font-family: 'TH Sarabun PSK', 'THSarabunNew', 'Sarabun', sans-serif;
        }}
        .memo-meta-table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 12pt;
            line-height: 1.5;
            margin-top: 2mm;
        }}
        .memo-meta-table td {{
            padding: 1mm 0;
            vertical-align: bottom;
        }}
        .memo-meta-label {{
            font-size: 14pt;
            font-weight: bold;
            white-space: nowrap;
        }}
        .memo-meta-dots {{
            border-bottom: 1px dotted #000;
            padding-left: 2mm;
        }}
        .memo-content {{
            margin-top: 3mm;
            font-size: 12pt;
            line-height: 1.5;
        }}
        .memo-p {{
            text-indent: 2.5cm;
            margin: 0 0 2.5mm 0;
            text-align: justify;
            text-justify: inter-cluster;
            line-height: 1.5;
            font-size: 12pt;
        }}
        .memo-att-container {{
            margin: 0 0 2.5mm 0;
            font-size: 12pt;
            line-height: 1.5;
        }}
        .memo-att-row {{
            display: flex;
            justify-content: space-between;
            align-items: baseline;
            width: 100%;
        }}
        .memo-att-unit {{
            padding-left: 0;
            margin-bottom: 1mm;
        }}

        /* UNIVERSAL SIGNATURE BLOCK (RULE-UNIVERSAL-SIGNATURE-BLOCK-V1) */
        .sig-container {{
            margin-top: 10mm;
            display: flex;
            justify-content: flex-end;
            width: 100%;
        }}
        .sig-block {{
            display: inline-flex;
            flex-direction: column;
            align-items: flex-start; /* คำรับรองชิดซ้ายของบล็อก */
            position: relative;
            min-width: 220px;
            margin-right: 25mm;
            font-size: 12pt;
            line-height: 1.0;
        }}
        .sig-verify {{
            margin-bottom: 24pt; /* 1 Enter พร้อมระยะห่างหลังย่อหน้าตามแบบฉบับราชการ */
            font-size: 12pt;
            line-height: 1.0;
            text-align: left;
        }}
        .sig-name-pos-group {{
            display: inline-flex;
            flex-direction: column;
            align-items: center; /* ชื่อและตำแหน่งจัดกึ่งกลางร่วมกัน */
            position: relative;
            margin-left: 65px; /* เยื้องกลุ่มชื่อ-ตำแหน่งเข้ามาทางขวาให้สัมพันธ์กับคำรับรอง */
        }}
        .sig-name-wrapper {{
            position: relative;
            display: inline-block;
            white-space: nowrap;
            margin-top: 24pt; /* พื้นที่สำหรับลงลายมือชื่อจริง */
        }}
        .sig-rank-anchor {{
            position: absolute;
            right: 100%; /* จุด . ท้ายยศ ชนขอบซ้ายของกรอบชื่อ ซึ่งคือ ( พอดีเป๊ะ 100% */
            bottom: calc(100% + 18px);
            white-space: nowrap;
            font-size: 12pt;
            line-height: 1.0;
        }}
        .sig-name {{
            margin: 0;
            line-height: 1.15;
            text-align: center;
            white-space: nowrap;
            font-size: 12pt;
            display: block;
        }}
        .sig-pos {{
            margin: 0;
            line-height: 1.15;
            text-align: center;
            white-space: nowrap;
            font-size: 12pt;
            display: block;
        }}
        .sig-date {{
            margin: 0;
            line-height: 1.15;
            text-align: center;
            white-space: nowrap;
            font-size: 12pt;
            display: block;
        }}

        /* PAGE 1 ROUTING FOOTER */
        .memo-footer-routing {{
            position: absolute;
            bottom: 20mm;
            right: 20mm;
            font-size: 12pt;
            line-height: 1.5;
        }}

        /* ================= PAGE 2: LEVELS ================= */
        #page-2 {{
            padding: 20mm 15mm 15mm 15mm;
        }}
        .p2-table {{
            width: 100%;
            table-layout: fixed;
            border-collapse: collapse;
            margin-top: 5mm;
            font-size: 12pt;
            line-height: 1.5;
        }}
        .p2-table th, .p2-table td {{
            border: 1px solid #000;
            padding: 1.5mm 1mm;
            height: 9mm;
        }}
        .p2-table th {{
            font-weight: normal;
            background-color: #fcfcfc;
        }}
        .p2-date-col {{
            width: 14.0%;
            white-space: nowrap;
        }}
        .p2-gun-th {{
            width: 12.0%;
            font-size: 12pt;
            line-height: 1.2;
        }}
        .p2-item-col {{
            width: 14.0%;
        }}
        .p2-examiner-col {{
            width: 14.0%;
            line-height: 1.2;
        }}
        .p2-examiner-cell {{
            vertical-align: middle;
            font-size: 12pt;
        }}
        .p2-title-box {{
            margin-top: 14mm;
            text-align: center;
            font-size: 12pt;
            line-height: 1.5;
            width: 100%;
        }}

        /* ================= PAGE 3: RECOIL ================= */
        .p3-title-container {{
            text-align: center;
            margin-bottom: 4mm;
            line-height: 1.5;
        }}
        .p3-auth-title {{
            font-size: 12pt;
        }}
        .p3-main-title {{
            font-size: 12pt;
        }}
        .p3-unit-title {{
            font-size: 12pt;
        }}
        .p3-table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 12pt;
            line-height: 1.5;
        }}
        .p3-table th, .p3-table td {{
            border: 1px solid #000;
            padding: 1.5mm 1mm;
        }}
        .p3-table th {{
            font-weight: normal;
            vertical-align: middle;
            background-color: #fafafa;
        }}
        .p3-gun-th {{
            line-height: 1.2;
            font-size: 12pt;
        }}
        .p3-sub-th {{
            font-size: 12pt;
        }}
        .p3-serial {{
            font-size: 12pt;
        }}
        .p3-cmd-th {{
            font-size: 12pt;
        }}
        .p3-val-cell {{
            height: 18mm;
            vertical-align: middle;
            font-size: 12pt;
            line-height: 1.4;
        }}
        .p3-cell-date {{
            margin-bottom: 1.5mm;
            white-space: nowrap;
            font-size: 12pt;
        }}
        .p3-officer-cell {{
            height: 22mm;
            vertical-align: top;
            padding: 1.5mm 1mm;
            font-size: 9pt;
            line-height: 1.15;
        }}
        .p3-officer-rank {{
            text-align: left;
            font-size: 9pt;
            padding-left: 1mm;
            line-height: 1.15;
        }}
        .p3-officer-name {{
            text-align: center;
            font-size: 9pt;
            margin-top: 8mm;
            white-space: nowrap;
            letter-spacing: -0.2px;
            line-height: 1.15;
        }}
    </style>
</head>
<body>

    <!-- PAGE 1: MEMO (A4 PORTRAIT) -->
    <div class="page-portrait" id="page-1">
        <div class="memo-copy-label">{p1["copy_type"]}</div>
        
        <div class="memo-header-container">
            <img class="memo-garuda" src="data:image/png;base64,{GARUDA_B64}" alt="ครุฑ">
            <div class="memo-title">{p1["title"]}</div>
        </div>

        <div style="margin-top: 2mm; font-size: 12pt; line-height: 1.5;">
            <div style="display: flex; align-items: baseline; margin-bottom: 1.5mm;">
                <span class="memo-meta-label" style="margin-right: 2mm;">{p1["official_department"]}</span>
                <div class="memo-meta-dots" style="flex: 1;">{p1["department_value"]}</div>
            </div>
            <div style="display: flex; align-items: baseline; margin-bottom: 1.5mm;">
                <div style="display: flex; align-items: baseline; flex: 1.2; margin-right: 6mm;">
                    <span class="memo-meta-label" style="margin-right: 2mm;">{p1["doc_number_label"]}</span>
                    <div class="memo-meta-dots" style="flex: 1;">{p1["doc_number_value"]}</div>
                </div>
                <div style="display: flex; align-items: baseline; flex: 1;">
                    <span class="memo-meta-label" style="margin-right: 2mm;">{p1["doc_date_label"]}</span>
                    <div class="memo-meta-dots" style="flex: 1;">&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;ก.ย. ๖๙</div>
                </div>
            </div>
            <div style="display: flex; align-items: baseline; margin-bottom: 1.5mm;">
                <span class="memo-meta-label" style="margin-right: 2mm;">{p1["subject_label"]}</span>
                <div style="flex: 1; padding-left: 2mm;">รายงานผลการตรวจสอบอาการถอยและจำลองการยิง ป. และเครื่องให้ทาง สูง-ต่ำ ซ้าย-ขวา ป.</div>
            </div>
        </div>

        <div class="memo-content">
            <div style="margin-bottom: 2mm;">
                <span class="font-bold">{p1["recipient_label"]}</span>&nbsp;&nbsp;{p1["recipient_value"]}
            </div>
            <div style="margin-bottom: 2mm;">
                <span class="font-bold">{p1["reference_label"]}</span>&nbsp;&nbsp;{p1["reference_value"]}
            </div>

            <div class="memo-att-container">
                <table style="width: 100%; border-collapse: collapse; font-size: 12pt; line-height: 1.5;">
                    <tr>
                        <td style="width: 28mm; vertical-align: top; white-space: nowrap;"><span class="font-bold">สิ่งที่ส่งมาด้วย</span></td>
                        <td style="vertical-align: top;">๑. ตารางการทดสอบอาการถอย ป.</td>
                        <td style="width: 28mm; vertical-align: top; text-align: right; white-space: nowrap;">จำนวน  ๑  ฉบับ</td>
                    </tr>
                    <tr>
                        <td></td>
                        <td style="vertical-align: top; padding-top: 1mm;">๒. ผลการตรวจระดับเครื่องให้ทาง สูง-ต่ำ, ซ้าย-ขวา</td>
                        <td style="vertical-align: top; text-align: right; white-space: nowrap; padding-top: 1mm;">จำนวน  ๑  ฉบับ</td>
                    </tr>
                </table>
            </div>

            <p class="memo-p">ตามอ้างถึง ป.๔ พัน.๔ ร้อย.๒ ได้ดำเนินการทดสอบอาการถอย ป. และตรวจระดับเครื่องช่วยให้ทาง สูง-ต่ำ, ซ้าย-ขวา ห้วงระหว่าง ก.ค. – ก.ย. ๖๙ เมื่อ ๓๐ ก.ย. ๖๙ เพื่อรายงานให้หน่วยเหนือทราบ รายละเอียดตามสิ่งที่ส่งมาด้วย</p>
            <p class="memo-p">จึงเรียนมาเพื่อกรุณาทราบ</p>
        </div>

        <!-- SIGNATURE BLOCK (พิกัด 130 มม. นับจากขอบซ้าย A4 210 มม., เคาะ 5 Enter จากบรรทัดสุดท้าย) -->
        <div style="margin-top: 55pt; padding-left: 105mm;">
            <div class="sig-block" style="margin-right: 0;">
                <div class="sig-name-pos-group" style="margin-left: 0;">
                    <div class="sig-name-wrapper" style="margin-top: 18pt;">
                        <span class="sig-rank-anchor">{p1["signature"]["rank"]}</span>
                        <div class="sig-name">({p1["signature"]["name"]})</div>
                    </div>
                    <div class="sig-pos">{p1["signature"]["position"]}</div>
                </div>
            </div>
        </div>

        <!-- ROUTING FOOTER -->
        <div class="memo-footer-routing">
            <div>{p1["routing_footer"][0]["rank"]}&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;{p1["routing_footer"][0]["action"]}&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;{p1["routing_footer"][0]["date"]}</div>
            <div>{p1["routing_footer"][1]["rank"]} {p1["routing_footer"][1]["name"]}&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;{p1["routing_footer"][1]["action"]}&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;30&nbsp;&nbsp;&nbsp;&nbsp;{p1["routing_footer"][1]["date"]}</div>
        </div>
    </div>

    <!-- PAGE 2: LEVELS (A4 PORTRAIT) -->
    <div class="page-portrait" id="page-2">
        <table class="p2-table">
            <thead>
                <tr>
                    <th rowspan="2" class="p2-date-col text-center font-bold" style="vertical-align: middle;">{p2["date_column"]}</th>
                    <th colspan="6" class="text-center font-bold">ผลการตรวจ</th>
                    <th rowspan="2" class="p2-examiner-col text-center font-bold" style="vertical-align: middle;">ผู้ตรวจ<br>สอบ</th>
                </tr>
                <tr>
                    <th class="p2-gun-th text-center"><div style="white-space: nowrap;">{p2["columns"][0]["name"]}</div><div>{p2["columns"][0]["serial"]}</div></th>
                    <th class="p2-gun-th text-center"><div style="white-space: nowrap;">{p2["columns"][1]["name"]}</div><div>{p2["columns"][1]["serial"]}</div></th>
                    <th class="p2-gun-th text-center"><div style="white-space: nowrap;">{p2["columns"][2]["name"]}</div><div>{p2["columns"][2]["serial"]}</div></th>
                    <th class="p2-gun-th text-center"><div style="white-space: nowrap;">{p2["columns"][3]["name"]}</div><div>{p2["columns"][3]["serial"]}</div></th>
                    <th class="p2-gun-th text-center"><div style="white-space: nowrap;">{p2["columns"][4]["name"]}</div><div>{p2["columns"][4]["serial"]}</div></th>
                    <th class="p2-gun-th text-center"><div style="white-space: nowrap;">{p2["columns"][5]["name"]}</div><div>{p2["columns"][5]["serial"]}</div></th>
                </tr>
            </thead>
            <tbody>
                {p2_rows_html}
            </tbody>
        </table>

        <div class="p2-title-box">
            <div>{p2["title"]}</div>
            <div>{p2["subtitle"]}</div>
            <div>{p2["period"]}</div>
        </div>
    </div>

    <!-- PAGE 3: RECOIL (A4 LANDSCAPE) -->
    <div class="page-landscape" id="page-3">
        <div class="p3-title-container">
            <div class="p3-auth-title">{p3["title_authority"]}</div>
            <div class="p3-main-title">{p3["title_main"]}</div>
            <div class="p3-unit-title">{p3["unit"]}</div>
        </div>

        <table class="p3-table">
            <thead>
                <tr>
                    <th rowspan="2" style="width: {p3['columns'][0]['width']};" class="text-center font-bold">ครั้งที่</th>
                    <th rowspan="2" style="width: {p3['columns'][1]['width']};" class="text-center font-bold">รายการ</th>
                    <th colspan="8" class="text-center font-bold">ผู้รับผิดชอบ</th>
                </tr>
                <tr>
                    {p3_cols_th}
                </tr>
            </thead>
            <tbody>
                <tr>
                    <td class="text-center font-bold">{r1["seq"]}</td>
                    <td class="text-center">
                        <div>{r1["date_label"]}</div>
                        <div>{r1["result_label"]}</div>
                    </td>
                    {p3_r1_cells}
                </tr>
                <tr>
                    <td></td>
                    <td class="text-center font-bold">{r2["role"]}</td>
                    {p3_r2_cells}
                </tr>
            </tbody>
        </table>

        <!-- SIGNATURE BLOCK (เลื่อนชิดมาทางขวา พิกัดอิง 250 มม. / ขอบขวา 20 มม.) -->
        <div style="position: absolute; right: 20mm; top: 125mm;">
            <div class="sig-block" style="margin-right: 0;">
                <div class="sig-verify">{p3["signature"]["certified"]}</div>
                <div class="sig-name-pos-group">
                    <div class="sig-name-wrapper">
                        <span class="sig-rank-anchor">{p3["signature"]["rank"]}</span>
                        <div class="sig-name">({p3["signature"]["name"]})</div>
                    </div>
                    <div class="sig-pos">{p3["signature"]["position"]}</div>
                    <div class="sig-date">{p3["signature"]["date"]}</div>
                </div>
            </div>
        </div>
    </div>

</body>
</html>
"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"Generated HTML: {output_path}")


# -----------------------------------------------------------------------------
# 2. GENERATE DOCX (2 Sections: Portrait then Landscape)
# -----------------------------------------------------------------------------
def set_cell_margins(cell, top=100, bottom=100, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b = OxmlElement(f'w:{edge}')
            b.set(qn('w:val'), edge_data.get('val', 'single'))
            b.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), edge_data.get('color', '000000'))
            tcBorders.append(b)
        else:
            b = OxmlElement(f'w:{edge}')
            b.set(qn('w:val'), 'none')
            tcBorders.append(b)
    tcPr.append(tcBorders)

def set_run_font(run, font_name="TH Sarabun PSK", size_pt=17, bold=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rPr.append(rFonts)

def generate_docx(output_path):
    p1 = bp["page_1_memo"]
    p2 = bp["page_2_levels_inspection"]
    p3 = bp["page_3_recoil_inspection"]

    doc = Document()

    # SECTION 1: A4 PORTRAIT
    sec1 = doc.sections[0]
    sec1.page_width = Inches(8.27)
    sec1.page_height = Inches(11.69)
    sec1.top_margin = Inches(0.98)    # 2.5 cm
    sec1.bottom_margin = Inches(0.79) # 2.0 cm
    sec1.left_margin = Inches(0.98)   # 2.5 cm
    sec1.right_margin = Inches(0.79)  # 2.0 cm

    border_all = {'val': 'single', 'sz': 4, 'color': '000000'}

    # --- PAGE 1: MEMO ---
    p_copy = doc.add_paragraph()
    p_copy.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_copy = p_copy.add_run(p1["copy_type"])
    set_run_font(r_copy, "TH Sarabun PSK", 17)

    tbl_head = doc.add_table(rows=1, cols=2)
    tbl_head.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_head.autofit = False
    c_img = tbl_head.cell(0, 0)
    c_img.width = Inches(1.5)
    p_img = c_img.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if os.path.exists(GARUDA_PATH):
        p_img.add_run().add_picture(GARUDA_PATH, width=Inches(1.2))

    c_title = tbl_head.cell(0, 1)
    c_title.width = Inches(5.0)
    p_title = c_title.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(14)
    r_title = p_title.add_run(p1["title"])
    set_run_font(r_title, "TH Sarabun PSK", 24, bold=True)

    set_cell_border(c_img)
    set_cell_border(c_title)

    # Meta Lines
    meta_lines = [
        (p1["official_department"], p1["department_value"]),
        (f'{p1["doc_number_label"]}  {p1["doc_number_value"]}', f'{p1["doc_date_label"]}  {p1["doc_date_value"]}'),
        (p1["subject_label"], "รายงานผลการตรวจสอบอาการถอยและจำลองการยิง ป. และเครื่องให้ทาง สูง-ต่ำ, ซ้าย-ขวา ป."),
        (p1["recipient_label"], p1["recipient_value"]),
        (p1["reference_label"], p1["reference_value"])
    ]

    for label, val in meta_lines:
        p_m = doc.add_paragraph()
        p_m.paragraph_format.space_after = Pt(2)
        p_m.paragraph_format.line_spacing = 1.5
        r_lbl = p_m.add_run(label + "  ")
        set_run_font(r_lbl, "TH Sarabun PSK", 14, bold=True)
        r_val = p_m.add_run(val)
        set_run_font(r_val, "TH Sarabun PSK", 12)

    # Attachments
    tbl_att = doc.add_table(rows=2, cols=2)
    tbl_att.style = None
    
    cell_a0 = tbl_att.cell(0, 0)
    p_a0 = cell_a0.paragraphs[0]
    p_a0.paragraph_format.space_after = Pt(2)
    p_a0.paragraph_format.line_spacing = 1.5
    r_att_lbl = p_a0.add_run(p1["attachments_label"] + "  ")
    set_run_font(r_att_lbl, "TH Sarabun PSK", 12, bold=True)
    r_att1 = p_a0.add_run(f'๑. {p1["attachments"][0]["title"]}')
    set_run_font(r_att1, "TH Sarabun PSK", 12)

    cell_a1 = tbl_att.cell(0, 1)
    p_a1 = cell_a1.paragraphs[0]
    p_a1.paragraph_format.space_after = Pt(2)
    p_a1.paragraph_format.line_spacing = 1.5
    p_a1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_q1 = p_a1.add_run(p1["attachments"][0]["quantity"])
    set_run_font(r_q1, "TH Sarabun PSK", 12)

    cell_b0 = tbl_att.cell(1, 0)
    p_b0 = cell_b0.paragraphs[0]
    p_b0.paragraph_format.space_after = Pt(2)
    p_b0.paragraph_format.line_spacing = 1.5
    p_b0.paragraph_format.left_indent = Inches(0.83)
    r_att2 = p_b0.add_run(f'๒. {p1["attachments"][1]["title"]}')
    set_run_font(r_att2, "TH Sarabun PSK", 12)

    cell_b1 = tbl_att.cell(1, 1)
    p_b1 = cell_b1.paragraphs[0]
    p_b1.paragraph_format.space_after = Pt(2)
    p_b1.paragraph_format.line_spacing = 1.5
    p_b1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_q2 = p_b1.add_run(p1["attachments"][1]["quantity"])
    set_run_font(r_q2, "TH Sarabun PSK", 12)

    # Paragraphs
    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.first_line_indent = Inches(0.98) # 2.5cm
    p_p1.paragraph_format.space_after = Pt(4)
    p_p1.paragraph_format.line_spacing = 1.5
    p_p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_p1 = p_p1.add_run("ตามอ้างถึง ป.๔ พัน.๔ ร้อย.๒ ได้ดำเนินการทดสอบอาการถอย ป. และตรวจระดับเครื่องช่วยให้ทาง สูง-ต่ำ, ซ้าย-ขวา ห้วงระหว่าง ก.ค. – ก.ย. ๖๙ เมื่อ ๓๐ ก.ย. ๖๙ เพื่อรายงานให้หน่วยเหนือทราบ รายละเอียดตามสิ่งที่ส่งมาด้วย")
    set_run_font(r_p1, "TH Sarabun PSK", 12)

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.first_line_indent = Inches(0.98)
    p_p2.paragraph_format.space_after = Pt(4)
    p_p2.paragraph_format.line_spacing = 1.5
    r_p2 = p_p2.add_run("จึงเรียนมาเพื่อกรุณาทราบ")
    set_run_font(r_p2, "TH Sarabun PSK", 12)

    # Signature Block Page 1
    p_sig1 = doc.add_paragraph()
    p_sig1.paragraph_format.space_before = Pt(24)
    p_sig1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_s1 = p_sig1.add_run(f'{p1["signature"]["rank"]}\n({p1["signature"]["name"]})\n{p1["signature"]["position"]}')
    set_run_font(r_s1, "TH Sarabun PSK", 12)

    # Routing Footer Page 1
    p_rout = doc.add_paragraph()
    p_rout.paragraph_format.space_before = Pt(40)
    p_rout.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_rt = p_rout.add_run(f'{p1["routing_footer"][0]["rank"]}                   {p1["routing_footer"][0]["action"]}              {p1["routing_footer"][0]["date"]}\n{p1["routing_footer"][1]["rank"]} {p1["routing_footer"][1]["name"]}         {p1["routing_footer"][1]["action"]}      30    {p1["routing_footer"][1]["date"]}')
    set_run_font(r_rt, "TH Sarabun PSK", 12)

    # --- PAGE 2: LEVELS (A4 PORTRAIT) ---
    doc.add_page_break()

    tbl_p2 = doc.add_table(rows=6, cols=8)
    tbl_p2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_p2.autofit = False

    col_w_p2 = [Inches(0.85), Inches(0.85), Inches(0.85), Inches(0.85), Inches(0.85), Inches(0.85), Inches(0.85), Inches(1.10)]

    # Header Row 0
    hdr0 = tbl_p2.rows[0]
    hdr0.cells[0].merge(tbl_p2.rows[1].cells[0])
    hdr0.cells[0].text = p2["date_column"]

    hdr0.cells[1].merge(hdr0.cells[6])
    hdr0.cells[1].text = "ผลการตรวจ"

    hdr0.cells[7].merge(tbl_p2.rows[1].cells[7])
    hdr0.cells[7].text = "ผู้ตรวจสอบ"

    # Header Row 1 Sub headers
    for idx, col_info in enumerate(p2["columns"], start=1):
        c = tbl_p2.rows[1].cells[idx]
        c.text = f'{col_info["name"]}\n{col_info["serial"]}'

    # Data Rows (2 to 5)
    for r_idx, r_data in enumerate(p2["rows"], start=2):
        row = tbl_p2.rows[r_idx]
        row.cells[0].text = r_data["item"]
        for c_idx, val in enumerate(r_data["values"], start=1):
            row.cells[c_idx].text = val

    # Merge examiner cell on rows 2 to 5
    tbl_p2.rows[2].cells[7].merge(tbl_p2.rows[5].cells[7])
    tbl_p2.rows[2].cells[7].text = p2["examiner"]

    # Style Table 2 cells
    for row in tbl_p2.rows:
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_w_p2[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell, top=border_all, bottom=border_all, left=border_all, right=border_all)
            set_cell_margins(cell, top=60, bottom=60, left=40, right=40)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    set_run_font(r, "TH Sarabun PSK", 12)

    # Page 2 Descriptive Box
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(36)
    p_desc.paragraph_format.line_spacing = 1.5
    p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_d = p_desc.add_run(f'{p2["title"]}\n{p2["subtitle"]}\n{p2["period"]}')
    set_run_font(r_d, "TH Sarabun PSK", 12)

    # --- SECTION 2: A4 LANDSCAPE (PAGE 3: RECOIL) ---
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.page_width = Inches(11.69)
    sec2.page_height = Inches(8.27)
    sec2.top_margin = Inches(0.60)
    sec2.bottom_margin = Inches(0.60)
    sec2.left_margin = Inches(0.70)
    sec2.right_margin = Inches(0.70)

    # Page 3 Title
    p_p3_head = doc.add_paragraph()
    p_p3_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_p3_head.paragraph_format.space_after = Pt(10)
    p_p3_head.paragraph_format.line_spacing = 1.5
    r_p3h = p_p3_head.add_run(f'{p3["title_authority"]}\n{p3["title_main"]}\n{p3["unit"]}')
    set_run_font(r_p3h, "TH Sarabun PSK", 12)

    # Page 3 Table
    tbl_p3 = doc.add_table(rows=4, cols=10)
    tbl_p3.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_p3.autofit = False

    col_w_p3 = [Inches(0.55), Inches(1.30), Inches(1.10), Inches(1.10), Inches(1.10), Inches(1.10), Inches(1.10), Inches(1.10), Inches(0.95), Inches(0.95)]

    # Header Row 0
    h3_0 = tbl_p3.rows[0]
    h3_0.cells[0].merge(tbl_p3.rows[1].cells[0])
    h3_0.cells[0].text = "ครั้งที่"

    h3_0.cells[1].merge(tbl_p3.rows[1].cells[1])
    h3_0.cells[1].text = "รายการ"

    h3_0.cells[2].merge(h3_0.cells[9])
    h3_0.cells[2].text = "ผู้รับผิดชอบ"

    # Header Row 1 Sub headers
    for idx, c_info in enumerate(p3["columns"][2:], start=2):
        cell = tbl_p3.rows[1].cells[idx]
        if "serial" in c_info:
            cell.text = f'{c_info["title"]}\nหมายเลข\n{c_info["serial"]}'
        else:
            cell.text = c_info["title"]

    # Data Row 2 (Results)
    r2_row = tbl_p3.rows[2]
    r2_row.cells[0].text = p3["row_1_result"]["seq"]
    r2_row.cells[1].text = f'{p3["row_1_result"]["date_label"]}\n{p3["row_1_result"]["result_label"]}'
    for c_idx, res in enumerate(p3["row_1_result"]["results"], start=2):
        cell = r2_row.cells[c_idx]
        cell.text = f'{res["date"]}\n{res["status"]}'

    # Data Row 3 (Officers)
    r3_row = tbl_p3.rows[3]
    r3_row.cells[0].text = ""
    r3_row.cells[1].text = p3["row_2_officers"]["role"]
    
    officers_list = p3["row_2_officers"]["officers"] + [
        p3["row_2_officers"]["coy_commander"],
        p3["row_2_officers"]["bn_commander"]
    ]
    for c_idx, off in enumerate(officers_list, start=2):
        cell = r3_row.cells[c_idx]
        cell.text = "" # Clear default text
        p_rk = cell.paragraphs[0]
        p_rk.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_rk.paragraph_format.space_before = Pt(0)
        p_rk.paragraph_format.space_after = Pt(14)
        r_rk = p_rk.add_run(off["rank"])
        set_run_font(r_rk, "TH Sarabun PSK", 9)

        p_nm = cell.add_paragraph()
        p_nm.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_nm.paragraph_format.space_before = Pt(0)
        p_nm.paragraph_format.space_after = Pt(0)
        r_nm = p_nm.add_run(f'({off["name"]})')
        set_run_font(r_nm, "TH Sarabun PSK", 9)

    # Style Table 3 cells
    for row_idx, row in enumerate(tbl_p3.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_w_p3[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell, top=border_all, bottom=border_all, left=border_all, right=border_all)
            set_cell_margins(cell, top=60, bottom=60, left=30, right=30)
            
            # Specific styling for officer row (row 3, columns 2 to 9)
            if row_idx == 3 and c_idx >= 2:
                for p_idx, p in enumerate(cell.paragraphs):
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        set_run_font(r, "TH Sarabun PSK", 9)
            else:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        color = RGBColor(211, 47, 47) if "อยู่ระหว่างดำเนินการซ่อม" in r.text else None
                        bold = True if "อยู่ระหว่างดำเนินการซ่อม" in r.text else False
                        set_run_font(r, "TH Sarabun PSK", 12, bold=bold, color=color)

    # Page 3 Signature Block
    p_sig3 = doc.add_paragraph()
    p_sig3.paragraph_format.space_before = Pt(28)
    p_sig3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s3 = p_sig3.add_run(f'{p3["signature"]["certified"]}\n\n{p3["signature"]["rank"]}\n({p3["signature"]["name"]})\n{p3["signature"]["position"]}\n{p3["signature"]["date"]}')
    set_run_font(r_s3, "TH Sarabun PSK", 12)

    doc.save(output_path)
    print(f"Generated DOCX: {output_path}")


# -----------------------------------------------------------------------------
# 3. GENERATE XLSX (3 Sheets)
# -----------------------------------------------------------------------------
def generate_xlsx(output_path):
    p1 = bp["page_1_memo"]
    p2 = bp["page_2_levels_inspection"]
    p3 = bp["page_3_recoil_inspection"]

    wb = openpyxl.Workbook()
    thin = Side(border_style="thin", color="000000")
    thin_border = Border(top=thin, left=thin, right=thin, bottom=thin)

    font_regular = Font(name="TH Sarabun PSK", size=12)
    font_bold = Font(name="TH Sarabun PSK", size=12, bold=True)
    font_large_bold = Font(name="TH Sarabun PSK", size=18, bold=True)
    font_red_bold = Font(name="TH Sarabun PSK", size=12, bold=True, color="D32F2F")

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # --- SHEET 1: บันทึกข้อความ ---
    ws1 = wb.active
    ws1.title = "บันทึกข้อความ"
    ws1.views.sheetView[0].showGridLines = True

    ws1["C2"] = p1["copy_type"]
    ws1["C2"].font = font_regular
    ws1["C2"].alignment = align_center

    ws1["C4"] = p1["title"]
    ws1["C4"].font = font_large_bold
    ws1["C4"].alignment = align_center

    memo_rows = [
        ("ส่วนราชการ", p1["department_value"]),
        ("ที่", f'{p1["doc_number_value"]}       วันที่  {p1["doc_date_value"]}'),
        ("เรื่อง", p1["subject_value"]),
        ("เรียน", p1["recipient_value"]),
        ("อ้างถึง", p1["reference_value"]),
        ("สิ่งที่ส่งมาด้วย", f'๑. {p1["attachments"][0]["title"]}   {p1["attachments"][0]["quantity"]}'),
        ("", f'๒. {p1["attachments"][1]["title"]}   {p1["attachments"][1]["quantity"]}'),
        ("", p1["paragraphs"][0]),
        ("", p1["paragraphs"][1])
    ]

    for idx, (lbl, val) in enumerate(memo_rows, start=6):
        ws1[f"B{idx}"] = lbl
        ws1[f"B{idx}"].font = font_bold
        ws1[f"C{idx}"] = val
        ws1[f"C{idx}"].font = font_regular
        ws1.row_dimensions[idx].height = 28

    ws1["E18"] = f'{p1["signature"]["rank"]}\n({p1["signature"]["name"]})\n{p1["signature"]["position"]}'
    ws1["E18"].font = font_regular
    ws1["E18"].alignment = align_center

    ws1["D22"] = f'{p1["routing_footer"][0]["rank"]}          {p1["routing_footer"][0]["action"]}        {p1["routing_footer"][0]["date"]}'
    ws1["D22"].font = font_regular
    ws1["D23"] = f'{p1["routing_footer"][1]["rank"]} {p1["routing_footer"][1]["name"]}   {p1["routing_footer"][1]["action"]}  {p1["routing_footer"][1]["day"]}  {p1["routing_footer"][1]["date"]}'
    ws1["D23"].font = font_regular

    ws1.column_dimensions["A"].width = 5
    ws1.column_dimensions["B"].width = 18
    ws1.column_dimensions["C"].width = 50
    ws1.column_dimensions["D"].width = 30
    ws1.column_dimensions["E"].width = 25

    # --- SHEET 2: ผลการตรวจระดับเครื่องให้ทาง ---
    ws2 = wb.create_sheet(title="ผลการตรวจระดับเครื่องให้ทาง")
    ws2.views.sheetView[0].showGridLines = True

    # Header Row 2
    ws2["A2"] = p2["date_column"]
    ws2["A2"].font = font_bold
    ws2["A2"].alignment = align_center
    ws2["A2"].border = thin_border
    ws2.merge_cells("A2:A3")

    ws2["B2"] = "ผลการตรวจ"
    ws2["B2"].font = font_bold
    ws2["B2"].alignment = align_center
    ws2["B2"].border = thin_border
    ws2.merge_cells("B2:G2")

    ws2["H2"] = "ผู้ตรวจสอบ"
    ws2["H2"].font = font_bold
    ws2["H2"].alignment = align_center
    ws2["H2"].border = thin_border
    ws2.merge_cells("H2:H3")

    for col_idx, col_data in enumerate(p2["columns"], start=2):
        col_letter = get_column_letter(col_idx)
        cell = ws2[f"{col_letter}3"]
        cell.value = f'{col_data["name"]}\n{col_data["serial"]}'
        cell.font = font_regular
        cell.alignment = align_center
        cell.border = thin_border

    # Data Rows (4 to 7)
    for r_idx, r_data in enumerate(p2["rows"], start=4):
        ws2[f"A{r_idx}"] = r_data["item"]
        ws2[f"A{r_idx}"].font = font_bold
        ws2[f"A{r_idx}"].alignment = align_center
        ws2[f"A{r_idx}"].border = thin_border
        for c_idx, val in enumerate(r_data["values"], start=2):
            col_letter = get_column_letter(c_idx)
            cell = ws2[f"{col_letter}{r_idx}"]
            cell.value = val
            cell.font = font_regular
            cell.alignment = align_center
            cell.border = thin_border

    ws2["H4"] = p2["examiner"]
    ws2["H4"].font = font_bold
    ws2["H4"].alignment = align_center
    ws2["H4"].border = thin_border
    ws2.merge_cells("H4:H7")

    # Titles below
    ws2["B10"] = p2["title"]
    ws2["B10"].font = font_bold
    ws2["B10"].alignment = align_center
    ws2.merge_cells("B10:G10")

    ws2["B11"] = p2["subtitle"]
    ws2["B11"].font = font_bold
    ws2["B11"].alignment = align_center
    ws2.merge_cells("B11:G11")

    ws2["B12"] = p2["period"]
    ws2["B12"].font = font_bold
    ws2["B12"].alignment = align_center
    ws2.merge_cells("B12:G12")

    ws2.column_dimensions["A"].width = 14
    for c in ["B", "C", "D", "E", "F", "G"]:
        ws2.column_dimensions[c].width = 16
    ws2.column_dimensions["H"].width = 18

    # --- SHEET 3: ตารางทดสอบอาการถอย ---
    ws3 = wb.create_sheet(title="ตารางทดสอบอาการถอย")
    ws3.views.sheetView[0].showGridLines = True

    ws3["A1"] = p3["title_authority"]
    ws3["A1"].font = font_bold
    ws3["A1"].alignment = align_center
    ws3.merge_cells("A1:J1")

    ws3["A2"] = p3["title_main"]
    ws3["A2"].font = font_bold
    ws3["A2"].alignment = align_center
    ws3.merge_cells("A2:J2")

    ws3["A3"] = p3["unit"]
    ws3["A3"].font = font_bold
    ws3["A3"].alignment = align_center
    ws3.merge_cells("A3:J3")

    # Headers
    ws3["A5"] = "ครั้งที่"
    ws3["A5"].font = font_bold
    ws3["A5"].alignment = align_center
    ws3["A5"].border = thin_border
    ws3.merge_cells("A5:A6")

    ws3["B5"] = "รายการ"
    ws3["B5"].font = font_bold
    ws3["B5"].alignment = align_center
    ws3["B5"].border = thin_border
    ws3.merge_cells("B5:B6")

    ws3["C5"] = "ผู้รับผิดชอบ"
    ws3["C5"].font = font_bold
    ws3["C5"].alignment = align_center
    ws3["C5"].border = thin_border
    ws3.merge_cells("C5:J5")

    for col_idx, c_info in enumerate(p3["columns"][2:], start=3):
        col_letter = get_column_letter(col_idx)
        cell = ws3[f"{col_letter}6"]
        if "serial" in c_info:
            cell.value = f'{c_info["title"]}\nหมายเลข\n{c_info["serial"]}'
        else:
            cell.value = c_info["title"]
        cell.font = font_regular
        cell.alignment = align_center
        cell.border = thin_border

    # Row 7 Results
    ws3["A7"] = p3["row_1_result"]["seq"]
    ws3["A7"].font = font_bold
    ws3["A7"].alignment = align_center
    ws3["A7"].border = thin_border

    ws3["B7"] = f'{p3["row_1_result"]["date_label"]}\n{p3["row_1_result"]["result_label"]}'
    ws3["B7"].font = font_regular
    ws3["B7"].alignment = align_center
    ws3["B7"].border = thin_border

    for col_idx, res in enumerate(p3["row_1_result"]["results"], start=3):
        col_letter = get_column_letter(col_idx)
        cell = ws3[f"{col_letter}7"]
        cell.value = f'{res["date"]}\n{res["status"]}'
        cell.font = font_red_bold if res["is_repair"] else font_regular
        cell.alignment = align_center
        cell.border = thin_border

    ws3["I7"].border = thin_border
    ws3["J7"].border = thin_border
    ws3.row_dimensions[7].height = 45

    # Row 8 Officers
    ws3["A8"].border = thin_border
    ws3["B8"] = p3["row_2_officers"]["role"]
    ws3["B8"].font = font_bold
    ws3["B8"].alignment = align_center
    ws3["B8"].border = thin_border

    for col_idx, off in enumerate(p3["row_2_officers"]["officers"], start=3):
        col_letter = get_column_letter(col_idx)
        cell = ws3[f"{col_letter}8"]
        cell.value = f'{off["rank"]}\n({off["name"]})'
        cell.font = font_regular
        cell.alignment = align_center
        cell.border = thin_border

    ws3["I8"] = f'{p3["row_2_officers"]["coy_commander"]["rank"]}\n({p3["row_2_officers"]["coy_commander"]["name"]})'
    ws3["I8"].font = font_regular
    ws3["I8"].alignment = align_center
    ws3["I8"].border = thin_border

    ws3["J8"] = f'{p3["row_2_officers"]["bn_commander"]["rank"]}\n({p3["row_2_officers"]["bn_commander"]["name"]})'
    ws3["J8"].font = font_regular
    ws3["J8"].alignment = align_center
    ws3["J8"].border = thin_border
    ws3.row_dimensions[8].height = 55

    # Signature Block
    ws3["E11"] = f'{p3["signature"]["certified"]}\n\n{p3["signature"]["rank"]}\n({p3["signature"]["name"]})\n{p3["signature"]["position"]}\n{p3["signature"]["date"]}'
    ws3["E11"].font = font_regular
    ws3["E11"].alignment = align_center
    ws3.merge_cells("E11:F11")
    ws3.row_dimensions[11].height = 100

    ws3.column_dimensions["A"].width = 8
    ws3.column_dimensions["B"].width = 18
    for c in ["C", "D", "E", "F", "G", "H"]:
        ws3.column_dimensions[c].width = 16
    ws3.column_dimensions["I"].width = 14
    ws3.column_dimensions["J"].width = 14

    wb.save(output_path)
    print(f"Generated XLSX: {output_path}")


# -----------------------------------------------------------------------------
# MAIN COMPILER EXECUTION
# -----------------------------------------------------------------------------
if __name__ == "__main__":
    html_file = os.path.join(os.path.dirname(__file__), "รายงานผลการตรวจสอบ_ไตรมาส_Rebuild_Thai.html")
    docx_file = os.path.join(os.path.dirname(__file__), "รายงานผลการตรวจสอบ_ไตรมาส_Rebuild_Thai.docx")
    xlsx_file = os.path.join(os.path.dirname(__file__), "รายงานผลการตรวจสอบ_ไตรมาส_Rebuild_Thai.xlsx")

    generate_html(html_file)
    generate_docx(docx_file)
    generate_xlsx(xlsx_file)

    print("\nAll formats compiled successfully!")
