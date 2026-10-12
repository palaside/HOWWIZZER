# -*- coding: utf-8 -*-
"""
Superclone Production-Grade Engine for ทบ.๔๖๘-๒๐๔ (REF-TB-468-204-V1)
Form: บัญชีหมายเลขอาวุธ (คส. ๔)
Specifications:
- 6-Page A4 Portrait (210mm x 297mm)
- Single Unified Master Grid: 3 Columns
  - Col 1: หมายเลข (width: 20%)
  - Col 2: รายการ (width: 65%)
  - Col 3: หมายเหตุ (width: 15%)
- Spacing: Table top separated by 1 Enter from Title
- Certification block:
  - Line 1: หมดรายการรวม [จำนวน] [หน่วย] ขอรับรองว่า
  - Line 2: [รายการ] มีความยาวไม่เกินบรรทัดที่ 1
  - Line 3: จำนวน [จำนวน] [หน่วย] ตามหมายเลข
  - Line 4: อยู่ครบถ้วน ถูกต้อง ทุกหมายเลข
  - All 4 lines centered in Column 2 (รายการ)
- Signer: Universal Red Line Standard
  - Rank endpoint '.' aligns exactly with '(' of name
  - Name and position centered together
  - Signer: ร.อ. (กฤติน  นราพงษ์) ผบ.ป.๔ พัน.๔ ร้อย.๒
"""

import os, sys, json
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

WORKDIR = r"d:\Project\WEPAPP"

# Helper function for Thai digits
def to_th(val):
    if val is None or val == "":
        return ""
    s = str(val)
    mapping = {
        '0': '๐', '1': '๑', '2': '๒', '3': '๓', '4': '๔',
        '5': '๕', '6': '๖', '7': '๗', '8': '๘', '9': '๙'
    }
    return "".join(mapping.get(c, c) for c in s)

# Blueprint Data Model (SSOT)
blueprint = {
    "metadata": {
        "form_id": "ทบ. ๔๖๘ – ๒๐๔",
        "sub_id": "( คส. ๔ )",
        "form_name": "บัญชีหมายเลขอาวุธ",
        "unit": "ผบ.ป.๔ พัน.๔ ร้อย.๒",
        "signer": {
            "rank": "ร.อ.",
            "name": "(กฤติน  นราพงษ์)",
            "position": "ผบ.ป.๔ พัน.๔ ร้อย.๒"
        },
        "columns": [
            {"id": "num", "label": "หมายเลข", "width_pct": 20.0},
            {"id": "desc", "label": "รายการ", "width_pct": 65.0},
            {"id": "remark", "label": "หมายเหตุ", "width_pct": 15.0}
        ]
    },
    "pages": [
        {
            "page_num": 1,
            "category_name": "ปบค.๙๕ ขนาด ๑๐๕ มม. เอ็ม. ๑๐๑ เอ. ๑",
            "short_name": "ปบค.๙๕ ขนาด ๑๐๕ มม. แบบ เอ็ม. ๑๐๑ เอ. ๑",
            "count": 6,
            "count_th": "๖",
            "unit": "กระบอก",
            "items": [
                {"num": "๒๔๐๘๐", "name": "ปบค.๙๕ ขนาด ๑๐๕ มม. เอ็ม. ๑๐๑ เอ. ๑", "remark": "(ส่งซ่อม)"},
                {"num": "๑๙๖๐๙", "name": "ปบค.๙๕ ขนาด ๑๐๕ มม. เอ็ม. ๑๐๑ เอ. ๑", "remark": ""},
                {"num": "๒๐๖๗๐", "name": "ปบค.๙๕ ขนาด ๑๐๕ มม. เอ็ม. ๑๐๑ เอ. ๑", "remark": ""},
                {"num": "๒๓๕๕๔", "name": "ปบค.๙๕ ขนาด ๑๐๕ มม. เอ็ม. ๑๐๑ เอ. ๑", "remark": ""},
                {"num": "๒๔๒๘๐", "name": "ปบค.๙๕ ขนาด ๑๐๕ มม. เอ็ม. ๑๐๑ เอ. ๑", "remark": ""},
                {"num": "๘๕๑๖", "name": "ปบค.๙๕ ขนาด ๑๐๕ มม. เอ็ม. ๑๐๑ เอ. ๑", "remark": ""}
            ]
        },
        {
            "page_num": 2,
            "category_name": "กล้องเล็งพาโนรามิค",
            "short_name": "กล้องพาโนรามิค",
            "count": 9,
            "count_th": "๙",
            "unit": "กล้อง",
            "items": [
                {"num": "๐๒๙๔", "name": "กล้องเล็งพาโนรามิค", "remark": ""},
                {"num": "๐๒๙๖", "name": "กล้องเล็งพาโนรามิค", "remark": ""},
                {"num": "๐๒๙๘", "name": "กล้องเล็งพาโนรามิค", "remark": ""},
                {"num": "๐๒๙๙", "name": "กล้องเล็งพาโนรามิค", "remark": ""},
                {"num": "๒๔๑๖๕", "name": "กล้องเล็งพาโนรามิค", "remark": ""},
                {"num": "๕๗๗๕๖", "name": "กล้องเล็งพาโนรามิค", "remark": ""},
                {"num": "๔๓๕", "name": "กล้องเล็งพาโนรามิค", "remark": "อะไหล่"},
                {"num": "๔๓๑", "name": "กล้องเล็งพาโนรามิค", "remark": "อะไหล่"},
                {"num": "๔๙๒๙๕", "name": "กล้องเล็งพาโนรามิค", "remark": "อะไหล่"}
            ]
        },
        {
            "page_num": 3,
            "category_name": "กล้องข้อศอก",
            "short_name": "กล้องข้อศอก",
            "count": 7,
            "count_th": "๗",
            "unit": "กล้อง",
            "items": [
                {"num": "๑๓๘๒๗", "name": "กล้องข้อศอก", "remark": ""},
                {"num": "๒๖๕๓", "name": "กล้องข้อศอก", "remark": ""},
                {"num": "๓๕๘๙๗", "name": "กล้องข้อศอก", "remark": ""},
                {"num": "๓๖๐๐๘", "name": "กล้องข้อศอก", "remark": ""},
                {"num": "๓๖๔๗๔", "name": "กล้องข้อศอก", "remark": ""},
                {"num": "๒๓๓๒", "name": "กล้องข้อศอก", "remark": ""},
                {"num": "๓๖๐๗๓", "name": "กล้องข้อศอก", "remark": "อะไหล่"}
            ]
        },
        {
            "page_num": 4,
            "category_name": "เครื่องตั้งมุมยิงประณีต",
            "short_name": "เครื่องตั้งมุมยิงประณีต",
            "count": 6,
            "count_th": "๖",
            "unit": "เครื่อง",
            "items": [
                {"num": "๑๗๙๑๐๒", "name": "เครื่องตั้งมุมยิงประณีต", "remark": ""},
                {"num": "๓๐๑๙๒", "name": "เครื่องตั้งมุมยิงประณีต", "remark": ""},
                {"num": "๒๘๑๘๙", "name": "เครื่องตั้งมุมยิงประณีต", "remark": ""},
                {"num": "๙๖๑๕๗", "name": "เครื่องตั้งมุมยิงประณีต", "remark": ""},
                {"num": "๑๔๗๘๓๐", "name": "เครื่องตั้งมุมยิงประณีต", "remark": ""},
                {"num": "๑๔๗๓๔๙", "name": "เครื่องตั้งมุมยิงประณีต", "remark": ""}
            ]
        },
        {
            "page_num": 5,
            "category_name": "ฐานกล้องเล็ง",
            "short_name": "ฐานกล้องเล็ง",
            "count": 9,
            "count_th": "๙",
            "unit": "อัน",
            "items": [
                {"num": "๐๒๙๑", "name": "ฐานกล้องเล็ง", "remark": ""},
                {"num": "๐๒๘๓", "name": "ฐานกล้องเล็ง", "remark": ""},
                {"num": "0290", "name": "ฐานกล้องเล็ง", "remark": ""},
                {"num": "๐๔๒๔", "name": "ฐานกล้องเล็ง", "remark": ""},
                {"num": "๐๔๓๕", "name": "ฐานกล้องเล็ง", "remark": ""},
                {"num": "๐๔๓๑", "name": "ฐานกล้องเล็ง", "remark": ""},
                {"num": "๑๒๐๓๓", "name": "ฐานกล้องเล็ง", "remark": "อะไหล่"},
                {"num": "๒๗๐๒๖", "name": "ฐานกล้องเล็ง", "remark": "อะไหล่"},
                {"num": "๗๑๑๕", "name": "ฐานกล้องเล็ง", "remark": "อะไหล่"}
            ]
        },
        {
            "page_num": 6,
            "category_name": "คอลลิเมเตอร์",
            "short_name": "คอลลิเมเตอร์",
            "count": 5,
            "count_th": "๕",
            "unit": "เครื่อง",
            "items": [
                {"num": "๙๑๔๘", "name": "คอลลิเมเตอร์", "remark": ""},
                {"num": "๙๗๗๔", "name": "คอลลิเมเตอร์", "remark": ""},
                {"num": "๙๓๕๐", "name": "คอลลิเมเตอร์", "remark": ""},
                {"num": "๑๘๘๒", "name": "คอลลิเมเตอร์", "remark": ""},
                {"num": "๑๘๘๓", "name": "คอลลิเมเตอร์", "remark": ""}
            ]
        }
    ]
}

# 1. Save SSOT Blueprint JSON
json_path = os.path.join(WORKDIR, "table_layout_468_204.json")
with open(json_path, "w", encoding="utf-8") as f:
    json.dump(blueprint, f, ensure_ascii=False, indent=2)
print("Saved SSOT Blueprint:", json_path)


# ----------------------------------------------------------------------
# 2. HTML Generation (A4 Portrait, 3 Columns: 20%, 65%, 15%)
# ----------------------------------------------------------------------
def generate_html():
    pages_html = ""
    for p in blueprint["pages"]:
        rows_tr = ""
        for it in p["items"]:
            rows_tr += f"""
            <tr class="data-row">
                <td class="col-num text-center">{it['num']}</td>
                <td class="col-desc text-center">{it['name']}</td>
                <td class="col-remark text-center">{it['remark']}</td>
            </tr>"""

        # Certification block (4 lines centered in Column 2)
        # Line 1: หมดรายการรวม [จำนวน] [หน่วย] ขอรับรองว่า
        # Line 2: [รายการ]
        # Line 3: จำนวน [จำนวน] [หน่วย] ตามหมายเลข
        # Line 4: อยู่ครบถ้วน ถูกต้อง ทุกหมายเลข
        cert_block_html = f"""
        <div class="cert-block text-center">
            <div class="cert-line">หมดรายการรวม {p['count_th']} {p['unit']} ขอรับรองว่า</div>
            <div class="cert-line cert-title">{p['short_name']}</div>
            <div class="cert-line">จำนวน {p['count_th']} {p['unit']} ตามหมายเลข</div>
            <div class="cert-line">อยู่ครบถ้วน ถูกต้อง ทุกหมายเลข</div>
        </div>
        """

        # Universal Signature Block
        sig_block_html = f"""
        <div class="sig-section">
            <div class="sig-name-pos-group">
                <div class="sig-name-wrapper">
                    <span class="sig-rank-anchor val">{blueprint['metadata']['signer']['rank']}</span>
                    <div class="sig-name val">{blueprint['metadata']['signer']['name']}</div>
                </div>
                <div class="sig-pos val">{blueprint['metadata']['signer']['position']}</div>
            </div>
        </div>
        """

        pages_html += f"""
        <div class="sheet-container">
            <!-- หัวมุมบนขวา: ทบ. ๔๖๘ – ๒๐๔ / ( คส. ๔ ) -->
            <div class="top-form-header">
                <div>{blueprint['metadata']['form_id']}</div>
                <div>{blueprint['metadata']['sub_id']}</div>
            </div>

            <!-- ชื่อแบบฟอร์ม: บัญชีหมายเลขอาวุธ -->
            <div class="form-title font-bold text-center">
                {blueprint['metadata']['form_name']}
            </div>

            <!-- ตารางหลัก 3 คอลัมน์: ห่าง 1 Enter จากชื่อแบบฟอร์ม -->
            <table class="military-weapon-table">
                <colgroup>
                    <col style="width: 20%;">
                    <col style="width: 65%;">
                    <col style="width: 15%;">
                </colgroup>
                <thead>
                    <tr>
                        <th class="col-num">หมายเลข</th>
                        <th class="col-desc">รายการ</th>
                        <th class="col-remark">หมายเหตุ</th>
                    </tr>
                </thead>
                <tbody>
                    {rows_tr}
                    <!-- พื้นที่เนื้อหาด้านล่าง: คำรับรอง + ลายเซ็น -->
                    <tr class="cert-sig-row">
                        <td class="col-num"></td>
                        <td class="col-desc cert-cell">
                            {cert_block_html}
                            {sig_block_html}
                        </td>
                        <td class="col-remark"></td>
                    </tr>
                </tbody>
            </table>
        </div>
        """

    full_html = f"""<!DOCTYPE html>
<html lang="th">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>ทบ. ๔๖๘-๒๐๔ บัญชีหมายเลขอาวุธ (คส. ๔)</title>
    <link href="https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap" rel="stylesheet">
    <style>
        @page {{
            size: A4 portrait;
            margin: 15mm 15mm;
        }}
        * {{
            box-sizing: border-box;
            font-family: 'TH Sarabun PSK', Sarabun, sans-serif;
        }}
        body {{
            margin: 0;
            padding: 20px;
            background: #e2e8f0;
            color: #000;
            font-size: 16pt;
            line-height: 1.15;
            -webkit-print-color-adjust: exact;
            print-color-adjust: exact;
            font-feature-settings: "locl" 1, "numr" 1;
            font-variant-numeric: tabular-nums;
        }}

        .sheet-container {{
            width: 210mm;
            min-height: 297mm;
            margin: 0 auto 30px auto;
            background: #fff;
            padding: 18mm 18mm 20mm 18mm;
            box-shadow: 0 4px 12px rgba(0,0,0,0.15);
            position: relative;
            page-break-after: always;
        }}

        /* Top Right Header */
        .top-form-header {{
            text-align: right;
            font-size: 16pt;
            line-height: 1.25;
            margin-bottom: 12px;
            padding-right: 2px;
        }}

        /* Center Title */
        .form-title {{
            text-align: center;
            font-size: 16pt;
            font-weight: bold;
            margin-top: 6px;
            margin-bottom: 24px; /* เว้น 1 Enter อย่างชัดเจน (ประมาณ 20-24pt) */
        }}

        /* Main 3-Column Table */
        .military-weapon-table {{
            width: 100%;
            table-layout: fixed;
            border-collapse: collapse;
            border: 1px solid #000;
            font-size: 16pt;
        }}
        .military-weapon-table th, .military-weapon-table td {{
            border-left: 1px solid #000;
            border-right: 1px solid #000;
            padding: 4px 8px;
            vertical-align: top;
            line-height: 1.2;
        }}
        .military-weapon-table thead th {{
            border-top: 1px solid #000;
            border-bottom: 1px solid #000;
            font-weight: normal; /* ตัวปกติ ตามมาตรฐาน */
            font-size: 16pt;
            text-align: center;
            padding: 6px 4px;
        }}

        /* Data Rows */
        .data-row td {{
            border-bottom: none;
            height: 26px;
        }}

        /* Empty space row for Cert & Sig */
        .cert-sig-row td {{
            border-bottom: 1px solid #000;
            min-height: 480px;
            padding-bottom: 30px;
        }}

        /* Certification Block */
        .cert-cell {{
            padding-top: 36px !important;
        }}
        .cert-block {{
            margin: 0 auto;
            text-align: center;
            line-height: 1.35;
            font-size: 16pt;
        }}
        .cert-line {{
            margin: 0;
            text-align: center;
        }}
        .cert-title {{
            white-space: nowrap;
        }}

        /* Universal Signature Engine (RULE-UNIVERSAL-SIGNATURE-BLOCK-V1) */
        .sig-section {{
            margin-top: 36px;
            display: flex;
            justify-content: center;
            align-items: center;
        }}
        .sig-name-pos-group {{
            display: inline-flex;
            flex-direction: column;
            align-items: center;
            position: relative;
            margin-left: 45px;
        }}
        .sig-name-wrapper {{
            position: relative;
            display: inline-block;
            white-space: nowrap;
            margin-top: 24px;
        }}
        .sig-rank-anchor {{
            position: absolute;
            right: 100%;
            bottom: calc(100% + 8px);
            white-space: nowrap;
            font-size: 16pt;
            line-height: 1.0;
        }}
        .sig-name {{
            margin: 0;
            line-height: 1.15;
            text-align: center;
            white-space: nowrap;
            font-size: 16pt;
            display: block;
        }}
        .sig-pos {{
            margin: 0;
            margin-top: 0; /* Space After 0 pt / 1 Enter */
            line-height: 1.15;
            text-align: center;
            white-space: nowrap;
            font-size: 16pt;
        }}

        .font-bold {{ font-weight: bold; }}
        .text-center {{ text-align: center; }}
        .text-left {{ text-align: left; }}
        .text-right {{ text-align: right; }}

        @media print {{
            body {{ background: transparent; padding: 0; }}
            .sheet-container {{
                box-shadow: none;
                margin: 0;
                padding: 10mm 12mm;
                width: 100%;
                min-height: auto;
            }}
        }}
    </style>
</head>
<body>
    {pages_html}
</body>
</html>
    """

    html_path = os.path.join(WORKDIR, "ทบ_468_204_Rebuild_Thai.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(full_html)
    print("Saved HTML:", html_path)
    return html_path

generate_html()


# ----------------------------------------------------------------------
# 3. Microsoft Excel Generation (.xlsx)
# ----------------------------------------------------------------------
def generate_excel():
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet

    font_regular = Font(name="TH Sarabun PSK", size=16, bold=False)
    font_bold = Font(name="TH Sarabun PSK", size=16, bold=True)
    thin_side = Side(border_style="thin", color="000000")
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    for p in blueprint["pages"]:
        ws = wb.create_sheet(title=f"แผ่นที่ {p['page_num']}")
        ws.views.sheetView[0].showGridLines = True
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_margins.left = 0.6
        ws.page_margins.right = 0.6
        ws.page_margins.top = 0.6
        ws.page_margins.bottom = 0.6

        # Column widths: Col A (หมายเลข): 18, Col B (รายการ): 58, Col C (หมายเหตุ): 14
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 58
        ws.column_dimensions['C'].width = 14

        # Row 1 & 2: Top Form ID
        ws["C1"] = blueprint['metadata']['form_id']
        ws["C1"].font = font_regular; ws["C1"].alignment = align_right
        ws["C2"] = blueprint['metadata']['sub_id']
        ws["C2"].font = font_regular; ws["C2"].alignment = align_right

        # Row 3: Blank
        # Row 4: Title
        ws["B4"] = blueprint['metadata']['form_name']
        ws["B4"].font = font_bold; ws["B4"].alignment = align_center

        # Row 5: 1 Enter Gap
        # Row 6: Header row
        ws["A6"] = "หมายเลข"
        ws["B6"] = "รายการ"
        ws["C6"] = "หมายเหตุ"
        for col in ['A', 'B', 'C']:
            c = ws[f"{col}6"]
            c.font = font_regular; c.alignment = align_center
            c.border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
            c.fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")

        # Data rows
        curr_r = 7
        for it in p["items"]:
            c_a = ws[f"A{curr_r}"]; c_a.value = it['num']; c_a.font = font_regular; c_a.alignment = align_center
            c_b = ws[f"B{curr_r}"]; c_b.value = it['name']; c_b.font = font_regular; c_b.alignment = align_center
            c_c = ws[f"C{curr_r}"]; c_c.value = it['remark']; c_c.font = font_regular; c_c.alignment = align_center
            for col in ['A', 'B', 'C']:
                ws[f"{col}{curr_r}"].border = Border(left=thin_side, right=thin_side)
            curr_r += 1

        # Spacing rows inside table
        for _ in range(2):
            for col in ['A', 'B', 'C']:
                ws[f"{col}{curr_r}"].border = Border(left=thin_side, right=thin_side)
            curr_r += 1

        # Cert block (4 rows centered in Col B)
        cert_lines = [
            f"หมดรายการรวม {p['count_th']} {p['unit']} ขอรับรองว่า",
            p['short_name'],
            f"จำนวน {p['count_th']} {p['unit']} ตามหมายเลข",
            "อยู่ครบถ้วน ถูกต้อง ทุกหมายเลข"
        ]
        for cl in cert_lines:
            ws[f"B{curr_r}"].value = cl
            ws[f"B{curr_r}"].font = font_regular
            ws[f"B{curr_r}"].alignment = align_center
            for col in ['A', 'B', 'C']:
                ws[f"{col}{curr_r}"].border = Border(left=thin_side, right=thin_side)
            curr_r += 1

        # Space before signature
        for _ in range(2):
            for col in ['A', 'B', 'C']:
                ws[f"{col}{curr_r}"].border = Border(left=thin_side, right=thin_side)
            curr_r += 1

        # Signature rows
        sig_lines = [
            f"      {blueprint['metadata']['signer']['rank']}",
            f"      {blueprint['metadata']['signer']['name']}",
            f"      {blueprint['metadata']['signer']['position']}"
        ]
        for sl in sig_lines:
            ws[f"B{curr_r}"].value = sl
            ws[f"B{curr_r}"].font = font_regular
            ws[f"B{curr_r}"].alignment = align_center
            for col in ['A', 'B', 'C']:
                ws[f"{col}{curr_r}"].border = Border(left=thin_side, right=thin_side)
            curr_r += 1

        # Extra spacing to bottom of page
        for _ in range(8):
            for col in ['A', 'B', 'C']:
                ws[f"{col}{curr_r}"].border = Border(left=thin_side, right=thin_side)
            curr_r += 1

        # Bottom border row
        for col in ['A', 'B', 'C']:
            ws[f"{col}{curr_r}"].border = Border(bottom=thin_side, left=thin_side, right=thin_side)

    xlsx_path = os.path.join(WORKDIR, "ทบ_468_204_Rebuild_Thai.xlsx")
    wb.save(xlsx_path)
    print("Saved XLSX:", xlsx_path)

generate_excel()


# ----------------------------------------------------------------------
# 4. Microsoft Word Generation (.docx)
# ----------------------------------------------------------------------
def format_docx_run(run, font_name="TH Sarabun PSK", size_pt=16, bold=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}" w:eastAsia="{font_name}"/>')
    rPr.append(rFonts)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top="none", bottom="none", left="single", right="single", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="000000"/><w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="000000"/><w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="000000"/><w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="000000"/></w:tcBorders>')
    tcPr.append(borders)

def generate_word():
    doc = docx.Document()
    
    # Configure A4 Portrait (8.27" x 11.69")
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)

    for p_idx, p in enumerate(blueprint["pages"]):
        if p_idx > 0:
            doc.add_page_break()

        # Top Form ID Right
        p_hdr = doc.add_paragraph()
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.paragraph_format.space_before = Pt(0)
        p_hdr.paragraph_format.space_after = Pt(0)
        p_hdr.paragraph_format.line_spacing = 1.15
        r_fid = p_hdr.add_run(f"{blueprint['metadata']['form_id']}\n{blueprint['metadata']['sub_id']}")
        format_docx_run(r_fid, size_pt=16, bold=False)

        # Title Center
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(6)
        p_title.paragraph_format.space_after = Pt(24) # 1 Enter gap
        r_t = p_title.add_run(blueprint['metadata']['form_name'])
        format_docx_run(r_t, size_pt=16, bold=True)

        # Table 3 columns: total width = 6.87 inches (8.27 - 1.4)
        # Col 1: 20% = 1.37 in, Col 2: 65% = 4.47 in, Col 3: 15% = 1.03 in
        tbl_widths = [Inches(1.37), Inches(4.47), Inches(1.03)]
        
        num_rows = 1 + len(p['items']) + 1 # header + items + cert/sig row
        table = doc.add_table(rows=num_rows, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Set column widths & header
        headers = ["หมายเลข", "รายการ", "หมายเหตุ"]
        for c_idx in range(3):
            cell = table.cell(0, c_idx)
            cell.width = tbl_widths[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            set_cell_borders(cell, top="single", bottom="single", left="single", right="single")
            p_c = cell.paragraphs[0]; p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_c.paragraph_format.space_after = Pt(0)
            r_c = p_c.add_run(headers[c_idx])
            format_docx_run(r_c, size_pt=16, bold=False)

        # Data rows
        for it_idx, it in enumerate(p['items']):
            row_idx = 1 + it_idx
            # Col 0
            c0 = table.cell(row_idx, 0)
            c0.width = tbl_widths[0]
            set_cell_margins(c0, top=40, bottom=40, left=80, right=80)
            set_cell_borders(c0, top="none", bottom="none", left="single", right="single")
            p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(it['num'])
            format_docx_run(r0, size_pt=16, bold=False)

            # Col 1
            c1 = table.cell(row_idx, 1)
            c1.width = tbl_widths[1]
            set_cell_margins(c1, top=40, bottom=40, left=80, right=80)
            set_cell_borders(c1, top="none", bottom="none", left="single", right="single")
            p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(it['name'])
            format_docx_run(r1, size_pt=16, bold=False)

            # Col 2
            c2 = table.cell(row_idx, 2)
            c2.width = tbl_widths[2]
            set_cell_margins(c2, top=40, bottom=40, left=80, right=80)
            set_cell_borders(c2, top="none", bottom="none", left="single", right="single")
            p2 = c2.paragraphs[0]; p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(it['remark'])
            format_docx_run(r2, size_pt=16, bold=False)

        # Bottom Cert & Sig Row
        cert_row_idx = num_rows - 1
        c_b0 = table.cell(cert_row_idx, 0); c_b0.width = tbl_widths[0]
        set_cell_margins(c_b0, top=100, bottom=400, left=80, right=80)
        set_cell_borders(c_b0, top="none", bottom="single", left="single", right="single")

        c_b2 = table.cell(cert_row_idx, 2); c_b2.width = tbl_widths[2]
        set_cell_margins(c_b2, top=100, bottom=400, left=80, right=80)
        set_cell_borders(c_b2, top="none", bottom="single", left="single", right="single")

        # Cell 1 (Cert & Sig)
        c_b1 = table.cell(cert_row_idx, 1); c_b1.width = tbl_widths[1]
        set_cell_margins(c_b1, top=240, bottom=400, left=80, right=80)
        set_cell_borders(c_b1, top="none", bottom="single", left="single", right="single")
        
        # Cert paragraphs
        cert_lines = [
            f"หมดรายการรวม {p['count_th']} {p['unit']} ขอรับรองว่า",
            p['short_name'],
            f"จำนวน {p['count_th']} {p['unit']} ตามหมายเลข",
            "อยู่ครบถ้วน ถูกต้อง ทุกหมายเลข"
        ]
        for l_idx, cl in enumerate(cert_lines):
            p_cl = c_b1.paragraphs[0] if l_idx == 0 else c_b1.add_paragraph()
            p_cl.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cl.paragraph_format.space_after = Pt(0)
            p_cl.paragraph_format.line_spacing = 1.2
            r_cl = p_cl.add_run(cl)
            format_docx_run(r_cl, size_pt=16, bold=False)

        # Space before sig
        p_sp = c_b1.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(20)

        # Universal Signature Block in DOCX
        axis_tab = Inches(2.23) # Center axis of 4.47 inches
        p_rk = c_b1.add_paragraph()
        p_rk.paragraph_format.tab_stops.add_tab_stop(axis_tab, WD_TAB_ALIGNMENT.RIGHT)
        p_rk.paragraph_format.space_after = Pt(14)
        r_rk = p_rk.add_run(f"\t{blueprint['metadata']['signer']['rank']}")
        format_docx_run(r_rk, size_pt=16, bold=False)

        p_nm = c_b1.add_paragraph()
        p_nm.paragraph_format.tab_stops.add_tab_stop(axis_tab, WD_TAB_ALIGNMENT.LEFT)
        p_nm.paragraph_format.space_after = Pt(0)
        r_nm = p_nm.add_run(f"\t{blueprint['metadata']['signer']['name']}")
        format_docx_run(r_nm, size_pt=16, bold=False)

        p_pos = c_b1.add_paragraph()
        p_pos.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_pos.paragraph_format.space_after = Pt(0)
        r_pos = p_pos.add_run(blueprint['metadata']['signer']['position'])
        format_docx_run(r_pos, size_pt=16, bold=False)

    docx_path = os.path.join(WORKDIR, "ทบ_468_204_Rebuild_Thai.docx")
    doc.save(docx_path)
    print("Saved DOCX:", docx_path)

generate_word()
