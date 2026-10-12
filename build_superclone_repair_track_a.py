#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SUPERCLONE Production Engine
Document: ทบ.๔๐๐-๐๗๑ แบบ ก ใบติดตามใบส่งซ่อม (การเบิกเกิน ๔๕ - ๙๐ วัน)
Authority: ระเบียบกองทัพบกว่าด้วยการส่งกำลังสิ่งอุปกรณ์ ประเภท ๒ และ ๔ พ.ศ. ๒๕๓๔
Unit: ป.๔ พัน.๔ ร้อย.๒
Formats: HTML (A4 Landscape Single Page), DOCX (Single Page Landscape), XLSX (Single Sheet)
Standard: RULE-PROD-DOC-ENGINE-V1 / RULE-UNIVERSAL-SIGNATURE-BLOCK-V1 (Red Line Standard)
Ref Model: REF-MEMO-INQUIRY-45-V1
"""

import os
import sys
import json

if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from signature_engine import (
    render_signature_html,
    render_signature_docx,
    render_signature_xlsx,
    is_signature_block
)

# Load SSOT Blueprint
BLUEPRINT_PATH = os.path.join(os.path.dirname(__file__), "table_layout_repair_track_a.json")
with open(BLUEPRINT_PATH, "r", encoding="utf-8") as f:
    bp = json.load(f)

p1_data = bp["page_1_form"]


# ==============================================================================
# 1. HTML GENERATOR (Pixel-Perfect A4 Landscape Single-Page Rebuild)
# ==============================================================================
def generate_html():
    out_html_path = os.path.join(os.path.dirname(__file__), "ใบติดตามใบส่งซ่อมแบบ_ก_Rebuild_Thai.html")
    
    # Signatures
    left_sig_data = p1_data["signatures"][0]
    right_sig_data = p1_data["signatures"][1]
    
    left_sig_html = render_signature_html({
        "rank": left_sig_data["rank"],
        "name": left_sig_data["name"],
        "position": left_sig_data["position"],
        "date": left_sig_data.get("date", left_sig_data.get("date_text", ""))
    }, margin_left_px=45, gap_pt=20)
    
    right_sig_html = render_signature_html({
        "label": right_sig_data.get("label", "ผู้ตรวจสอบ"),
        "rank": right_sig_data["rank"],
        "name": right_sig_data["name"],
        "position": right_sig_data["position"],
        "date": right_sig_data.get("date", right_sig_data.get("date_text", ""))
    }, margin_left_px=45, gap_pt=20)
    
    # Table rows
    table_rows_html = ""
    for r in p1_data["table"]["rows"]:
        table_rows_html += f"""
        <tr>
            <td style="text-align: center; vertical-align: middle;">{r["seq"]}</td>
            <td style="text-align: left; vertical-align: middle; padding-left: 3mm; line-height: 1.25;">{r["item_desc"]}</td>
            <td style="text-align: center; vertical-align: middle;">{r["item_number"]}</td>
            <td style="text-align: center; vertical-align: middle;">{r["doc_ref"]}</td>
            <td style="text-align: center; vertical-align: middle;">{r["doc_date"]}</td>
            <td style="text-align: left; vertical-align: middle;">{r["notes"]}</td>
        </tr>
        """
        
    blank_rows_count = p1_data["table"].get("blank_rows_count", 9)
    for _ in range(blank_rows_count):
        table_rows_html += """
        <tr>
            <td>&nbsp;</td>
            <td>&nbsp;</td>
            <td>&nbsp;</td>
            <td>&nbsp;</td>
            <td>&nbsp;</td>
            <td>&nbsp;</td>
        </tr>
        """
        
    html_content = f"""<!DOCTYPE html>
<html lang="th">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{bp["title"]}</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Sarabun:ital,wght@0,100;0,200;0,300;0,400;0,600;0,700;1,400&display=swap" rel="stylesheet">
    <style>
        @page landscape-page {{
            size: A4 landscape;
            margin: 0;
        }}
        * {{
            box-sizing: border-box;
            -webkit-print-color-adjust: exact;
            print-color-adjust: exact;
        }}
        body {{
            margin: 0;
            padding: 20px 0;
            background-color: #525659;
            font-family: 'Sarabun', 'TH Sarabun PSK', 'THSarabunNew', sans-serif;
            font-size: 12pt;
            font-weight: 400;
            font-feature-settings: "locl" 1, "numr" 1;
            font-variant-numeric: tabular-nums;
            color: #000000;
            line-height: 1.35;
            display: flex;
            flex-direction: column;
            align-items: center;
        }}
        .page-landscape {{
            width: 297mm;
            min-height: 210mm;
            height: 210mm;
            padding: 8mm 12mm 6mm 15mm;
            background-color: #ffffff;
            margin-bottom: 20px;
            box-shadow: 0 4px 10px rgba(0, 0, 0, 0.3);
            position: relative;
            page: landscape-page;
            overflow: hidden;
            display: flex;
            flex-direction: column;
            justify-content: space-between;
        }}
        @media print {{
            body {{
                background-color: transparent;
                padding: 0;
            }}
            .page-landscape {{
                margin-bottom: 0;
                box-shadow: none;
                page-break-after: always;
                break-after: page;
            }}
        }}

        /* Header Layout */
        .form-top-header {{
            display: flex;
            justify-content: space-between;
            align-items: flex-start;
            margin-bottom: 1.5mm;
        }}
        .doc-title {{
            font-size: 18pt;
            font-weight: bold;
            text-align: center;
            margin-left: 50mm;
        }}
        .header-top-right {{
            text-align: right;
            font-size: 12pt;
            line-height: 1.2;
        }}
        .form-code {{
            font-weight: bold;
        }}
        .form-subcode {{
            margin-top: 1mm;
        }}

        /* Meta block Left */
        .meta-block {{
            margin-bottom: 2mm;
            font-size: 12pt;
            line-height: 1.45;
        }}
        .meta-line {{
            margin-bottom: 0.5mm;
        }}
        .meta-label {{
            font-weight: bold;
            display: inline-block;
            width: 12mm;
        }}
        .instruction-line {{
            padding-left: 20mm;
            margin-top: 1.5mm;
            margin-bottom: 2mm;
            font-size: 12.5pt;
        }}

        /* Grid Table (6 columns) */
        .grid-table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 11.5pt;
            margin-bottom: 2mm;
        }}
        .grid-table th, .grid-table td {{
            border: 1px solid #000000;
            padding: 1.2mm 1.5mm;
            height: 6.8mm;
        }}
        .grid-table th {{
            text-align: center;
            font-weight: bold;
            background-color: #ffffff;
            vertical-align: middle;
            line-height: 1.2;
        }}

        /* Universal Signature Engine Styles (Red Line Standard) */
        .sig-block {{
            position: relative;
            display: inline-block;
            font-family: inherit;
        }}
        .sig-verify {{
            font-weight: bold;
            font-size: 12pt;
            text-align: left;
            margin-bottom: 20pt;
        }}
        .sig-name-pos-group {{
            position: relative;
            display: flex;
            flex-direction: column;
            align-items: center;
        }}
        .sig-name-wrapper {{
            position: relative;
            display: inline-block;
            white-space: nowrap;
            margin-top: 20pt;
        }}
        .sig-rank-anchor {{
            position: absolute;
            right: 100%;
            bottom: calc(100% + 18px);
            white-space: nowrap;
            font-size: 12pt;
            line-height: 1.0;
            margin-right: 2px;
        }}
        .sig-name {{
            display: inline-block;
            font-weight: normal;
            font-size: 12pt;
            line-height: 1.15;
            white-space: nowrap;
        }}
        .sig-pos {{
            text-align: center;
            white-space: pre-line;
            line-height: 1.2;
            margin-top: 2px;
            font-size: 12pt;
        }}
        .sig-date {{
            text-align: center;
            margin-top: 3px;
            font-size: 12pt;
            white-space: nowrap;
        }}

        .double-sig-row {{
            display: flex;
            justify-content: space-between;
            align-items: flex-start;
            margin-top: 3mm;
            padding: 0 15mm;
        }}
        .sig-col-left {{
            width: 45%;
            display: flex;
            justify-content: center;
        }}
        .sig-col-right {{
            width: 45%;
            display: flex;
            justify-content: center;
        }}

        /* Footer Note */
        .footer-note {{
            text-align: center;
            margin-top: 2mm;
            margin-bottom: 1mm;
            font-size: 10.5pt;
        }}
    </style>
</head>
<body>

    <!-- ==================== PAGE: ใบติดตามใบส่งซ่อม (A4 Landscape 1 หน้าจบ) ==================== -->
    <div class="page-landscape">
        <div>
            <!-- Form Top Header -->
            <div class="form-top-header">
                <div></div>
                <div class="doc-title">{p1_data["title"]}</div>
                <div class="header-top-right">
                    <div class="form-code">{p1_data["form_code"]}</div>
                    <div class="form-subcode">{p1_data["form_type"]}</div>
                </div>
            </div>

            <!-- Meta block -->
            <div class="meta-block">
                <div class="meta-line">{p1_data["header_fields"]["doc_number"]}</div>
                <div class="meta-line">{p1_data["header_fields"]["subject"]}</div>
                <div class="meta-line">
                    <span class="meta-label">ถึง</span>
                    <span>{p1_data["header_fields"]["to"]}</span>
                </div>
                <div class="meta-line">
                    <span class="meta-label">จาก</span>
                    <span>{p1_data["header_fields"]["from"]}</span>
                </div>
                <div class="instruction-line">
                    {p1_data["instruction"]}
                </div>
            </div>

            <!-- Table (6 columns) -->
            <table class="grid-table">
                <thead>
                    <tr>
                        <th style="width: 5%;">ลำดับ</th>
                        <th style="width: 21%;">รายการ สป.</th>
                        <th style="width: 17%;">หมายเลข สป.</th>
                        <th style="width: 18%; line-height: 1.15; padding: 1mm 1mm;"><span style="white-space: nowrap;">ที่ใบส่งซ่อม ที่ทะเบียนหน่วย</span><br><span>จ่าย</span></th>
                        <th style="width: 18%; white-space: nowrap; padding: 1mm 1mm;">วันที่ส่งหรือวางใบส่งซ่อม</th>
                        <th style="width: 21%; white-space: nowrap; padding: 1mm 1mm;">หมายเหตุ  (บันทึกโดยหน่วยจ่าย)</th>
                    </tr>
                </thead>
                <tbody>
                    {table_rows_html}
                </tbody>
            </table>
        </div>

        <div>
            <!-- Double Signature Block (Universal Red Line Standard) -->
            <div class="double-sig-row">
                <div class="sig-col-left">
                    {left_sig_html}
                </div>
                <div class="sig-col-right">
                    {right_sig_html}
                </div>
            </div>

            <!-- Footer note (ด้านล่างในหน้าเดียวกันเหมือนรูปที่ 1) -->
            <div class="footer-note">
                {p1_data["footer_note"]}
            </div>
        </div>
    </div>

</body>
</html>
"""
    with open(out_html_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"[HTML] Built: {out_html_path}")
    return out_html_path


# ==============================================================================
# 2. DOCX GENERATOR (Pixel-Perfect 1-Page Landscape Document)
# ==============================================================================
def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', '000000'))
            tcBorders.append(element)
        else:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), 'none')
            tcBorders.append(element)
    tcPr.append(tcBorders)

def generate_docx():
    out_docx_path = os.path.join(os.path.dirname(__file__), "ใบติดตามใบส่งซ่อมแบบ_ก_Rebuild_Thai.docx")
    doc = Document()
    
    # Section Setup: A4 Landscape (1 Page)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11.69)  # 297mm
    sec.page_height = Inches(8.27)  # 210mm
    sec.top_margin = Inches(0.40)
    sec.bottom_margin = Inches(0.35)
    sec.left_margin = Inches(0.60)
    sec.right_margin = Inches(0.50)
    
    font_name = "TH Sarabun PSK"
    
    def add_p(text="", size=13, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=0):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if text:
            r = p.add_run(text)
            r.font.name = font_name
            r.font.size = Pt(size)
            r.font.bold = bold
        return p

    # Top Header Table (Title Center, Form Code Right)
    t_hdr = doc.add_table(rows=2, cols=2)
    t_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hdr.autofit = False
    
    t_hdr.columns[0].width = Inches(7.5)
    t_hdr.columns[1].width = Inches(3.0)
    
    # Row 0
    c_title = t_hdr.cell(0, 0)
    p_t = c_title.paragraphs[0]
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(0)
    p_t.paragraph_format.space_after = Pt(0)
    r_t = p_t.add_run(p1_data["title"])
    r_t.font.name = font_name
    r_t.font.size = Pt(18)
    r_t.font.bold = True
    
    c_code = t_hdr.cell(0, 1)
    p_c = c_code.paragraphs[0]
    p_c.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_c.paragraph_format.space_before = Pt(0)
    p_c.paragraph_format.space_after = Pt(0)
    r_c = p_c.add_run(p1_data["form_code"])
    r_c.font.name = font_name
    r_c.font.size = Pt(13)
    r_c.font.bold = True
    
    # Row 1
    c_subcode = t_hdr.cell(1, 1)
    p_sc = c_subcode.paragraphs[0]
    p_sc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sc.paragraph_format.space_before = Pt(0)
    p_sc.paragraph_format.space_after = Pt(0)
    r_sc = p_sc.add_run(p1_data["form_type"])
    r_sc.font.name = font_name
    r_sc.font.size = Pt(12)
    
    # Remove borders from t_hdr
    for row in t_hdr.rows:
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell, top=0, bottom=0, left=0, right=0)
            
    # Field Rows Left
    add_p(p1_data["header_fields"]["doc_number"], size=12, space_before=2, space_after=1)
    add_p(p1_data["header_fields"]["subject"], size=12, space_before=1, space_after=1)
    
    # ถึง & จาก
    p_to = add_p("", size=12, space_before=1, space_after=1)
    r_to_lbl = p_to.add_run("ถึง    ")
    r_to_lbl.font.name = font_name
    r_to_lbl.font.size = Pt(12)
    r_to_lbl.font.bold = True
    r_to_val = p_to.add_run(p1_data["header_fields"]["to"])
    r_to_val.font.name = font_name
    r_to_val.font.size = Pt(12)
    
    p_from = add_p("", size=12, space_before=1, space_after=2)
    r_from_lbl = p_from.add_run("จาก  ")
    r_from_lbl.font.name = font_name
    r_from_lbl.font.size = Pt(12)
    r_from_lbl.font.bold = True
    r_from_val = p_from.add_run(p1_data["header_fields"]["from"])
    r_from_val.font.name = font_name
    r_from_val.font.size = Pt(12)
    
    # Instruction Text with indent
    p_inst = add_p(f"        {p1_data['instruction']}", size=12.5, space_before=2, space_after=4)
    
    # Main Table (6 cols: 5%, 21%, 17%, 18%, 18%, 21%)
    col_widths = [Inches(0.53), Inches(2.25), Inches(1.80), Inches(1.95), Inches(1.95), Inches(2.25)]
    total_rows = 1 + len(p1_data["table"]["rows"]) + p1_data["table"].get("blank_rows_count", 9)
    t_main = doc.add_table(rows=total_rows, cols=6)
    t_main.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_main.autofit = False
    
    # Headers
    headers = [
        "ลำดับ",
        "รายการ สป.",
        "หมายเลข สป.",
        "ที่ใบส่งซ่อม ที่ทะเบียนหน่วย\nจ่าย",
        "วันที่ส่งหรือวางใบส่งซ่อม",
        "หมายเหตุ  (บันทึกโดยหน่วยจ่าย)"
    ]
    
    hdr_row = t_main.rows[0]
    for idx, (title, w) in enumerate(zip(headers, col_widths)):
        cell = hdr_row.cells[idx]
        cell.width = w
        set_cell_border(cell, top={'sz': 4}, bottom={'sz': 4}, left={'sz': 4}, right={'sz': 4})
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(title)
        r.font.name = font_name
        r.font.size = Pt(11.5)
        r.font.bold = True
        
    # Data Row (Row 1)
    data_row = t_main.rows[1]
    row_data = p1_data["table"]["rows"][0]
    row_values = [
        (row_data["seq"], WD_ALIGN_PARAGRAPH.CENTER),
        (row_data["item_desc"].replace("<br>", "\n"), WD_ALIGN_PARAGRAPH.LEFT),
        (row_data["item_number"], WD_ALIGN_PARAGRAPH.CENTER),
        (row_data["doc_ref"], WD_ALIGN_PARAGRAPH.CENTER),
        (row_data["doc_date"], WD_ALIGN_PARAGRAPH.CENTER),
        (row_data["notes"], WD_ALIGN_PARAGRAPH.LEFT)
    ]
    
    for idx, ((val, align), w) in enumerate(zip(row_values, col_widths)):
        cell = data_row.cells[idx]
        cell.width = w
        set_cell_border(cell, top={'sz': 4}, bottom={'sz': 4}, left={'sz': 4}, right={'sz': 4})
        set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(val)
        r.font.name = font_name
        r.font.size = Pt(11.5)
        
    # Blank Rows (Row 2 to total_rows - 1)
    for r_idx in range(2, total_rows):
        b_row = t_main.rows[r_idx]
        for c_idx, w in enumerate(col_widths):
            cell = b_row.cells[c_idx]
            cell.width = w
            set_cell_border(cell, top={'sz': 4}, bottom={'sz': 4}, left={'sz': 4}, right={'sz': 4})
            set_cell_margins(cell, top=40, bottom=40, left=40, right=40)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run("")
            r.font.name = font_name
            r.font.size = Pt(11.5)
            
    # Spacer before signatures
    add_p("", size=6, space_before=0, space_after=4)
    
    # Double Signature Block Table (Left & Right)
    t_sig = doc.add_table(rows=1, cols=2)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sig.autofit = False
    t_sig.columns[0].width = Inches(5.2)
    t_sig.columns[1].width = Inches(5.2)
    
    for row in t_sig.rows:
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell, top=0, bottom=0, left=0, right=0)
            
    # Render Left Sig (Commander)
    c_left = t_sig.cell(0, 0)
    left_sig_data = p1_data["signatures"][0]
    render_signature_docx(
        c_left,
        {
            "rank": left_sig_data["rank"],
            "name": left_sig_data["name"],
            "position": left_sig_data["position"],
            "date": left_sig_data.get("date", left_sig_data.get("date_text", ""))
        },
        axis_inch=2.0,
        font_name=font_name,
        font_size=12,
        space_after_rank_pt=22
    )
    
    # Render Right Sig (Inspector)
    c_right = t_sig.cell(0, 1)
    right_sig_data = p1_data["signatures"][1]
    render_signature_docx(
        c_right,
        {
            "label": right_sig_data.get("label", "ผู้ตรวจสอบ"),
            "rank": right_sig_data["rank"],
            "name": right_sig_data["name"],
            "position": right_sig_data["position"],
            "date": right_sig_data.get("date", right_sig_data.get("date_text", ""))
        },
        axis_inch=2.2,
        font_name=font_name,
        font_size=12,
        space_after_rank_pt=22
    )

    # Footer note on same page
    p_ft = add_p(p1_data["footer_note"], size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=0)
    
    doc.save(out_docx_path)
    print(f"[DOCX] Built: {out_docx_path}")
    return out_docx_path


# ==============================================================================
# 3. XLSX GENERATOR (Pixel-Perfect Single-Sheet Document)
# ==============================================================================
def generate_xlsx():
    out_xlsx_path = os.path.join(os.path.dirname(__file__), "ใบติดตามใบส่งซ่อมแบบ_ก_Rebuild_Thai.xlsx")
    wb = openpyxl.Workbook()
    
    # Sheet 1: ใบติดตามใบส่งซ่อม
    ws1 = wb.active
    ws1.title = "ใบติดตามใบส่งซ่อม"
    ws1.views.sheetView[0].showGridLines = True
    
    # Page setup
    ws1.page_setup.orientation = ws1.ORIENTATION_LANDSCAPE
    ws1.page_setup.paperSize = ws1.PAPERSIZE_A4
    ws1.page_margins.left = 0.5
    ws1.page_margins.right = 0.4
    ws1.page_margins.top = 0.4
    ws1.page_margins.bottom = 0.3
    
    font_name = "TH Sarabun PSK"
    f_title = Font(name=font_name, size=16, bold=True)
    f_header = Font(name=font_name, size=12, bold=True)
    f_normal = Font(name=font_name, size=12)
    f_footer = Font(name=font_name, size=10.5)
    
    thin_border_side = Side(border_style="thin", color="000000")
    thin_box = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    # Columns: A (Seq 5%), B (Desc 21%), C (Num 17%), D (Ref 18%), E (Date 18%), F (Notes 21%)
    col_widths = {
        "A": 7,
        "B": 25,
        "C": 20,
        "D": 23,
        "E": 23,
        "F": 27
    }
    for col, width in col_widths.items():
        ws1.column_dimensions[col].width = width
        
    # Row 1: Header Top Right (Form Code) & Title Center
    ws1["A1"] = p1_data["title"]
    ws1["A1"].font = f_title
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.merge_cells("A1:E1")
    
    ws1["F1"] = p1_data["form_code"]
    ws1["F1"].font = f_header
    ws1["F1"].alignment = Alignment(horizontal="right", vertical="center")
    
    ws1["F2"] = p1_data["form_type"]
    ws1["F2"].font = f_normal
    ws1["F2"].alignment = Alignment(horizontal="right", vertical="center")
    
    # Fields Left (Rows 3 - 6)
    ws1["A3"] = p1_data["header_fields"]["doc_number"]
    ws1["A3"].font = f_normal
    ws1.merge_cells("A3:D3")
    
    ws1["A4"] = p1_data["header_fields"]["subject"]
    ws1["A4"].font = f_normal
    ws1.merge_cells("A4:D4")
    
    ws1["A5"] = f"ถึง    {p1_data['header_fields']['to']}"
    ws1["A5"].font = f_normal
    ws1.merge_cells("A5:D5")
    
    ws1["A6"] = f"จาก  {p1_data['header_fields']['from']}"
    ws1["A6"].font = f_normal
    ws1.merge_cells("A6:D6")
        
    # Instruction Text (Row 7)
    ws1["A7"] = f"        {p1_data['instruction']}"
    ws1["A7"].font = f_normal
    ws1["A7"].alignment = Alignment(horizontal="left", vertical="center")
    ws1.merge_cells("A7:F7")
    
    # Main Table Headers (Row 9)
    headers = [
        ("A", "ลำดับ"),
        ("B", "รายการ สป."),
        ("C", "หมายเลข สป."),
        ("D", "ที่ใบส่งซ่อม ที่ทะเบียนหน่วย\nจ่าย"),
        ("E", "วันที่ส่งหรือวางใบส่งซ่อม"),
        ("F", "หมายเหตุ  (บันทึกโดยหน่วยจ่าย)")
    ]
    ws1.row_dimensions[9].height = 32
    for col, txt in headers:
        cell = ws1[f"{col}9"]
        cell.value = txt
        cell.font = f_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_box
        
    # Data Row (Row 10)
    ws1.row_dimensions[10].height = 42
    r_data = p1_data["table"]["rows"][0]
    data_cells = [
        ("A", r_data["seq"], Alignment(horizontal="center", vertical="center")),
        ("B", r_data["item_desc"].replace("<br>", "\n"), Alignment(horizontal="left", vertical="center", wrap_text=True)),
        ("C", r_data["item_number"], Alignment(horizontal="center", vertical="center")),
        ("D", r_data["doc_ref"], Alignment(horizontal="center", vertical="center")),
        ("E", r_data["doc_date"], Alignment(horizontal="center", vertical="center")),
        ("F", r_data["notes"], Alignment(horizontal="left", vertical="center"))
    ]
    for col, val, align in data_cells:
        cell = ws1[f"{col}10"]
        cell.value = val
        cell.font = f_normal
        cell.alignment = align
        cell.border = thin_box
        
    # Blank rows (Row 11 to 19)
    for row_idx in range(11, 20):
        ws1.row_dimensions[row_idx].height = 20
        for col, _ in headers:
            cell = ws1[f"{col}{row_idx}"]
            cell.value = ""
            cell.border = thin_box
            
    # Signatures at bottom (Row 21+)
    # Left Signature: Cols A & B (ร.อ. กฤติน นราพงษ์)
    left_sig_data = p1_data["signatures"][0]
    render_signature_xlsx(
        ws1,
        {
            "rank": left_sig_data["rank"],
            "name": left_sig_data["name"],
            "position": left_sig_data["position"],
            "date": left_sig_data.get("date", left_sig_data.get("date_text", ""))
        },
        start_row=22,
        rank_col="A",
        name_col="B",
        font_name=font_name,
        font_size=12
    )
    
    # Right Signature: Cols E & F (ผู้ตรวจสอบ จ.ส.อ. ณัฐชัย รักษาวงษ์)
    right_sig_data = p1_data["signatures"][1]
    render_signature_xlsx(
        ws1,
        {
            "label": right_sig_data.get("label", "ผู้ตรวจสอบ"),
            "rank": right_sig_data["rank"],
            "name": right_sig_data["name"],
            "position": right_sig_data["position"],
            "date": right_sig_data.get("date", right_sig_data.get("date_text", ""))
        },
        start_row=21,
        rank_col="E",
        name_col="F",
        font_name=font_name,
        font_size=12
    )
    
    # Footer note (Row 27)
    ws1["A27"] = p1_data["footer_note"]
    ws1["A27"].font = f_footer
    ws1["A27"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.merge_cells("A27:F27")
    
    wb.save(out_xlsx_path)
    print(f"[XLSX] Built: {out_xlsx_path}")
    return out_xlsx_path


if __name__ == "__main__":
    generate_html()
    generate_docx()
    generate_xlsx()
    print("All SUPERCLONE documents generated successfully!")
