"""
Clone Informal Equipment Inventory from Excel using SUPERCLONE standard.
Extracts:
- Section 1: สิ่งอุปกรณ์ในโรงเก็บปืนใหญ่ (ลำดับที่ 1 - 43)
- Section 2: สิ่งอุปกรณ์ในคลัง บก.โรงเก็บปืนใหญ่ (ลำดับที่ 1 - 29)
Saves to:
- backend/data/informal_equipment_inventory.json
- references/informal_equipment/informal_equipment_blueprint.json
"""
import openpyxl
import json
import os
from pathlib import Path

WORKSPACE_ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = WORKSPACE_ROOT / "บัญชีรายการและสิ่งอุปกรณ์.xlsx"
DATA_OUT = WORKSPACE_ROOT / "backend" / "data" / "informal_equipment_inventory.json"
BLUEPRINT_OUT = WORKSPACE_ROOT / "references" / "informal_equipment" / "informal_equipment_blueprint.json"
BLUEPRINT_OUT.parent.mkdir(parents=True, exist_ok=True)
DATA_OUT.parent.mkdir(parents=True, exist_ok=True)

def parse_excel():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb['Sheet1']

    section_1_items = []
    section_2_items = []

    # Section 1: Rows 5 to 63
    current_parent_1 = ""
    current_seq_1 = 0
    sub_counter_1 = 0

    for r in range(5, 64):
        seq = ws.cell(r, 1).value
        # If header repeated inside
        if str(seq).strip() == 'ลำดับ':
            continue

        raw_name = ws.cell(r, 2).value
        raw_name2 = ws.cell(r, 3).value
        raw_name3 = ws.cell(r, 4).value

        # Determine item name
        name_parts = [str(x).strip() for x in [raw_name, raw_name2, raw_name3] if x is not None and str(x).strip() != '']
        if not name_parts:
            continue
        
        is_sub = False
        display_name = ""
        if name_parts[0] in ['-', '- ']:
            is_sub = True
            display_name = " ".join(name_parts[1:])
        elif name_parts[0].startswith('-'):
            is_sub = True
            display_name = name_parts[0].lstrip('-').strip() + (" " + " ".join(name_parts[1:]) if len(name_parts) > 1 else "")
        else:
            display_name = " ".join(name_parts)

        ctrl = ws.cell(r, 6).value or 0
        curr = ws.cell(r, 7).value or 0
        short = ws.cell(r, 8).value or 0
        surp = ws.cell(r, 9).value or 0
        repair = ws.cell(r, 10).value or 0
        loan = ws.cell(r, 11).value or 0
        bal = ws.cell(r, 12).value if ws.cell(r, 12).value is not None else curr
        remarks = ws.cell(r, 13).value or ""

        if seq is not None and str(seq).strip() != '':
            try:
                current_seq_1 = int(seq)
            except:
                pass
            current_parent_1 = display_name
            sub_counter_1 = 0
            item_id = f"gs-{current_seq_1:02d}"
        else:
            sub_counter_1 += 1
            item_id = f"gs-{current_seq_1:02d}-{sub_counter_1}"

        item_obj = {
            "id": item_id,
            "section": "gun_shed",
            "section_title": "สิ่งอุปกรณ์ในโรงเก็บปืนใหญ่",
            "seq": current_seq_1 if (seq is not None and str(seq).strip() != '') else "",
            "display_seq": str(current_seq_1) if (seq is not None and str(seq).strip() != '') else f"- {display_name}",
            "name": display_name,
            "full_name": f"{current_parent_1} ({display_name})" if is_sub else display_name,
            "parent_name": current_parent_1,
            "is_sub": is_sub,
            "control_amount": int(ctrl) if isinstance(ctrl, (int, float)) else 0,
            "physical_amount": int(curr) if isinstance(curr, (int, float)) else 0,
            "deficit_amount": int(short) if isinstance(short, (int, float)) else 0,
            "surplus_amount": int(surp) if isinstance(surp, (int, float)) else 0,
            "repair_amount": int(repair) if isinstance(repair, (int, float)) else 0,
            "loan_amount": int(loan) if isinstance(loan, (int, float)) else 0,
            "balance_amount": int(bal) if isinstance(bal, (int, float)) else (int(curr) if isinstance(curr, (int, float)) else 0),
            "remarks": str(remarks).strip() if remarks else ""
        }
        section_1_items.append(item_obj)

    # Section 2: Rows 72 to 115
    current_parent_2 = ""
    current_seq_2 = 0
    sub_counter_2 = 0

    for r in range(72, 116):
        seq = ws.cell(r, 1).value
        if str(seq).strip() == 'ลำดับ':
            continue

        raw_name = ws.cell(r, 2).value
        raw_name2 = ws.cell(r, 3).value
        raw_name3 = ws.cell(r, 4).value

        name_parts = [str(x).strip() for x in [raw_name, raw_name2, raw_name3] if x is not None and str(x).strip() != '']
        if not name_parts:
            continue

        is_sub = False
        display_name = ""
        if name_parts[0] in ['-', '- ']:
            is_sub = True
            display_name = " ".join(name_parts[1:])
        elif name_parts[0].startswith('-'):
            is_sub = True
            display_name = name_parts[0].lstrip('-').strip() + (" " + " ".join(name_parts[1:]) if len(name_parts) > 1 else "")
        else:
            display_name = " ".join(name_parts)

        ctrl = ws.cell(r, 6).value or 0
        curr = ws.cell(r, 7).value or 0
        short = ws.cell(r, 8).value or 0
        surp = ws.cell(r, 9).value or 0
        repair = ws.cell(r, 10).value or 0
        loan = ws.cell(r, 11).value or 0
        bal = ws.cell(r, 12).value if ws.cell(r, 12).value is not None else curr
        remarks = ws.cell(r, 13).value or ""

        if seq is not None and str(seq).strip() != '':
            try:
                current_seq_2 = int(seq)
            except:
                pass
            current_parent_2 = display_name
            sub_counter_2 = 0
            item_id = f"hq-{current_seq_2:02d}"
        else:
            sub_counter_2 += 1
            item_id = f"hq-{current_seq_2:02d}-{sub_counter_2}"

        item_obj = {
            "id": item_id,
            "section": "hq_warehouse",
            "section_title": "สิ่งอุปกรณ์ในคลัง บก.โรงเก็บปืนใหญ่",
            "seq": current_seq_2 if (seq is not None and str(seq).strip() != '') else "",
            "display_seq": str(current_seq_2) if (seq is not None and str(seq).strip() != '') else f"- {display_name}",
            "name": display_name,
            "full_name": f"{current_parent_2} ({display_name})" if is_sub else display_name,
            "parent_name": current_parent_2,
            "is_sub": is_sub,
            "control_amount": int(ctrl) if isinstance(ctrl, (int, float)) else 0,
            "physical_amount": int(curr) if isinstance(curr, (int, float)) else 0,
            "deficit_amount": int(short) if isinstance(short, (int, float)) else 0,
            "surplus_amount": int(surp) if isinstance(surp, (int, float)) else 0,
            "repair_amount": int(repair) if isinstance(repair, (int, float)) else 0,
            "loan_amount": int(loan) if isinstance(loan, (int, float)) else 0,
            "balance_amount": int(bal) if isinstance(bal, (int, float)) else (int(curr) if isinstance(curr, (int, float)) else 0),
            "remarks": str(remarks).strip() if remarks else ""
        }
        section_2_items.append(item_obj)

    inventory_data = {
        "metadata": {
            "title": "บัญชีคุมยุทโธปกรณ์ ป.4 พัน.4 ร้อย.2 (การสำรวจแบบไม่เป็นทางการ)",
            "unit": "ป.๔ พัน.๔ ร้อย.๒",
            "source_file": "บัญชีรายการและสิ่งอุปกรณ์.xlsx",
            "status": "informal_survey",
            "is_official_catalog": False,
            "note": "การสำรวจเพื่อรายงาน ผบ.ร้อย ทราบ ไม่นับรวมกับบัญชีสิ่งอุปกรณ์ทางการเดิม",
            "total_section_1_items": len(section_1_items),
            "total_section_2_items": len(section_2_items)
        },
        "sections": {
            "gun_shed": {
                "id": "gun_shed",
                "title": "สิ่งอุปกรณ์ในโรงเก็บปืนใหญ่",
                "range_desc": "ลำดับที่ 1 - 43",
                "items": section_1_items
            },
            "hq_warehouse": {
                "id": "hq_warehouse",
                "title": "สิ่งอุปกรณ์ในคลัง บก.โรงเก็บปืนใหญ่",
                "range_desc": "ลำดับที่ 1 - 29",
                "items": section_2_items
            }
        },
        "loans": [
            {
                "loan_id": "LN-2569-001",
                "item_id": "gs-30",
                "item_name": "จอบขุด",
                "section": "gun_shed",
                "section_title": "สิ่งอุปกรณ์ในโรงเก็บปืนใหญ่",
                "borrower_rank": "ส.อ.",
                "borrower_name": "สมชาย หมายมั่น",
                "borrower_fullname": "ส.อ. สมชาย หมายมั่น",
                "borrower_position": "พลปืน ป.หมู่ 1",
                "quantity": 1,
                "loan_date": "2026-11-02",
                "expected_return_date": "2026-11-04",
                "actual_return_date": "2026-11-05",
                "status": "returned_damaged",
                "is_overdue": False,
                "damage_info": {
                    "has_damage": True,
                    "cause": "การขุดลอกคูระบายน้ำรอบโรงเก็บปืนใหญ่แล้วกระทบหินแข็งใต้ผิวดินอย่างรุนแรง",
                    "symptom": "ด้ามจอบหักสะบั้นและตัวใบจอบบิดงอผิดรูป ไม่สามารถใช้งานขุดต่อได้",
                    "reported_to_commander": False,
                    "memo_generated": False
                }
            }
        ]
    }

    # Write Data JSON
    with open(DATA_OUT, 'w', encoding='utf-8') as f:
        json.dump(inventory_data, f, ensure_ascii=False, indent=2)
    print(f"Saved Inventory Data to: {DATA_OUT}")

    # Write Blueprint JSON
    blueprint = {
        "blueprint_version": "SUPERCLONE-V1.0",
        "title": "Informal Equipment Survey Blueprint",
        "source": "บัญชีรายการและสิ่งอุปกรณ์.xlsx",
        "sections_count": 2,
        "section_1_count": len(section_1_items),
        "section_2_count": len(section_2_items),
        "columns": [
            {"id": "seq", "title": "ลำดับ", "width": "50px"},
            {"id": "name", "title": "รายการ", "width": "220px"},
            {"id": "control_amount", "title": "ยอดบัญชีคุม", "width": "80px"},
            {"id": "physical_amount", "title": "ยอดที่มี", "width": "80px"},
            {"id": "deficit_amount", "title": "ขาด", "width": "60px"},
            {"id": "surplus_amount", "title": "เกิน", "width": "60px"},
            {"id": "repair_amount", "title": "ส่งซ่อม", "width": "60px"},
            {"id": "loan_amount", "title": "เบิกยืม", "width": "60px"},
            {"id": "balance_amount", "title": "คงเหลือ", "width": "80px"},
            {"id": "remarks", "title": "หมายเหตุ", "width": "120px"}
        ]
    }
    with open(BLUEPRINT_OUT, 'w', encoding='utf-8') as f:
        json.dump(blueprint, f, ensure_ascii=False, indent=2)
    print(f"Saved Blueprint to: {BLUEPRINT_OUT}")

if __name__ == '__main__':
    parse_excel()
