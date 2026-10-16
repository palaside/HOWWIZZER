import sys
import os
import json
import shutil
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from signature_engine import render_signature_html, render_signature_docx

WORKSPACE = Path(__file__).resolve().parent.parent
DATA_JSON_PATH = WORKSPACE / "backend" / "data" / "informal_equipment_inventory.json"
OUT_HTML = WORKSPACE / "บัญชีรายการและสิ่งอุปกรณ์_Rebuild_Thai.html"
OUT_DOCX = WORKSPACE / "บัญชีรายการและสิ่งอุปกรณ์_Rebuild_Thai.docx"
OUT_XLSX = WORKSPACE / "บัญชีรายการและสิ่งอุปกรณ์.xlsx"

FRONTEND_DOCS = WORKSPACE / "frontend" / "static" / "docs"
STATIC_DOCS = WORKSPACE / "static" / "docs"

def to_thai_num(val):
    if val is None or val == "":
        return ""
    th_digits = "๐๑๒๓๔๕๖๗๘๙"
    return "".join(th_digits[int(c)] if c.isdigit() else c for c in str(val))

def load_data():
    with open(DATA_JSON_PATH, "r", encoding="utf-8") as f:
        return json.load(f)

def build_html(data):
    s1_items = data["sections"]["gun_shed"]["items"]
    s2_items = data["sections"]["hq_warehouse"]["items"]

    # Calculate sums
    def calc_sums(items):
        return {
            "ctrl": sum(i.get("control_amount", 0) for i in items if not i.get("is_sub")),
            "curr": sum(i.get("physical_amount", 0) for i in items if not i.get("is_sub")),
            "short": sum(i.get("deficit_amount", 0) for i in items if not i.get("is_sub")),
            "surp": sum(i.get("surplus_amount", 0) for i in items if not i.get("is_sub")),
            "repair": sum(i.get("repair_amount", 0) for i in items if not i.get("is_sub")),
            "loan": sum(i.get("loan_amount", 0) for i in items if not i.get("is_sub")),
            "bal": sum(i.get("balance_amount", 0) for i in items if not i.get("is_sub")),
        }

    sum1 = calc_sums(s1_items)
    sum2 = calc_sums(s2_items)

    def render_rows(items):
        rows = []
        for itm in items:
            is_sub = itm.get("is_sub", False)
            seq = to_thai_num(itm.get("seq", "")) if not is_sub else ""
            prefix = "&nbsp;&nbsp;&nbsp;&nbsp;- " if is_sub else ""
            name = prefix + itm.get("name", "")
            ctrl = to_thai_num(itm.get("control_amount", 0))
            curr = to_thai_num(itm.get("physical_amount", 0))
            short = to_thai_num(itm.get("deficit_amount", 0)) if itm.get("deficit_amount") else "-"
            surp = to_thai_num(itm.get("surplus_amount", 0)) if itm.get("surplus_amount") else "-"
            repair = to_thai_num(itm.get("repair_amount", 0)) if itm.get("repair_amount") else "-"
            loan = to_thai_num(itm.get("loan_amount", 0)) if itm.get("loan_amount") else "-"
            bal = to_thai_num(itm.get("balance_amount", 0))
            rem = itm.get("remarks", "")

            row_style = "background-color: #fcfcfc;" if is_sub else ""
            name_weight = "300" if is_sub else "600"
            rows.append(f"""
            <tr style="{row_style}">
                <td style="text-align: center; font-weight: {name_weight};">{seq}</td>
                <td style="text-align: left; font-weight: {name_weight};">{name}</td>
                <td style="text-align: center;">{ctrl}</td>
                <td style="text-align: center;">{curr}</td>
                <td style="text-align: center;">{short}</td>
                <td style="text-align: center;">{surp}</td>
                <td style="text-align: center;">{repair}</td>
                <td style="text-align: center;">{loan}</td>
                <td style="text-align: center; font-weight: 600;">{bal}</td>
                <td style="text-align: left;">{rem}</td>
            </tr>
            """)
        return "".join(rows)

    sig_block = {
        "title": "ตรวจถูกต้อง",
        "rank": "ร.อ.",
        "name": "ชยานันต์ คงสัตย์",
        "position": "ผบ.ร้อย.ป.๒",
        "alignment": "right",
        "space_before_pt": 24
    }
    sig_html = render_signature_html(sig_block)

    html_content = f"""<!DOCTYPE html>
<html lang="th">
<head>
    <meta charset="UTF-8">
    <title>บัญชีรายการและสิ่งอุปกรณ์ (โรงเก็บปืนใหญ่ และ คลัง บก.โรงเก็บปืนใหญ่)</title>
    <link href="https://fonts.googleapis.com/css2?family=Sarabun:ital,wght@0,300;0,400;0,600;0,700&display=swap" rel="stylesheet">
    <style>
        @page {{
            size: A4 landscape;
            margin: 15mm 15mm 15mm 15mm;
        }}
        * {{
            box-sizing: border-box;
            -webkit-print-color-adjust: exact;
            print-color-adjust: exact;
        }}
        body {{
            font-family: 'TH Sarabun PSK', 'THSarabunPSK', 'TH Sarabun New', 'Sarabun', sans-serif !important;
            margin: 0;
            padding: 20px;
            background-color: #525659;
            color: #000;
            display: flex;
            flex-direction: column;
            align-items: center;
        }}
        .sheet {{
            background: #fff;
            width: 297mm;
            min-height: 210mm;
            padding: 18mm 18mm 18mm 20mm;
            box-shadow: 0 4px 15px rgba(0,0,0,0.3);
            margin-bottom: 20px;
            position: relative;
        }}
        @media print {{
            body {{
                background: none;
                padding: 0;
            }}
            .sheet {{
                box-shadow: none;
                margin: 0;
                width: 100%;
                page-break-after: always;
            }}
        }}
        .header-title {{
            text-align: center;
            font-size: 20pt;
            font-weight: 600;
            margin-bottom: 4px;
        }}
        .header-sub {{
            text-align: center;
            font-size: 16pt;
            font-weight: 600;
            margin-bottom: 14px;
        }}
        .sec-title {{
            font-size: 16pt;
            font-weight: 600;
            margin-top: 14px;
            margin-bottom: 6px;
            background-color: #f1f5f9;
            padding: 4px 10px;
            border-left: 4px solid #0284c7;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 13.5pt;
            line-height: 1.25;
            margin-bottom: 12px;
        }}
        th, td {{
            border: 1px solid #475569;
            padding: 4px 6px;
        }}
        th {{
            background-color: #f8fafc;
            font-weight: 600;
            text-align: center;
        }}
        td {{
            font-weight: 300;
        }}
        .total-row td {{
            font-weight: 600;
            background-color: #f8fafc;
        }}
        .sig-container {{
            margin-top: 24px;
            display: flex;
            justify-content: flex-end;
            padding-right: 40px;
        }}
    </style>
</head>
<body>

    <div class="sheet">
        <div class="header-title">บัญชีรายการและสิ่งอุปกรณ์</div>
        <div class="header-sub">กองร้อยทหารปืนใหญ่ที่ ๒ กองพันทหารปืนใหญ่ที่ ๔ กรมทหารปืนใหญ่ที่ ๔ (ป.๔ พัน.๔ ร้อย.๒)</div>

        <div class="sec-title">หมวดที่ ๑ : สิ่งอุปกรณ์ในโรงเก็บปืนใหญ่ (ลำดับ ๑ - ๔๓)</div>
        <table>
            <thead>
                <tr>
                    <th style="width: 45px;">ลำดับ</th>
                    <th>รายการ</th>
                    <th style="width: 75px;">ยอดบัญชีคุม</th>
                    <th style="width: 75px;">ยอดที่มี</th>
                    <th style="width: 55px;">ขาด</th>
                    <th style="width: 55px;">เกิน</th>
                    <th style="width: 55px;">ส่งซ่อม</th>
                    <th style="width: 55px;">เบิกยืม</th>
                    <th style="width: 75px;">คงเหลือ</th>
                    <th style="width: 130px;">หมายเหตุ</th>
                </tr>
            </thead>
            <tbody>
                {render_rows(s1_items)}
                <tr class="total-row">
                    <td colspan="2" style="text-align: center;">รวมหมวดที่ ๑</td>
                    <td style="text-align: center;">{to_thai_num(sum1['ctrl'])}</td>
                    <td style="text-align: center;">{to_thai_num(sum1['curr'])}</td>
                    <td style="text-align: center;">{to_thai_num(sum1['short'])}</td>
                    <td style="text-align: center;">{to_thai_num(sum1['surp'])}</td>
                    <td style="text-align: center;">{to_thai_num(sum1['repair'])}</td>
                    <td style="text-align: center;">{to_thai_num(sum1['loan'])}</td>
                    <td style="text-align: center;">{to_thai_num(sum1['bal'])}</td>
                    <td></td>
                </tr>
            </tbody>
        </table>

        <div class="sec-title" style="margin-top: 24px;">หมวดที่ ๒ : สิ่งอุปกรณ์ในคลัง บก.โรงเก็บปืนใหญ่ (ลำดับ ๑ - ๒๙)</div>
        <table>
            <thead>
                <tr>
                    <th style="width: 45px;">ลำดับ</th>
                    <th>รายการ</th>
                    <th style="width: 75px;">ยอดบัญชีคุม</th>
                    <th style="width: 75px;">ยอดที่มี</th>
                    <th style="width: 55px;">ขาด</th>
                    <th style="width: 55px;">เกิน</th>
                    <th style="width: 55px;">ส่งซ่อม</th>
                    <th style="width: 55px;">เบิกยืม</th>
                    <th style="width: 75px;">คงเหลือ</th>
                    <th style="width: 130px;">หมายเหตุ</th>
                </tr>
            </thead>
            <tbody>
                {render_rows(s2_items)}
                <tr class="total-row">
                    <td colspan="2" style="text-align: center;">รวมหมวดที่ ๒</td>
                    <td style="text-align: center;">{to_thai_num(sum2['ctrl'])}</td>
                    <td style="text-align: center;">{to_thai_num(sum2['curr'])}</td>
                    <td style="text-align: center;">{to_thai_num(sum2['short'])}</td>
                    <td style="text-align: center;">{to_thai_num(sum2['surp'])}</td>
                    <td style="text-align: center;">{to_thai_num(sum2['repair'])}</td>
                    <td style="text-align: center;">{to_thai_num(sum2['loan'])}</td>
                    <td style="text-align: center;">{to_thai_num(sum2['bal'])}</td>
                    <td></td>
                </tr>
            </tbody>
        </table>

        <div class="sig-container">
            {sig_html}
        </div>
    </div>

</body>
</html>
"""
    with open(OUT_HTML, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"Generated HTML: {OUT_HTML}")

def build_docx(data):
    s1_items = data["sections"]["gun_shed"]["items"]
    s2_items = data["sections"]["hq_warehouse"]["items"]

    doc = Document()
    
    # Page setup: A4 Landscape
    sections = doc.sections
    for s in sections:
        s.page_width = Cm(29.7)
        s.page_height = Cm(21.0)
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.5)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(1.5)

    def set_font(run, name="TH Sarabun PSK", size_pt=14, bold=False):
        run.font.name = name
        run.font.size = Pt(size_pt)
        run.bold = bold
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{name}" w:hAnsi="{name}" w:cs="{name}"/>')
        rPr.append(rFonts)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    r = p_title.add_run("บัญชีรายการและสิ่งอุปกรณ์")
    set_font(r, size_pt=18, bold=True)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("กองร้อยทหารปืนใหญ่ที่ ๒ กองพันทหารปืนใหญ่ที่ ๔ กรมทหารปืนใหญ่ที่ ๔ (ป.๔ พัน.๔ ร้อย.๒)")
    set_font(r_sub, size_pt=15, bold=True)

    def add_table(title, items):
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(8)
        p_sec.paragraph_format.space_after = Pt(4)
        r_sec = p_sec.add_run(title)
        set_font(r_sec, size_pt=15, bold=True)

        col_widths = [Cm(1.2), Cm(6.5), Cm(2.0), Cm(2.0), Cm(1.5), Cm(1.5), Cm(1.5), Cm(1.5), Cm(2.0), Cm(4.5)]
        table = doc.add_table(rows=1, cols=10)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        headers = ["ลำดับ", "รายการ", "ยอดบัญชีคุม", "ยอดที่มี", "ขาด", "เกิน", "ส่งซ่อม", "เบิกยืม", "คงเหลือ", "หมายเหตุ"]
        hdr_cells = table.rows[0].cells
        for idx, text in enumerate(headers):
            hdr_cells[idx].text = text
            hdr_cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in hdr_cells[idx].paragraphs[0].runs:
                set_font(run, size_pt=13, bold=True)

        for itm in items:
            is_sub = itm.get("is_sub", False)
            seq = to_thai_num(itm.get("seq", "")) if not is_sub else ""
            prefix = "    - " if is_sub else ""
            name = prefix + itm.get("name", "")
            ctrl = to_thai_num(itm.get("control_amount", 0))
            curr = to_thai_num(itm.get("physical_amount", 0))
            short = to_thai_num(itm.get("deficit_amount", 0)) if itm.get("deficit_amount") else "-"
            surp = to_thai_num(itm.get("surplus_amount", 0)) if itm.get("surplus_amount") else "-"
            repair = to_thai_num(itm.get("repair_amount", 0)) if itm.get("repair_amount") else "-"
            loan = to_thai_num(itm.get("loan_amount", 0)) if itm.get("loan_amount") else "-"
            bal = to_thai_num(itm.get("balance_amount", 0))
            rem = itm.get("remarks", "")

            row_cells = table.add_row().cells
            row_data = [seq, name, ctrl, curr, short, surp, repair, loan, bal, rem]
            for idx, val in enumerate(row_data):
                row_cells[idx].text = val
                align = WD_ALIGN_PARAGRAPH.LEFT if idx in [1, 9] else WD_ALIGN_PARAGRAPH.CENTER
                row_cells[idx].paragraphs[0].alignment = align
                for run in row_cells[idx].paragraphs[0].runs:
                    set_font(run, size_pt=12.5, bold=(idx == 1 and not is_sub))

    add_table("หมวดที่ ๑ : สิ่งอุปกรณ์ในโรงเก็บปืนใหญ่ (ลำดับ ๑ - ๔๓)", s1_items)
    add_table("หมวดที่ ๒ : สิ่งอุปกรณ์ในคลัง บก.โรงเก็บปืนใหญ่ (ลำดับ ๑ - ๒๙)", s2_items)

    # Signature
    sig_block = {
        "title": "ตรวจถูกต้อง",
        "rank": "ร.อ.",
        "name": "ชยานันต์ คงสัตย์",
        "position": "ผบ.ร้อย.ป.๒",
        "tab_stop_mm": 200,
        "space_before_pt": 24
    }
    render_signature_docx(doc, sig_block)

    doc.save(OUT_DOCX)
    print(f"Generated DOCX: {OUT_DOCX}")

def main():
    data = load_data()
    build_html(data)
    build_docx(data)

    # Copy files to frontend/static/docs and static/docs
    for dest in [FRONTEND_DOCS, STATIC_DOCS]:
        dest.mkdir(parents=True, exist_ok=True)
        shutil.copy2(OUT_HTML, dest / OUT_HTML.name)
        shutil.copy2(OUT_DOCX, dest / OUT_DOCX.name)
        if OUT_XLSX.exists():
            shutil.copy2(OUT_XLSX, dest / OUT_XLSX.name)
            # also copy with Rebuild_Thai name for uniform access
            shutil.copy2(OUT_XLSX, dest / "บัญชีรายการและสิ่งอุปกรณ์_Rebuild_Thai.xlsx")

    # Also copy to root as Rebuild_Thai.xlsx for consistency
    if OUT_XLSX.exists():
        shutil.copy2(OUT_XLSX, WORKSPACE / "บัญชีรายการและสิ่งอุปกรณ์_Rebuild_Thai.xlsx")

    print("All documents generated and copied to static docs folders successfully!")

if __name__ == "__main__":
    main()
