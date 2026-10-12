import sys
import pathlib

PDF_PATH = pathlib.Path(r"d:/Project/WEPAPP/docs/thai-docx.skill/ระบบเอกสาร/400-012 ผบ.ร้อย.pdf")
OUTPUT_DOCX = pathlib.Path(r"d:/Project/WEPAPP/400-012_ผบ.ร้อย_converted.docx")

try:
    import pdfplumber
    from docx import Document
except ImportError as e:
    print(f"Missing package: {e}")
    sys.exit(1)

if not PDF_PATH.is_file():
    print(f"PDF not found: {PDF_PATH}")
    sys.exit(1)

doc = Document()

with pdfplumber.open(str(PDF_PATH)) as pdf:
    for page_number, page in enumerate(pdf.pages, start=1):
        tables = page.extract_tables()
        if not tables:
            continue
        doc.add_heading(f"Page {page_number}", level=2)
        for table in tables:
            # Write table rows as tab‑separated text (no visible grid)
            for row in table:
                line = "\t".join(cell if cell is not None else "" for cell in row)
                doc.add_paragraph(line)
            doc.add_paragraph()  # blank line between tables

# Save the Word document
OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUTPUT_DOCX))
print(f"Word document created: {OUTPUT_DOCX}")
