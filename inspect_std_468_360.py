import fitz
import docx
import os

pdf_std = 'docs/thai-docx.skill/โรงปืน/Ref_เอกสารโรงปืน/Ref/02_แบบฟอร์มมาตรฐาน_Standard_DOCX_and_PDF/ทบ_468_360_ปรนนิบัติบำรุง.pdf'
if os.path.exists(pdf_std):
    doc = fitz.open(pdf_std)
    print("Standard PDF total pages:", len(doc))
    for i, p in enumerate(doc):
        pix = p.get_pixmap(dpi=150)
        pix.save(f'C:/Users/EVE/.gemini/antigravity-ide/brain/7ca9cd9b-6ae7-44b7-ab50-280c021eef14/std_468_360_p{i+1}.png')
        print(f"Saved std page {i+1}, rect: {p.rect}")

docx_std = 'docs/thai-docx.skill/โรงปืน/Ref_เอกสารโรงปืน/Ref/02_แบบฟอร์มมาตรฐาน_Standard_DOCX_and_PDF/ทบ_468_360_ปรนนิบัติบำรุง.docx'
if os.path.exists(docx_std):
    d = docx.Document(docx_std)
    out_lines = []
    out_lines.append(f"Standard DOCX Paragraphs: {len(d.paragraphs)}, Tables: {len(d.tables)}")
    for j, p in enumerate(d.paragraphs):
        if p.text.strip():
            out_lines.append(f"P{j}: {p.text.strip()}")
    for t_idx, table in enumerate(d.tables):
        out_lines.append(f"\n--- Table {t_idx}: {len(table.rows)} rows x {len(table.columns)} cols ---")
        for r_idx, r in enumerate(table.rows):
            row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
            out_lines.append(f"R{r_idx}: {row_txt[:8]}")
    with open("std_docx_inspect.txt", "w", encoding="utf-8") as f:
        f.write("\n".join(out_lines))
    print("Standard docx inspection written.")
