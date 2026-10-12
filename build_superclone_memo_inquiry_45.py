import os
import sys
import json
import base64

if sys.stdout:
    sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENTATION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import signature_engine

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
BLUEPRINT_PATH = os.path.join(BASE_DIR, "table_layout_memo_inquiry_45.json")
GARUDA_PATH = os.path.join(BASE_DIR, "garuda_vector_official.png")

with open(BLUEPRINT_PATH, "r", encoding="utf-8") as f:
    bp = json.load(f)

with open(GARUDA_PATH, "rb") as f:
    GARUDA_B64 = base64.b64encode(f.read()).decode("utf-8")

# -----------------------------------------------------------------------------
# 1. GENERATE HTML
# -----------------------------------------------------------------------------
def generate_html(output_path):
    p1 = bp["pages"][0]
    c1 = p1["content"]
    sig1 = c1["signature"]
    
    p2 = bp["pages"][1]
    p3 = bp["pages"][2]

    def render_tracking_page(p_data, page_id):
        h = p_data["header"]
        tbl = p_data["table"]
        sigs = p_data["signatures"]
        sig_cdr = next(s for s in sigs if s["role_key"] == "commander")
        sig_insp = next(s for s in sigs if s["role_key"] == "inspector")

        # Table rows HTML (6 columns)
        rows_html = ""
        for r in tbl["rows"]:
            rows_html += f"""
                <tr>
                    <td class="text-center">{r["no"]}</td>
                    <td class="text-left" style="padding-left: 3mm;">{r["item_name"]}</td>
                    <td class="text-center">{r["stock_no"]}</td>
                    <td class="text-center">{r["voucher_no"]}</td>
                    <td class="text-center">{r["submitted_date"]}</td>
                    <td class="text-left">{r["note"]}</td>
                </tr>
            """
        # Fill empty rows up to 10 rows for clean government format
        current_rows = len(tbl["rows"])
        for _ in range(max(0, 10 - current_rows)):
            rows_html += """
                <tr>
                    <td>&nbsp;</td>
                    <td>&nbsp;</td>
                    <td>&nbsp;</td>
                    <td>&nbsp;</td>
                    <td>&nbsp;</td>
                    <td>&nbsp;</td>
                </tr>
            """

        # HTML Signature Blocks via signature_engine.py
        sig_cdr_html = signature_engine.render_signature_html(sig_cdr, margin_left_px=45, gap_pt=22)
        sig_insp_html = signature_engine.render_signature_html(sig_insp, margin_left_px=45, gap_pt=22)

        return f"""
        <div class="page-legal-landscape" id="{page_id}">
            <!-- Form Top Header -->
            <div style="display: flex; justify-content: space-between; align-items: flex-start; margin-bottom: 2mm;">
                <div></div>
                <div style="font-size: 18pt; font-weight: bold; text-align: center; margin-left: 60mm;">
                    {p_data["title"]}
                </div>
                <div style="text-align: right; font-size: 12pt; line-height: 1.2;">
                    <div style="font-weight: bold;">{p_data["form_code"]}</div>
                    <div>{p_data["form_subcode"]}</div>
                </div>
            </div>

            <!-- Meta block -->
            <div style="margin-bottom: 3mm; font-size: 12pt; line-height: 1.5;">
                <div style="margin-bottom: 1mm;">{h["doc_number"]}</div>
                <div style="margin-bottom: 1mm;">{h["subject"]}</div>
                <div style="margin-bottom: 1mm;">
                    <span style="font-weight: bold; display: inline-block; width: 12mm;">{h["to_label"]}</span>
                    <span>{h["to_value"]}</span>
                </div>
                <div style="margin-bottom: 2mm;">
                    <span style="font-weight: bold; display: inline-block; width: 12mm;">{h["from_label"]}</span>
                    <span>{h["from_value"]}</span>
                </div>
                <div style="padding-left: 20mm; margin-top: 2mm; margin-bottom: 3mm;">
                    {h["instruction"]}
                </div>
            </div>

            <!-- Table (6 columns) -->
            <table class="grid-table">
                <thead>
                    <tr>
                        <th style="width: 5%;">ลำดับ</th>
                        <th style="width: 24%;">รายการ สป.</th>
                        <th style="width: 22%;">หมายเลข สป.</th>
                        <th style="width: 15%;">ที่ใบเบิก ที่ทะเบียนหน่วย<br>จ่าย</th>
                        <th style="width: 16%;">วันที่ส่งหรือวางใบเบิก</th>
                        <th style="width: 18%;">หมายเหตุ  (บันทึกโดยหน่วยจ่าย)</th>
                    </tr>
                </thead>
                <tbody>
                    {rows_html}
                </tbody>
            </table>

            <!-- Double Signature Block (Universal Red Line Standard) -->
            <div style="display: flex; justify-content: space-between; margin-top: 12mm; padding: 0 20mm;">
                <!-- Left: Commander (ผบ.ร้อย / นยส.) -->
                <div style="width: 45%; display: flex; justify-content: center;">
                    {sig_cdr_html}
                </div>

                <!-- Right: Inspector (ผู้ตรวจสอบ) -->
                <div style="width: 45%; display: flex; justify-content: center;">
                    {sig_insp_html}
                </div>
            </div>

            <!-- Footer note (เลื่อนขึ้นมา 1 บรรทัด) -->
            <div style="text-align: center; margin-top: 4mm; font-size: 10.5pt;">
                {p_data["footer_note"]}
            </div>
        </div>
        """

    p2_html = render_tracking_page(p2, "page-2")
    p3_html = render_tracking_page(p3, "page-3")

    sig1_html = signature_engine.render_signature_html(sig1, margin_left_px=45, gap_pt=22)

    # Admin footer on page 1
    admin_footer_html = f"""
    <div style="position: absolute; right: 15mm; bottom: 15mm; font-size: 11pt; line-height: 1.3; text-align: left;">
        <table style="border-collapse: collapse; font-size: 11pt;">
            <tr>
                <td style="padding-right: 4mm;">{c1["admin_footer"][0]["rank"]}</td>
                <td style="padding-right: 4mm;">{c1["admin_footer"][0]["role"]}</td>
                <td>{c1["admin_footer"][0]["date"]}</td>
            </tr>
            <tr>
                <td style="padding-right: 4mm;">{c1["admin_footer"][1]["rank"]}</td>
                <td style="padding-right: 4mm;">{c1["admin_footer"][1]["role"]}&nbsp;&nbsp;&nbsp;&nbsp;{c1["admin_footer"][1]["day"]}</td>
                <td>{c1["admin_footer"][1]["date"]}</td>
            </tr>
        </table>
    </div>
    """

    html_content = f"""<!DOCTYPE html>
<html lang="th">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{bp["document_title"]}</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Sarabun:ital,wght@0,100;0,200;0,300;0,400;0,600;0,700;1,400&display=swap" rel="stylesheet">
    <style>
        @page portrait-page {{
            size: A4 portrait;
            margin: 0;
        }}
        @page landscape-page {{
            size: 355.6mm 215.9mm; /* Legal Landscape */
            margin: 0;
        }}
        * {{
            box-sizing: border-box;
            -webkit-print-color-adjust: exact;
            print-color-adjust: exact;
            font-family: 'TH Sarabun PSK', 'THSarabunNew', 'Sarabun', sans-serif;
        }}
        body {{
            margin: 0;
            padding: 20px 0;
            background-color: #525659;
            font-family: 'TH Sarabun PSK', 'THSarabunNew', 'Sarabun', sans-serif;
            font-size: 14pt;
            font-weight: normal;
            font-feature-settings: "locl" 1, "numr" 1;
            font-variant-numeric: tabular-nums;
            color: #000000;
            line-height: 1.15;
            display: flex;
            flex-direction: column;
            align-items: center;
        }}
        
        /* PAGE 1: A4 PORTRAIT */
        .page-portrait {{
            width: 210mm;
            height: 297mm;
            background-color: #ffffff;
            margin-bottom: 20mm;
            box-shadow: 0 4px 15px rgba(0, 0, 0, 0.35);
            position: relative;
            padding: 15mm 15mm 15mm 25mm; /* Top 1.5cm, Right 1.5cm, Bottom 1.5cm, Left 2.5cm */
            overflow: hidden;
            page-break-after: always;
            page: portrait-page;
            font-size: 12pt;
            line-height: 1.5;
        }}

        /* PAGE 2 & 3: LEGAL LANDSCAPE */
        .page-legal-landscape {{
            width: 355.6mm;
            height: 215.9mm;
            background-color: #ffffff;
            margin-bottom: 20mm;
            box-shadow: 0 4px 15px rgba(0, 0, 0, 0.35);
            position: relative;
            padding: 10mm 15mm 10mm 15mm;
            overflow: hidden;
            page-break-after: always;
            page: landscape-page;
        }}

        @media print {{
            body {{
                background-color: #ffffff;
                padding: 0;
            }}
            .page-portrait {{
                margin: 0;
                box-shadow: none;
                page: portrait-page;
            }}
            .page-legal-landscape {{
                margin: 0;
                box-shadow: none;
                page: landscape-page;
            }}
        }}

        .text-center {{ text-align: center; }}
        .text-left {{ text-align: left; }}
        .text-right {{ text-align: right; }}
        .font-bold {{ font-weight: bold; }}

        /* MEMO HEADER */
        .memo-copy-label {{
            text-align: center;
            font-size: 11pt;
            margin-bottom: 2mm;
        }}
        .memo-header-container {{
            position: relative;
            width: 100%;
            height: 30mm;
            margin-top: 0mm;
        }}
        .memo-garuda {{
            position: absolute;
            left: 0;
            top: 0;
            width: 30mm;
            height: 30mm;
            object-fit: contain;
        }}
        .memo-title {{
            position: absolute;
            left: 0;
            right: 0;
            top: 6mm;
            text-align: center;
            font-size: 24pt;
            font-weight: bold;
            line-height: 1;
            font-family: 'TH Sarabun PSK', 'THSarabunNew', 'Sarabun', sans-serif;
        }}
        .memo-meta-label {{
            font-size: 14pt;
            font-weight: bold;
            white-space: nowrap;
        }}
        .memo-meta-dots {{
            border-bottom: 1px dotted #000;
            padding-left: 2mm;
        }}
        .memo-content {{
            margin-top: 3mm;
            font-size: 12pt;
            line-height: 1.5;
        }}
        .memo-p {{
            text-indent: 2.5cm;
            margin: 0 0 2.5mm 0;
            text-align: justify;
            text-justify: inter-cluster;
            line-height: 1.5;
            font-size: 12pt;
        }}

        /* UNIVERSAL SIGNATURE BLOCK (RULE-UNIVERSAL-SIGNATURE-BLOCK-V1) */
        .sig-block {{
            display: inline-flex;
            flex-direction: column;
            align-items: flex-start;
            position: relative;
            min-width: 220px;
            font-size: 12pt;
            line-height: 1.0;
        }}
        .sig-name-pos-group {{
            display: inline-flex;
            flex-direction: column;
            align-items: center;
            position: relative;
            margin-left: 0;
        }}
        .sig-name-wrapper {{
            position: relative;
            display: inline-block;
            white-space: nowrap;
            margin-top: 18pt;
        }}
        .sig-rank-anchor {{
            position: absolute;
            right: 100%;
            bottom: calc(100% + 18px);
            white-space: nowrap;
            font-size: 12pt;
            line-height: 1.0;
        }}
        .sig-name {{
            margin: 0;
            line-height: 1.15;
            text-align: center;
            white-space: nowrap;
            font-size: 12pt;
            display: block;
        }}
        .sig-pos {{
            margin: 0;
            line-height: 1.15;
            text-align: center;
            white-space: nowrap;
            font-size: 12pt;
            display: block;
        }}

        /* TABLE STYLES (LEGAL LANDSCAPE) */
        table.grid-table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 11pt;
            margin-top: 2mm;
        }}
        table.grid-table th, table.grid-table td {{
            border: 1px solid #000000;
            padding: 1.5mm 1.5mm;
            vertical-align: middle;
        }}
        table.grid-table th {{
            text-align: center;
            font-weight: bold;
            background-color: #ffffff;
        }}
    </style>
</head>
<body>

    <!-- PAGE 1: บันทึกข้อความ (A4 Portrait) -->
    <div class="page-portrait" id="page-1">
        <div class="memo-copy-label">{p1["copy_label"]}</div>
        <div class="memo-header-container">
            <img class="memo-garuda" src="data:image/png;base64,{GARUDA_B64}" alt="Garuda">
            <div class="memo-title">{c1["title"]}</div>
        </div>

        <div style="margin-top: 1mm; font-size: 12pt; line-height: 1.5;">
            <!-- ส่วนราชการ -->
            <div style="display: flex; align-items: baseline; margin-bottom: 1.5mm;">
                <span class="memo-meta-label" style="margin-right: 2mm;">{c1["official_department"]}</span>
                <div style="flex: 1; border-bottom: 1px dotted #000; padding-left: 2mm; padding-bottom: 1px;">{c1["department_value"]}</div>
            </div>
            <!-- ที่ และ วันที่ -->
            <div style="display: flex; align-items: baseline; margin-bottom: 1.5mm;">
                <div style="display: flex; align-items: baseline; width: 45%;">
                    <span class="memo-meta-label" style="margin-right: 2mm;">{c1["doc_number_label"]}</span>
                    <div style="flex: 1; border-bottom: 1px dotted #000; padding-left: 2mm; padding-bottom: 1px;">{c1["doc_number_value"]}</div>
                </div>
                <div style="display: flex; align-items: baseline; width: 55%; padding-left: 10mm;">
                    <span class="memo-meta-label" style="margin-right: 2mm;">{c1["doc_date_label"]}</span>
                    <div style="flex: 1; border-bottom: 1px dotted #000; padding-left: 2mm; padding-bottom: 1px;">&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;{c1["doc_date_month_year"]}</div>
                </div>
            </div>
            <!-- เรื่อง -->
            <div style="display: flex; align-items: baseline; margin-bottom: 1.5mm;">
                <span class="memo-meta-label" style="margin-right: 2mm;">{c1["subject_label"]}</span>
                <div style="flex: 1; border-bottom: 1px dotted #000; padding-left: 2mm; padding-bottom: 1px;">{c1["subject_value"]}</div>
            </div>
        </div>

        <div class="memo-content">
            <div style="margin-bottom: 2mm;">
                <span class="font-bold">{c1["recipient_label"]}</span>&nbsp;&nbsp;{c1["recipient_value"]}
            </div>
            <div style="margin-bottom: 2mm;">
                <span class="font-bold">{c1["reference_label"]}</span>&nbsp;&nbsp;{c1["reference_value"]}
            </div>

            <div style="margin: 0 0 2.5mm 0; font-size: 12pt; line-height: 1.5;">
                <table style="width: 100%; border-collapse: collapse; font-size: 12pt; line-height: 1.5;">
                    <tr>
                        <td style="width: 28mm; vertical-align: top; white-space: nowrap;"><span class="font-bold">{c1["attachments_label"]}</span></td>
                        <td style="vertical-align: top;">{c1["attachments"][0]["title"]}</td>
                        <td style="width: 35mm; vertical-align: top; text-align: right; white-space: nowrap;">{c1["attachments"][0]["quantity"]}</td>
                    </tr>
                </table>
            </div>

            <p class="memo-p">{c1["paragraphs"][0]}</p>
            <p class="memo-p">{c1["paragraphs"][1]}</p>
        </div>

        <!-- SIGNATURE BLOCK (RULE: เคาะ 5 Enter จากบรรทัดสุดท้าย, พิกัดซ้าย 130 มม. - Margin ซ้าย 25 มม. = padding-left 105 มม.) -->
        <div style="margin-top: 55pt; padding-left: 105mm;">
            {sig1_html}
        </div>

        <!-- ADMIN FOOTER BLOCK -->
        {admin_footer_html}
    </div>

    <!-- PAGE 2: ใบติดตามใบเบิก แผ่นที่ ๑ (Legal Landscape) -->
    {p2_html}

    <!-- PAGE 3: ใบติดตามใบเบิก แผ่นที่ ๒ (Legal Landscape) -->
    {p3_html}

</body>
</html>
    """

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"Generated HTML: {output_path}")


# -----------------------------------------------------------------------------
# 2. GENERATE DOCX
# -----------------------------------------------------------------------------
def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["val", "color", "sz", "space"]:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))

def set_run_font(run, font_name="TH Sarabun PSK", size_pt=12, bold=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:cs'), font_name)

def generate_docx(output_path):
    doc = Document()

    # SECTION 1: Page 1 A4 Portrait
    sec1 = doc.sections[0]
    sec1.orientation = WD_ORIENTATION.PORTRAIT
    sec1.page_width = Inches(8.27)
    sec1.page_height = Inches(11.69)
    sec1.top_margin = Inches(0.59) # 1.5 cm
    sec1.right_margin = Inches(0.59) # 1.5 cm
    sec1.bottom_margin = Inches(0.59) # 1.5 cm
    sec1.left_margin = Inches(0.98) # 2.5 cm

    p1 = bp["pages"][0]
    c1 = p1["content"]
    sig1 = c1["signature"]

    # Top Copy Label
    p_cp = doc.add_paragraph()
    p_cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cp.paragraph_format.space_after = Pt(2)
    r_cp = p_cp.add_run(p1["copy_label"])
    set_run_font(r_cp, "TH Sarabun PSK", 11)

    # Header Table
    tbl_h = doc.add_table(rows=1, cols=2)
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_img = tbl_h.cell(0, 0)
    c_img.width = Inches(1.3)
    p_img = c_img.paragraphs[0]
    p_img.paragraph_format.space_after = Pt(0)
    r_img = p_img.add_run()
    r_img.add_picture(GARUDA_PATH, width=Inches(1.18))

    c_title = tbl_h.cell(0, 1)
    c_title.width = Inches(5.4)
    p_t = c_title.paragraphs[0]
    p_t.paragraph_format.space_before = Pt(14)
    p_t.paragraph_format.space_after = Pt(0)
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_t.add_run(c1["title"])
    set_run_font(r_t, "TH Sarabun PSK", 24, bold=True)

    set_cell_border(c_img)
    set_cell_border(c_title)

    # Meta Lines
    meta_lines = [
        (c1["official_department"], c1["department_value"]),
        (f'{c1["doc_number_label"]}  {c1["doc_number_value"]}', f'{c1["doc_date_label"]}            {c1["doc_date_month_year"]}'),
        (c1["subject_label"], c1["subject_value"]),
        (c1["recipient_label"], c1["recipient_value"]),
        (c1["reference_label"], c1["reference_value"])
    ]

    for label, val in meta_lines:
        p_m = doc.add_paragraph()
        p_m.paragraph_format.space_after = Pt(2)
        p_m.paragraph_format.line_spacing = 1.5
        r_lbl = p_m.add_run(label + "  ")
        set_run_font(r_lbl, "TH Sarabun PSK", 14, bold=True)
        r_val = p_m.add_run(val)
        set_run_font(r_val, "TH Sarabun PSK", 12)

    # Attachments
    p_att = doc.add_paragraph()
    p_att.paragraph_format.space_after = Pt(2)
    p_att.paragraph_format.line_spacing = 1.5
    r_al = p_att.add_run(c1["attachments_label"] + "  ")
    set_run_font(r_al, "TH Sarabun PSK", 14, bold=True)
    r_av = p_att.add_run(f'{c1["attachments"][0]["title"]}     {c1["attachments"][0]["quantity"]}')
    set_run_font(r_av, "TH Sarabun PSK", 12)

    # Paragraphs
    for para in c1["paragraphs"]:
        p_p = doc.add_paragraph()
        p_p.paragraph_format.first_line_indent = Inches(0.98)
        p_p.paragraph_format.space_after = Pt(4)
        p_p.paragraph_format.line_spacing = 1.5
        p_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_p = p_p.add_run(para)
        set_run_font(r_p, "TH Sarabun PSK", 12)

    # Signature Block via signature_engine.py
    # RULE-UNIVERSAL-SIGNATURE-BLOCK-V1: 5 Enters + 130mm from page left edge (4.13 inches from left margin)
    signature_engine.render_signature_docx(
        doc,
        sig1,
        axis_inch=4.13,
        font_name="TH Sarabun PSK",
        font_size=16,
        space_after_rank_pt=22,
        space_before_block_pt=55
    )

    # Admin Footer Block at bottom right
    p_adm1 = doc.add_paragraph()
    p_adm1.paragraph_format.space_before = Pt(30)
    p_adm1.paragraph_format.space_after = Pt(0)
    p_adm1.paragraph_format.tab_stops.add_tab_stop(Inches(4.5), WD_TAB_ALIGNMENT.LEFT)
    p_adm1.paragraph_format.tab_stops.add_tab_stop(Inches(5.5), WD_TAB_ALIGNMENT.LEFT)
    p_adm1.paragraph_format.tab_stops.add_tab_stop(Inches(6.4), WD_TAB_ALIGNMENT.LEFT)
    r_adm1 = p_adm1.add_run(f"\t{c1['admin_footer'][0]['rank']}\t{c1['admin_footer'][0]['role']}\t{c1['admin_footer'][0]['date']}")
    set_run_font(r_adm1, "TH Sarabun PSK", 11)

    p_adm2 = doc.add_paragraph()
    p_adm2.paragraph_format.space_before = Pt(0)
    p_adm2.paragraph_format.space_after = Pt(0)
    p_adm2.paragraph_format.tab_stops.add_tab_stop(Inches(4.5), WD_TAB_ALIGNMENT.LEFT)
    p_adm2.paragraph_format.tab_stops.add_tab_stop(Inches(5.5), WD_TAB_ALIGNMENT.LEFT)
    p_adm2.paragraph_format.tab_stops.add_tab_stop(Inches(6.4), WD_TAB_ALIGNMENT.LEFT)
    r_adm2 = p_adm2.add_run(f"\t{c1['admin_footer'][1]['rank']}\t{c1['admin_footer'][1]['role']}  {c1['admin_footer'][1]['day']}\t{c1['admin_footer'][1]['date']}")
    set_run_font(r_adm2, "TH Sarabun PSK", 11)

    # FUNCTION TO RENDER LEGAL LANDSCAPE TRACKING FORM PAGES IN DOCX
    def add_docx_tracking_page(p_data):
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        sec.orientation = WD_ORIENTATION.LANDSCAPE
        sec.page_width = Inches(14.0) # US Legal
        sec.page_height = Inches(8.5)
        sec.top_margin = Inches(0.5)
        sec.right_margin = Inches(0.59)
        sec.bottom_margin = Inches(0.5)
        sec.left_margin = Inches(0.59)

        h = p_data["header"]
        tbl_data = p_data["table"]
        sigs = p_data["signatures"]
        sig_cdr = next(s for s in sigs if s["role_key"] == "commander")
        sig_insp = next(s for s in sigs if s["role_key"] == "inspector")

        # Top header
        p_top = doc.add_paragraph()
        p_top.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_fc = p_top.add_run(f'{p_data["form_code"]}\n{p_data["form_subcode"]}')
        set_run_font(r_fc, "TH Sarabun PSK", 12, bold=True)

        p_th = doc.add_paragraph()
        p_th.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_th.paragraph_format.space_after = Pt(2)
        r_th = p_th.add_run(p_data["title"])
        set_run_font(r_th, "TH Sarabun PSK", 18, bold=True)

        # Meta
        meta_items = [
            h["doc_number"],
            h["subject"],
            f'{h["to_label"]}   {h["to_value"]}',
            f'{h["from_label"]}   {h["from_value"]}',
            f'        {h["instruction"]}'
        ]
        for itm in meta_items:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.space_after = Pt(1)
            p_it.paragraph_format.line_spacing = 1.15
            r_it = p_it.add_run(itm)
            set_run_font(r_it, "TH Sarabun PSK", 12)

        # Grid Table (6 columns)
        cols_w = [Inches(0.6), Inches(3.2), Inches(2.7), Inches(1.9), Inches(1.9), Inches(2.5)]
        t_form = doc.add_table(rows=1, cols=6)
        t_form.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Header Row
        headers = ["ลำดับ", "รายการ สป.", "หมายเลข สป.", "ที่ใบเบิก ที่ทะเบียนหน่วย\nจ่าย", "วันที่ส่งหรือวางใบเบิก", "หมายเหตุ (บันทึกโดยหน่วยจ่าย)"]
        for c_i, h_txt in enumerate(headers):
            cell = t_form.cell(0, c_i)
            cell.width = cols_w[c_i]
            p_c = cell.paragraphs[0]
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_c = p_c.add_run(h_txt)
            set_run_font(r_c, "TH Sarabun PSK", 11, bold=True)
            set_cell_border(cell, top={"val": "single", "sz": "4", "color": "000000"},
                                  bottom={"val": "single", "sz": "4", "color": "000000"},
                                  left={"val": "single", "sz": "4", "color": "000000"},
                                  right={"val": "single", "sz": "4", "color": "000000"})

        # Data Rows
        for r_data in tbl_data["rows"]:
            row = t_form.add_row()
            vals = [r_data["no"], r_data["item_name"], r_data["stock_no"], r_data["voucher_no"], r_data["submitted_date"], r_data["note"]]
            aligns = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT]
            for c_i, val in enumerate(vals):
                cell = row.cells[c_i]
                cell.width = cols_w[c_i]
                p_c = cell.paragraphs[0]
                p_c.alignment = aligns[c_i]
                r_c = p_c.add_run(val)
                set_run_font(r_c, "TH Sarabun PSK", 11)
                set_cell_border(cell, top={"val": "single", "sz": "4", "color": "000000"},
                                      bottom={"val": "single", "sz": "4", "color": "000000"},
                                      left={"val": "single", "sz": "4", "color": "000000"},
                                      right={"val": "single", "sz": "4", "color": "000000"})

        # Pad rows
        for _ in range(max(0, 10 - len(tbl_data["rows"]))):
            row = t_form.add_row()
            for c_i in range(6):
                cell = row.cells[c_i]
                cell.width = cols_w[c_i]
                set_cell_border(cell, top={"val": "single", "sz": "4", "color": "000000"},
                                      bottom={"val": "single", "sz": "4", "color": "000000"},
                                      left={"val": "single", "sz": "4", "color": "000000"},
                                      right={"val": "single", "sz": "4", "color": "000000"})

        # Signatures 2 columns table
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(14)

        t_sig = doc.add_table(rows=1, cols=2)
        t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_left = t_sig.cell(0, 0)
        c_right = t_sig.cell(0, 1)
        c_left.width = Inches(6.0)
        c_right.width = Inches(6.0)

        # Left: Commander (ผบ.ร้อย / นยส.) — via signature_engine.py
        sig_cdr_copy = dict(sig_cdr)
        sig_cdr_copy["add_top_spacer"] = True
        signature_engine.render_signature_docx(
            c_left,
            sig_cdr_copy,
            axis_inch=3.0,
            font_name="TH Sarabun PSK",
            font_size=12,
            space_after_rank_pt=22,
            space_before_block_pt=0
        )

        # Right: Inspector (ผู้ตรวจสอบ) — via signature_engine.py
        signature_engine.render_signature_docx(
            c_right,
            sig_insp,
            axis_inch=3.0,
            font_name="TH Sarabun PSK",
            font_size=12,
            space_after_rank_pt=22,
            space_before_block_pt=0
        )

        set_cell_border(c_left)
        set_cell_border(c_right)

        # Footer Note (เลื่อนขึ้นมา 1 บรรทัด)
        p_fn = doc.add_paragraph()
        p_fn.paragraph_format.space_before = Pt(6)
        p_fn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_fn = p_fn.add_run(p_data["footer_note"])
        set_run_font(r_fn, "TH Sarabun PSK", 10.5)

    add_docx_tracking_page(bp["pages"][1])
    add_docx_tracking_page(bp["pages"][2])

    doc.save(output_path)
    print(f"Generated DOCX: {output_path}")


# -----------------------------------------------------------------------------
# 3. GENERATE XLSX
# -----------------------------------------------------------------------------
def generate_xlsx(output_path):
    wb = openpyxl.Workbook()
    # Sheet 1: บันทึกข้อความ
    ws1 = wb.active
    ws1.title = "บันทึกข้อความ"
    ws1.views.sheetView[0].showGridLines = True

    p1 = bp["pages"][0]
    c1 = p1["content"]
    sig1 = c1["signature"]

    font_regular = Font(name="TH Sarabun PSK", size=12)
    font_bold = Font(name="TH Sarabun PSK", size=12, bold=True)
    font_large_bold = Font(name="TH Sarabun PSK", size=18, bold=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws1["C2"] = p1["copy_label"]
    ws1["C2"].font = font_regular
    ws1["C2"].alignment = align_center

    ws1["C4"] = c1["title"]
    ws1["C4"].font = font_large_bold
    ws1["C4"].alignment = align_center

    memo_rows = [
        ("ส่วนราชการ", c1["department_value"]),
        ("ที่", f'{c1["doc_number_value"]}       วันที่            {c1["doc_date_month_year"]}'),
        ("เรื่อง", c1["subject_value"]),
        ("เรียน", c1["recipient_value"]),
        ("อ้างถึง", c1["reference_value"]),
        ("สิ่งที่ส่งมาด้วย", f'{c1["attachments"][0]["title"]}   {c1["attachments"][0]["quantity"]}'),
        ("", c1["paragraphs"][0]),
        ("", c1["paragraphs"][1])
    ]

    for idx, (lbl, val) in enumerate(memo_rows, start=6):
        ws1[f"B{idx}"] = lbl
        ws1[f"B{idx}"].font = font_bold
        ws1[f"C{idx}"] = val
        ws1[f"C{idx}"].font = font_regular
        ws1.row_dimensions[idx].height = 28

    # Signature via signature_engine.py
    signature_engine.render_signature_xlsx(
        ws1,
        sig1,
        start_row=18,
        rank_col="D",
        name_col="E",
        font_name="TH Sarabun PSK",
        font_size=12
    )

    # Admin footer in Excel
    ws1["D22"] = c1["admin_footer"][0]["rank"]
    ws1["D22"].font = font_regular
    ws1["E22"] = f"{c1['admin_footer'][0]['role']}    {c1['admin_footer'][0]['date']}"
    ws1["E22"].font = font_regular

    ws1["D23"] = c1["admin_footer"][1]["rank"]
    ws1["D23"].font = font_regular
    ws1["E23"] = f"{c1['admin_footer'][1]['role']}  {c1['admin_footer'][1]['day']}    {c1['admin_footer'][1]['date']}"
    ws1["E23"].font = font_regular

    ws1.column_dimensions["A"].width = 5
    ws1.column_dimensions["B"].width = 18
    ws1.column_dimensions["C"].width = 50
    ws1.column_dimensions["D"].width = 25
    ws1.column_dimensions["E"].width = 30

    # Sheet 2 & 3: ใบติดตามใบเบิก
    def add_xlsx_tracking_sheet(p_data, sheet_title):
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True
        h = p_data["header"]
        tbl = p_data["table"]
        sigs = p_data["signatures"]
        sig_cdr = next(s for s in sigs if s["role_key"] == "commander")
        sig_insp = next(s for s in sigs if s["role_key"] == "inspector")

        thin = Side(style='thin', color='000000')
        border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

        # Header
        ws["C2"] = p_data["title"]
        ws["C2"].font = font_large_bold
        ws["C2"].alignment = align_center

        ws["F2"] = f'{p_data["form_code"]}\n{p_data["form_subcode"]}'
        ws["F2"].font = font_bold
        ws["F2"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

        ws["B4"] = h["doc_number"]
        ws["B4"].font = font_regular
        ws["B5"] = h["subject"]
        ws["B5"].font = font_regular
        ws["B6"] = f'{h["to_label"]}   {h["to_value"]}'
        ws["B6"].font = font_regular
        ws["B7"] = f'{h["from_label"]}   {h["from_value"]}'
        ws["B7"].font = font_regular
        ws["B8"] = f'        {h["instruction"]}'
        ws["B8"].font = font_regular

        # Table Header (6 columns)
        headers = ["ลำดับ", "รายการ สป.", "หมายเลข สป.", "ที่ใบเบิก ที่ทะเบียนหน่วยจ่าย", "วันที่ส่งหรือวางใบเบิก", "หมายเหตุ (บันทึกโดยหน่วยจ่าย)"]
        cols = ["A", "B", "C", "D", "E", "F"]
        for col_letter, h_txt in zip(cols, headers):
            cell = ws[f"{col_letter}10"]
            cell.value = h_txt
            cell.font = font_bold
            cell.alignment = align_center
            cell.border = border_all
        ws.row_dimensions[10].height = 28

        # Table Data
        for r_idx, r_data in enumerate(tbl["rows"], start=11):
            vals = [r_data["no"], r_data["item_name"], r_data["stock_no"], r_data["voucher_no"], r_data["submitted_date"], r_data["note"]]
            for c_idx, (col_letter, val) in enumerate(zip(cols, vals)):
                cell = ws[f"{col_letter}{r_idx}"]
                cell.value = val
                cell.font = font_regular
                cell.border = border_all
                if c_idx in [0, 2, 3, 4]:
                    cell.alignment = align_center
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[r_idx].height = 22

        start_sig_row = 11 + len(tbl["rows"]) + 2

        # Signatures via signature_engine.py
        # Left signature in B:C
        signature_engine.render_signature_xlsx(
            ws,
            sig_cdr,
            start_row=start_sig_row,
            rank_col="B",
            name_col="C",
            font_name="TH Sarabun PSK",
            font_size=12
        )

        # Right signature in E:F
        signature_engine.render_signature_xlsx(
            ws,
            sig_insp,
            start_row=start_sig_row,
            rank_col="E",
            name_col="F",
            font_name="TH Sarabun PSK",
            font_size=12
        )

        # Footer note (เลื่อนขึ้นมา 1 แถว)
        ws[f"B{start_sig_row + 4}"] = p_data["footer_note"]
        ws[f"B{start_sig_row + 4}"].font = Font(name="TH Sarabun PSK", size=10.5)
        ws[f"B{start_sig_row + 4}"].alignment = align_center
        ws.merge_cells(f"B{start_sig_row + 4}:F{start_sig_row + 4}")

        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 28
        ws.column_dimensions["D"].width = 24
        ws.column_dimensions["E"].width = 22
        ws.column_dimensions["F"].width = 28

    add_xlsx_tracking_sheet(bp["pages"][1], "ใบติดตามใบเบิก_แผ่นที่1")
    add_xlsx_tracking_sheet(bp["pages"][2], "ใบติดตามใบเบิก_แผ่นที่2")

    wb.save(output_path)
    print(f"Generated XLSX: {output_path}")


if __name__ == "__main__":
    html_out = os.path.join(BASE_DIR, "ขอส่งใบติดตามทวงถาม_45_วัน_Rebuild_Thai.html")
    docx_out = os.path.join(BASE_DIR, "ขอส่งใบติดตามทวงถาม_45_วัน_Rebuild_Thai.docx")
    xlsx_out = os.path.join(BASE_DIR, "ขอส่งใบติดตามทวงถาม_45_วัน_Rebuild_Thai.xlsx")

    generate_html(html_out)
    generate_docx(docx_out)
    generate_xlsx(xlsx_out)
    print("All builds completed successfully!")
