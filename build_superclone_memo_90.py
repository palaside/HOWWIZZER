import os
import sys
import json
import base64
import shutil

if sys.stdout:
    sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENTATION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import signature_engine

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
BLUEPRINT_PATH = os.path.join(BASE_DIR, "table_layout_memo_90.json")
GARUDA_PATH = os.path.join(BASE_DIR, "garuda_vector_official.png")
if not os.path.exists(GARUDA_PATH):
    GARUDA_PATH = os.path.join(BASE_DIR, "garuda_memo.png")

with open(BLUEPRINT_PATH, "r", encoding="utf-8") as f:
    bp = json.load(f)

with open(GARUDA_PATH, "rb") as f:
    GARUDA_B64 = base64.b64encode(f.read()).decode("utf-8")

# -----------------------------------------------------------------------------
# 1. GENERATE HTML
# -----------------------------------------------------------------------------
def generate_html(output_path, is_copy=True):
    c = bp["content"]
    sig = c["signature"]

    copy_label_html = f'<div class="memo-copy-label">{c.get("copy_type", "( สำเนาคู่ฉบับ )")}</div>' if is_copy else '<div class="memo-copy-label" style="visibility: hidden;">&nbsp;</div>'

    sig_html = f"""
    <div class="sig-block" style="font-size: 17pt; line-height: 1.0;">
        <div class="sig-name-pos-group" style="margin-left: 45px;">
            <div class="sig-name-wrapper" style="margin-top: 22pt;">
                <span class="sig-rank-anchor" style="font-size: 17pt;">{sig['rank']}</span>
                <div class="sig-name" style="font-size: 17pt;">({sig['name']})</div>
            </div>
            <div class="sig-pos" style="font-size: 17pt;">{sig['position']}</div>
        </div>
    </div>
    """.strip()

    admin_footer_html = ""
    if is_copy and "routing_footer" in c:
        rf = c["routing_footer"]
        day0 = rf[0].get("day", "")
        day1 = rf[1].get("day", "")
        admin_footer_html = f"""
    <div class="memo-footer-staff">
        <div class="staff-grid">
            <div>{rf[0]["rank"]}</div>
            <div>{rf[0]["role"]}</div>
            <div style="text-align: right;">{day0}</div>
            <div style="padding-left: 2mm;">{rf[0]["date"]}</div>

            <div>{rf[1]["rank"]}</div>
            <div>{rf[1]["role"]}</div>
            <div style="text-align: right;">{day1}</div>
            <div style="padding-left: 2mm;">{rf[1]["date"]}</div>
        </div>
    </div>
        """

    html_content = f"""<!DOCTYPE html>
<html lang="th">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{bp["document_title"]}</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Sarabun:ital,wght@0,100;0,200;0,300;0,400;0,600;0,700;1,100;1,300;1,400;1,700&display=swap" rel="stylesheet">
    <style>
        /* ======================================================== */
        /* OFFICIAL GOVERNMENT TH SARABUN PSK ENGINE (PRIMARY SSOT) */
        /* ======================================================== */
        @font-face {{
            font-family: 'TH Sarabun PSK';
            font-style: normal;
            font-weight: 400;
            src: local('TH Sarabun PSK'), local('THSarabunPSK'),
                 url('/fonts/THSarabunPSK.ttf') format('truetype'),
                 url('fonts/THSarabunPSK.ttf') format('truetype'),
                 url('../fonts/THSarabunPSK.ttf') format('truetype'),
                 url('../../fonts/THSarabunPSK.ttf') format('truetype');
        }}
        @font-face {{
            font-family: 'TH Sarabun PSK';
            font-style: normal;
            font-weight: 300;
            src: local('TH Sarabun PSK'), local('THSarabunPSK'),
                 url('/fonts/THSarabunPSK.ttf') format('truetype'),
                 url('fonts/THSarabunPSK.ttf') format('truetype'),
                 url('../fonts/THSarabunPSK.ttf') format('truetype'),
                 url('../../fonts/THSarabunPSK.ttf') format('truetype');
        }}
        @font-face {{
            font-family: 'TH Sarabun PSK';
            font-style: normal;
            font-weight: 600;
            src: local('TH Sarabun PSK Bold'), local('THSarabunPSK-Bold'),
                 url('/fonts/THSarabunPSK-Bold.ttf') format('truetype'),
                 url('fonts/THSarabunPSK-Bold.ttf') format('truetype'),
                 url('../fonts/THSarabunPSK-Bold.ttf') format('truetype'),
                 url('../../fonts/THSarabunPSK-Bold.ttf') format('truetype');
        }}
        @font-face {{
            font-family: 'TH Sarabun PSK';
            font-style: normal;
            font-weight: 700;
            src: local('TH Sarabun PSK Bold'), local('THSarabunPSK-Bold'),
                 url('/fonts/THSarabunPSK-Bold.ttf') format('truetype'),
                 url('fonts/THSarabunPSK-Bold.ttf') format('truetype'),
                 url('../fonts/THSarabunPSK-Bold.ttf') format('truetype'),
                 url('../../fonts/THSarabunPSK-Bold.ttf') format('truetype');
        }}

        @page {{
            size: A4 portrait;
            margin: 0;
        }}
        * {{
            box-sizing: border-box;
            -webkit-print-color-adjust: exact;
            print-color-adjust: exact;
            font-family: 'TH Sarabun PSK', 'THSarabunPSK', 'TH Sarabun New', 'THSarabunNew', 'Sarabun', sans-serif;
        }}
        body {{
            margin: 0;
            padding: 20px 0;
            background-color: #525659;
            font-family: 'TH Sarabun PSK', 'THSarabunPSK', 'TH Sarabun New', 'THSarabunNew', 'Sarabun', sans-serif;
            font-size: 17pt;
            font-weight: normal;
            font-feature-settings: "locl" 1, "numr" 1;
            font-variant-numeric: tabular-nums;
            color: #000000;
            line-height: 1.25;
            display: flex;
            flex-direction: column;
            align-items: center;
        }}
        
        .page-portrait {{
            width: 210mm;
            height: 297mm;
            background-color: #ffffff;
            margin-bottom: 20mm;
            box-shadow: 0 4px 15px rgba(0, 0, 0, 0.35);
            position: relative;
            padding: 15mm 15mm 15mm 25mm; /* Top 1.5cm, Right 1.5cm, Bottom 1.5cm, Left 2.5cm */
            overflow: hidden;
            page-break-after: always;
            font-size: 17pt;
            line-height: 1.25;
        }}

        @media print {{
            body {{
                background-color: #ffffff;
                padding: 0;
            }}
            .page-portrait {{
                margin: 0;
                box-shadow: none;
            }}
        }}

        .text-center {{ text-align: center; }}
        .text-left {{ text-align: left; }}
        .text-right {{ text-align: right; }}
        .font-bold {{ font-weight: 600; }}

        /* MEMO HEADER */
        .memo-copy-label {{
            text-align: center;
            font-size: 13pt;
            margin-bottom: 2mm;
        }}
        .memo-header-container {{
            position: relative;
            width: 100%;
            height: 30mm;
            margin-top: 0mm;
        }}
        .memo-garuda {{
            position: absolute;
            left: 0;
            top: 0;
            width: 30mm;
            height: 30mm;
            object-fit: contain;
        }}
        .memo-title {{
            position: absolute;
            left: 0;
            right: 0;
            top: 6mm;
            text-align: center;
            font-size: 29pt;
            font-weight: 600;
            line-height: 1;
            font-family: 'TH Sarabun PSK', 'THSarabunNew', 'Sarabun', sans-serif;
        }}
        .memo-meta-label {{
            font-size: 20pt;
            font-weight: 600;
            line-height: 1.05;
            padding-bottom: 1.5px;
            white-space: nowrap;
        }}
        .memo-meta-dots {{
            border-bottom: 1px dotted #000;
            padding-left: 2mm;
            padding-bottom: 1.5px;
            font-size: 17pt;
            line-height: 1.05;
        }}
        .memo-content {{
            margin-top: 2mm;
            font-size: 17pt;
            line-height: 1.25;
        }}
        .memo-p {{
            text-indent: 2.5cm;
            margin: 0 0 1.5mm 0;
            text-align: justify;
            text-justify: inter-cluster; word-break: keep-all; overflow-wrap: break-word;
            line-height: 1.25;
            font-size: 17pt;
        }}

        /* UNIVERSAL SIGNATURE BLOCK (RULE-UNIVERSAL-SIGNATURE-BLOCK-V1) */
        .sig-block {{
            display: inline-flex;
            flex-direction: column;
            align-items: flex-start;
            position: relative;
            min-width: 220px;
            font-size: 17pt;
            line-height: 1.2;
        }}
        .sig-name-pos-group {{
            display: inline-flex;
            flex-direction: column;
            align-items: center;
            position: relative;
            margin-left: 0;
        }}
        .sig-name-wrapper {{
            position: relative;
            display: inline-block;
            white-space: nowrap;
            margin-top: 18pt;
        }}
        .sig-rank-anchor {{
            position: absolute;
            right: 100%;
            bottom: calc(100% + 18px);
            white-space: nowrap;
            font-size: 17pt;
            line-height: 1.2;
        }}
        .sig-name {{
            margin: 0;
            line-height: 1.2;
            text-align: center;
            white-space: nowrap;
            font-size: 17pt;
            display: block;
        }}
        .sig-pos {{
            margin: 0;
            line-height: 1.2;
            text-align: center;
            white-space: nowrap;
            font-size: 17pt;
            display: block;
        }}

        /* STAFF INITIALS FOOTER (RULE-STAFF-REVIEW-FOOTER-V1: Mandatory 4-Column Dual-Axis) */
        .memo-footer-staff {{
            margin-top: 22mm; /* อยู่ใต้ ผบ.ป.๔ พัน.๔ ร้อย.๒ ๕ เอนเตอร์ */
            display: flex;
            justify-content: flex-end;
        }}
        .staff-grid {{
            display: grid;
            grid-template-columns: 26mm 22mm 8mm 20mm;
            font-size: 13.5pt;
            font-weight: 300;
            line-height: 1.25;
            text-align: left;
        }}

        /* RULE 9: THIN FONT-WEIGHT STANDARD (font-weight: 300 for non-headings) */
        p, span, td, .memo-p, .memo-content, .memo-att-row, .memo-att-desc, .memo-att-qty,
        .sig-rank-anchor, .sig-name, .sig-pos, .staff-grid, .staff-grid div {{
            font-weight: 300;
        }}

        .memo-title, .memo-meta-label,
        .font-bold, b, strong, .sig-verify {{
            font-weight: 600;
        }}
    </style>
</head>
<body>

    <!-- บันทึกข้อความ (A4 Portrait) -->
    <div class="page-portrait" id="page-1">
        {copy_label_html}
        <div class="memo-header-container">
            <img class="memo-garuda" src="data:image/png;base64,{GARUDA_B64}" alt="ตราครุฑ">
            <div class="memo-title">{c["title"]}</div>
        </div>

        <div style="margin-top: 2mm; font-size: 17pt; line-height: 1.25;">
            <!-- ส่วนราชการ -->
            <div style="display: flex; align-items: baseline; margin-bottom: 1.5mm;">
                <span class="memo-meta-label" style="margin-right: 2mm;">{c["official_department"]}</span>
                <div class="memo-meta-dots" style="flex: 1;">{c["department_value"]}</div>
            </div>
            <!-- ที่ และ วันที่ -->
            <div style="display: flex; align-items: baseline; margin-bottom: 1.5mm;">
                <div style="display: flex; align-items: baseline; flex: 1.2; margin-right: 6mm;">
                    <span class="memo-meta-label" style="margin-right: 2mm;">{c["doc_number_label"]}</span>
                    <div class="memo-meta-dots" style="flex: 1;">{c["doc_number_value"]}</div>
                </div>
                <div style="display: flex; align-items: baseline; flex: 1;">
                    <span class="memo-meta-label" style="margin-right: 2mm;">{c["doc_date_label"]}</span>
                    <div class="memo-meta-dots" style="flex: 1;">&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;{c["doc_date_month_year"]}</div>
                </div>
            </div>
            <!-- เรื่อง -->
            <div style="display: flex; align-items: baseline; margin-bottom: 1.5mm;">
                <span class="memo-meta-label" style="margin-right: 2mm;">{c["subject_label"]}</span>
                <div class="memo-meta-dots" style="flex: 1;">{c["subject_value"]}</div>
            </div>
        </div>

        <div class="memo-content">
            <div style="margin-bottom: 2mm; font-size: 17pt;">
                <span class="font-bold">{c["recipient_label"]}</span>&nbsp;&nbsp;{c["recipient_value"]}
            </div>
            <div style="margin-bottom: 2mm; font-size: 17pt;">
                <span class="font-bold">{c["reference_label"]}</span>&nbsp;&nbsp;{c["reference_value"]}
            </div>

            <div style="margin: 0 0 2.5mm 0; font-size: 17pt; line-height: 1.25;">
                <table style="width: 100%; border-collapse: collapse; font-size: 17pt; line-height: 1.25;">
                    <tr>
                        <td style="width: 32mm; vertical-align: top; white-space: nowrap;"><span class="font-bold">{c["attachments_label"]}</span></td>
                        <td style="vertical-align: top; font-size: 17pt;">{c["attachments"][0]["title"]}</td>
                        <td style="width: 35mm; vertical-align: top; text-align: right; white-space: nowrap; font-size: 17pt;">{c["attachments"][0]["quantity"]}</td>
                    </tr>
                </table>
            </div>

            <p class="memo-p" style="font-size: 17pt; line-height: 1.25;">{c["paragraphs"][0]}</p>
            <p class="memo-p" style="font-size: 17pt; line-height: 1.25;">{c["paragraphs"][1]}</p>
        </div>

        <!-- SIGNATURE BLOCK (RULE: เคาะ 5 Enter จากบรรทัดสุดท้าย, พิกัดซ้าย 130 มม. - Margin ซ้าย 25 มม. = padding-left 105 มม.) -->
        <div style="margin-top: 55pt; padding-left: 105mm;">
            {sig_html}
        </div>

        <!-- ADMIN FOOTER BLOCK -->
        {admin_footer_html}
    </div>

</body>
</html>
    """
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"Generated HTML: {output_path}")


# -----------------------------------------------------------------------------
# 2. GENERATE DOCX
# -----------------------------------------------------------------------------
def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'<w:tcBorders %s/>' % nsdecls('w'))
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = parse_xml(r'<%s %s %s/>' % (tag, nsdecls('w'),
                ''.join(['w:{}="{}" '.format(k, v) for k, v in edge_data.items()])))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def set_run_font(run, font_name="TH Sarabun PSK", size_pt=17, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run._r.get_or_add_rPr().append(
        parse_xml(r'<w:rFonts %s w:ascii="%s" w:hAnsi="%s" w:cs="%s"/>' % (nsdecls('w'), font_name, font_name, font_name))
    )

def generate_docx(output_path):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENTATION.PORTRAIT
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.59)
    sec.right_margin = Inches(0.59)
    sec.bottom_margin = Inches(0.59)
    sec.left_margin = Inches(0.98) # 2.5 cm

    c = bp["content"]
    sig = c["signature"]

    # Header Copy Label
    p_copy = doc.add_paragraph()
    p_copy.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_copy.paragraph_format.space_after = Pt(2)
    r_copy = p_copy.add_run(c.get("copy_type", "( สำเนาคู่ฉบับ )"))
    set_run_font(r_copy, "TH Sarabun PSK", 13)

    # Garuda + Title Table
    t_hdr = doc.add_table(rows=1, cols=2)
    t_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_img = t_hdr.cell(0, 0)
    c_title = t_hdr.cell(0, 1)
    c_img.width = Inches(1.2)
    c_title.width = Inches(5.5)

    p_img = c_img.paragraphs[0]
    p_img.paragraph_format.space_after = Pt(0)
    r_img = p_img.add_run()
    if os.path.exists(GARUDA_PATH):
        r_img.add_picture(GARUDA_PATH, width=Inches(1.18), height=Inches(1.18))

    p_t = c_title.paragraphs[0]
    p_t.paragraph_format.space_before = Pt(14)
    p_t.paragraph_format.space_after = Pt(0)
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_t.add_run(c["title"])
    set_run_font(r_t, "TH Sarabun PSK", 29, bold=True)

    set_cell_border(c_img)
    set_cell_border(c_title)

    # Meta Lines
    meta_lines = [
        (c["official_department"], c["department_value"]),
        (f'{c["doc_number_label"]}  {c["doc_number_value"]}', f'{c["doc_date_label"]}            {c["doc_date_month_year"]}'),
        (c["subject_label"], c["subject_value"]),
        (c["recipient_label"], c["recipient_value"]),
        (c["reference_label"], c["reference_value"])
    ]

    for label, val in meta_lines:
        p_m = doc.add_paragraph()
        p_m.paragraph_format.space_after = Pt(2)
        p_m.paragraph_format.line_spacing = 1.25
        r_lbl = p_m.add_run(label + "  ")
        set_run_font(r_lbl, "TH Sarabun PSK", 20, bold=True)
        r_val = p_m.add_run(val)
        set_run_font(r_val, "TH Sarabun PSK", 17)

    # Attachments
    p_att = doc.add_paragraph()
    p_att.paragraph_format.space_after = Pt(2)
    p_att.paragraph_format.line_spacing = 1.25
    r_al = p_att.add_run(c["attachments_label"] + "  ")
    set_run_font(r_al, "TH Sarabun PSK", 17, bold=True)
    r_av = p_att.add_run(f'{c["attachments"][0]["title"]}     {c["attachments"][0]["quantity"]}')
    set_run_font(r_av, "TH Sarabun PSK", 17)

    # Paragraphs
    for para in c["paragraphs"]:
        p_p = doc.add_paragraph()
        p_p.paragraph_format.first_line_indent = Inches(0.98)
        p_p.paragraph_format.space_after = Pt(4)
        p_p.paragraph_format.line_spacing = 1.25
        p_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_p = p_p.add_run(para)
        set_run_font(r_p, "TH Sarabun PSK", 17)

    # Signature Block via signature_engine.py
    signature_engine.render_signature_docx(
        doc,
        sig,
        axis_inch=4.13,
        font_name="TH Sarabun PSK",
        font_size=17,
        space_after_rank_pt=22,
        space_before_block_pt=55
    )

    # Admin Footer Block
    if "routing_footer" in c:
        rf = c["routing_footer"]
        day0 = rf[0].get("day", "")
        day1 = rf[1].get("day", "")
        p_adm1 = doc.add_paragraph()
        p_adm1.paragraph_format.space_before = Pt(30)
        p_adm1.paragraph_format.space_after = Pt(0)
        p_adm1.paragraph_format.tab_stops.add_tab_stop(Inches(4.5), WD_TAB_ALIGNMENT.LEFT)
        p_adm1.paragraph_format.tab_stops.add_tab_stop(Inches(5.5), WD_TAB_ALIGNMENT.LEFT)
        p_adm1.paragraph_format.tab_stops.add_tab_stop(Inches(6.4), WD_TAB_ALIGNMENT.LEFT)
        r_adm1 = p_adm1.add_run(f"\t{rf[0]['rank']}\t{rf[0]['role']}\t{rf[0]['date']}")
        set_run_font(r_adm1, "TH Sarabun PSK", 13.5)

        p_adm2 = doc.add_paragraph()
        p_adm2.paragraph_format.space_before = Pt(0)
        p_adm2.paragraph_format.space_after = Pt(0)
        p_adm2.paragraph_format.tab_stops.add_tab_stop(Inches(4.5), WD_TAB_ALIGNMENT.LEFT)
        p_adm2.paragraph_format.tab_stops.add_tab_stop(Inches(5.5), WD_TAB_ALIGNMENT.LEFT)
        p_adm2.paragraph_format.tab_stops.add_tab_stop(Inches(6.4), WD_TAB_ALIGNMENT.LEFT)
        r_adm2 = p_adm2.add_run(f"\t{rf[1]['rank']}\t{rf[1]['role']}\t{day1} {rf[1]['date']}")
        set_run_font(r_adm2, "TH Sarabun PSK", 13.5)

    doc.save(output_path)
    print(f"Generated DOCX: {output_path}")


# -----------------------------------------------------------------------------
# 3. GENERATE XLSX
# -----------------------------------------------------------------------------
def generate_xlsx(output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "บันทึกข้อความ"
    ws.views.sheetView[0].showGridLines = True

    c = bp["content"]
    sig = c["signature"]

    font_large_bold = Font(name="TH Sarabun PSK", size=24, bold=True)
    font_bold = Font(name="TH Sarabun PSK", size=17, bold=True)
    font_regular = Font(name="TH Sarabun PSK", size=17)
    align_center = Alignment(horizontal="center", vertical="center")

    ws["A1"] = c.get("copy_type", "( สำเนาคู่ฉบับ )")
    ws["A1"].font = Font(name="TH Sarabun PSK", size=13)
    ws["A1"].alignment = align_center
    ws.merge_cells("A1:E1")

    ws["B3"] = c["title"]
    ws["B3"].font = font_large_bold
    ws["B3"].alignment = align_center
    ws.merge_cells("B3:D3")

    ws["B5"] = c["official_department"]
    ws["B5"].font = font_bold
    ws["C5"] = c["department_value"]
    ws["C5"].font = font_regular

    ws["B6"] = c["doc_number_label"]
    ws["B6"].font = font_bold
    ws["C6"] = c["doc_number_value"]
    ws["C6"].font = font_regular
    ws["D6"] = c["doc_date_label"]
    ws["D6"].font = font_bold
    ws["E6"] = c["doc_date_month_year"]
    ws["E6"].font = font_regular

    ws["B7"] = c["subject_label"]
    ws["B7"].font = font_bold
    ws["C7"] = c["subject_value"]
    ws["C7"].font = font_regular

    ws["B8"] = c["recipient_label"]
    ws["B8"].font = font_bold
    ws["C8"] = c["recipient_value"]
    ws["C8"].font = font_regular

    ws["B9"] = c["reference_label"]
    ws["B9"].font = font_bold
    ws["C9"] = c["reference_value"]
    ws["C9"].font = font_regular

    ws["B10"] = c["attachments_label"]
    ws["B10"].font = font_bold
    ws["C10"] = f"{c['attachments'][0]['title']}   {c['attachments'][0]['quantity']}"
    ws["C10"].font = font_regular

    ws["B12"] = c["paragraphs"][0]
    ws["B12"].font = font_regular
    ws["B13"] = c["paragraphs"][1]
    ws["B13"].font = font_regular

    # Signature via signature_engine.py
    signature_engine.render_signature_xlsx(
        ws,
        sig,
        start_row=16,
        rank_col="D",
        name_col="E",
        font_name="TH Sarabun PSK",
        font_size=17
    )

    if "routing_footer" in c:
        rf = c["routing_footer"]
        day0 = rf[0].get("day", "")
        day1 = rf[1].get("day", "")
        ws["D22"] = rf[0]["rank"]
        ws["D22"].font = font_regular
        ws["E22"] = f"{rf[0]['role']}    {rf[0]['date']}"
        ws["E22"].font = font_regular

        ws["D23"] = rf[1]["rank"]
        ws["D23"].font = font_regular
        ws["E23"] = f"{rf[1]['role']}  {day1}    {rf[1]['date']}"
        ws["E23"].font = font_regular

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 25
    ws.column_dimensions["E"].width = 30

    wb.save(output_path)
    print(f"Generated XLSX: {output_path}")


# -----------------------------------------------------------------------------
# MAIN COMPILER & SYNC
# -----------------------------------------------------------------------------
if __name__ == "__main__":
    html_main = os.path.join(BASE_DIR, "บันทึกข้อความ_90_วัน_Rebuild_Thai.html")
    html_copy = os.path.join(BASE_DIR, "บันทึกข้อความ_90_วัน_สำเนาคู่ฉบับ.html")
    html_orig = os.path.join(BASE_DIR, "บันทึกข้อความ_90_วัน_ฉบับจริง.html")
    docx_out = os.path.join(BASE_DIR, "บันทึกข้อความ_90_วัน_Rebuild_Thai.docx")
    xlsx_out = os.path.join(BASE_DIR, "บันทึกข้อความ_90_วัน_Rebuild_Thai.xlsx")

    # Generate files
    generate_html(html_main, is_copy=True)
    generate_html(html_copy, is_copy=True)
    generate_html(html_orig, is_copy=False)
    generate_docx(docx_out)
    generate_xlsx(xlsx_out)

    # Sync to approved reference and static directories
    sync_targets = [
        os.path.join(BASE_DIR, "references", "approved_memo_90"),
        os.path.join(BASE_DIR, "static", "docs"),
        os.path.join(BASE_DIR, "frontend", "static", "docs")
    ]

    files_to_sync = [
        "บันทึกข้อความ_90_วัน_Rebuild_Thai.html",
        "บันทึกข้อความ_90_วัน_สำเนาคู่ฉบับ.html",
        "บันทึกข้อความ_90_วัน_ฉบับจริง.html",
        "บันทึกข้อความ_90_วัน_Rebuild_Thai.docx",
        "บันทึกข้อความ_90_วัน_Rebuild_Thai.xlsx",
        "table_layout_memo_90.json"
    ]

    for target_dir in sync_targets:
        if os.path.exists(target_dir):
            for fname in files_to_sync:
                src = os.path.join(BASE_DIR, fname)
                if os.path.exists(src):
                    dst = os.path.join(target_dir, fname)
                    shutil.copy2(src, dst)
            print(f"Synced files to: {target_dir}")

    print("Successfully built and synchronized all files for บันทึกข้อความ_90_วัน!")
